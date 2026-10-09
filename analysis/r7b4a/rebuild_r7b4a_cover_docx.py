#!/usr/bin/env python3
"""Rebuild the one-page R7B4A Human Genomics cover letter DOCX."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


BASE = Path(r"H:\SCI2\YR1\5_manuscript\R7B4A_HumanGenomics_SubmissionInterface")
SOURCE = BASE / "cover_letter" / "R7B4A_HumanGenomics_cover_letter_DRAFT.md"
OUTPUT = BASE / "cover_letter" / "R7B4A_HumanGenomics_cover_letter_DRAFT.docx"


def add_page_field(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    paragraph._p.append(field)


def main() -> None:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(1.65)
    section.bottom_margin = Cm(1.65)
    section.left_margin = Cm(2.1)
    section.right_margin = Cm(2.1)
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(3)
    normal.paragraph_format.line_spacing = 1.0

    for raw in SOURCE.read_text(encoding="utf-8").splitlines():
        line = raw.strip()
        if not line:
            continue
        if line.startswith("# "):
            paragraph = doc.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_after = Pt(4)
            run = paragraph.add_run(line[2:].replace("—", "–"))
            run.bold = True
            run.font.name = "Arial"
            run.font.size = Pt(13)
            run.font.color.rgb = RGBColor(31, 90, 133)
        elif line.startswith("- "):
            paragraph = doc.add_paragraph(style="List Bullet")
            paragraph.add_run(line[2:])
            paragraph.paragraph_format.left_indent = Cm(0.45)
            paragraph.paragraph_format.first_line_indent = Cm(-0.25)
            paragraph.paragraph_format.space_after = Pt(0)
            for run in paragraph.runs:
                run.font.size = Pt(9.5)
        else:
            clean = line.replace("**", "")
            clean = re.sub(r"\*([^*]+)\*", r"\1", clean).replace("`", "")
            paragraph = doc.add_paragraph(clean)
            if line.startswith("**"):
                paragraph.runs[0].bold = True

    add_page_field(section.footer.paragraphs[0])
    doc.core_properties.title = "Cover letter for Human Genomics"
    doc.core_properties.comments = "Draft with explicit author-confirmation fields; not signed."
    doc.save(OUTPUT)
    print(f"COVER_DOCX_OK\t{OUTPUT}\t{OUTPUT.stat().st_size}")


if __name__ == "__main__":
    main()
