#!/usr/bin/env python3
"""Build Human Genomics submission-interface assets from the frozen R7B3A master.

The script adapts presentation and journal-facing metadata only. It does not
change the frozen scientific universe, posterior results, simulations, or
external/tissue evidence.
"""

from __future__ import annotations

import csv
import hashlib
import json
import re
import textwrap
import zipfile
from pathlib import Path

import matplotlib as mpl
import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch
from PIL import Image
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(r"H:\SCI2\YR1")
SOURCE = ROOT / "5_manuscript" / "R7B3A_ManuscriptV2"
OUT = ROOT / "5_manuscript" / "R7B4A_HumanGenomics_SubmissionInterface"
CODE = ROOT / "2_code" / "11_manuscript" / "R7B4A"
TAG = "r7b4a-human-genomics-interface-2026-10-10"
REPO = "https://github.com/1209433622cz-maker/igan-open-regulatory-evidence"

for directory in [
    OUT / "manuscript",
    OUT / "figures",
    OUT / "graphical_abstract",
    OUT / "supplement",
    OUT / "cover_letter",
    OUT / "submission",
    OUT / "qa",
    OUT / "results",
]:
    directory.mkdir(parents=True, exist_ok=True)


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def write_tsv(path: Path, fieldnames: list[str], rows: list[dict[str, object]]) -> None:
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames, delimiter="\t")
        writer.writeheader()
        writer.writerows(rows)


def section(text: str, name: str) -> str:
    match = re.search(rf"(?ms)^## {re.escape(name)}\s*$\n(.*?)(?=^## |\Z)", text)
    if not match:
        raise ValueError(f"Missing section: {name}")
    return match.group(1).strip()


source_md = (SOURCE / "manuscript" / "R7B3A_full_english_manuscript_v2.md").read_text(encoding="utf-8")
old_title = source_md.splitlines()[0].removeprefix("# ").strip()
new_title = (
    "Input-matched multi-signal colocalization clarifies immune-cell regulatory "
    "assignments at primary biliary cholangitis risk loci"
)

abstract_background = (
    "Regulatory assignments at primary biliary cholangitis (PBC) risk loci can change with variant coverage, "
    "disease-statistic definitions and local signal architecture. We evaluated whether bidirectional "
    "colocalization reclassification persisted after matching inputs while separating statistical sharing from "
    "biological mechanism. A PBC-wide registry contained 6,923 locus-gene-cell comparisons; 5,460 met the "
    "single-causal overlap requirement and a frozen high-information subset of 642 underwent source-matched "
    "multi-signal analysis. Four arms retained the historical screen (A0), restricted variant support (A1), "
    "matched disease statistics (A2) and compared A2 with the frozen multi-signal result (M). We also evaluated "
    "signal-pair identity, PF10/PF50 sensitivity, 486,000 simulation iterations and bounded molecular-QTL and "
    "donor-level liver evidence."
)
abstract_results = (
    "Support restriction and disease-statistic matching changed 15 and 19 categorical assignments, respectively. "
    "With matched inputs, 79 of 113 approximate-Bayes-factor H4 comparisons retained shared-pair support, whereas "
    "eight supported distinct pairs; ten of 452 H3 comparisons acquired shared-pair support. Final PF10 states "
    "comprised 92 H4-supported, 428 H3-supported, 59 model-sensitive and 63 uninformative comparisons. Thirty-two "
    "H4-supported comparisons also contained a qualifying distinct pair. Across the frozen simulation grid, mean "
    "false-H4 decisions in distinct-signal scenarios decreased from 12.41% to 8.56%, while shared-signal recovery "
    "decreased from 32.99% to 26.66%; only 160,782 iterations (33.08%) formed an evaluable signal pair. IL12RB2-NK "
    "and FCRL3 B-cell assignments received bounded cross-resource support. Both target-lineage pairs were detected "
    "in all five PBC and five control livers without corrected PBC-specific enrichment (both q=0.111)."
)
abstract_conclusions = (
    "Input matching explained part, but not all, of the differences in PBC regulatory assignments. Shared and "
    "distinct signal pairs, covariate-model sensitivity and discovery limitations should be reported separately. "
    "The results do not establish causal truth, general method superiority or tissue mediation."
)

abstract_block = f"""## Abstract

### Background

{abstract_background}

### Results

{abstract_results}

### Conclusions

{abstract_conclusions}

**Keywords:** primary biliary cholangitis; colocalization; single-cell eQTL; fine-mapping; linkage disequilibrium; signal pair; regulatory genomics; statistical genetics

## Background"""

manuscript = source_md.replace(f"# {old_title}", f"# {new_title}", 1)
manuscript = re.sub(r"\*\*Article type:\*\* Research article\s*", "**Article type:** Research\n", manuscript, count=1)
manuscript = re.sub(
    r"(?ms)^## Abstract\s*$.*?^\*\*Keywords:\*\*.*?\n\n## Background",
    abstract_block,
    manuscript,
    count=1,
)

availability = f"""### Availability of data and materials

PBC summary statistics are available through GWAS Catalog study GCST90061440. Study-derived PBC regional statistics and LD are available from the GJOKA archive. The OneK1K cohort, updated release and TenK10K resource are identified separately in their cited publications and Zenodo records. HRA008003 is available through GSA-Human. Code, compact aggregate outputs, manifests, frozen protocols, source-bound figures and submission-interface assets are versioned in the public repository at `{REPO}/tree/{TAG}` [18]. Additional files 1 and 2 contain the released supplementary workbook and machine-readable derived tables. Individual-level genotype, BAM, pseudobulk matrices and source archives governed by third-party licences are not redistributed; their source accessions and conditions are recorded in Supplementary Table S1.

### Competing interests"""
manuscript = re.sub(
    r"(?ms)^### Availability of data and materials\s*$.*?^### Competing interests",
    availability,
    manuscript,
    count=1,
)

manuscript = manuscript.replace(
    "(F) Stable signal-pair semantics, including 32 comparisons with qualifying shared and distinct pairs.",
    "(F) Stable signal-pair semantics: 60 stable-H4 comparisons had no qualifying H3 pair under the frozen thresholds, whereas 32 contained qualifying shared and distinct pairs.",
)

supp_details = section(manuscript, "Supplementary materials")
additional_files = f"""## Additional files

**Additional file 1.** `Additional_file_1_Supplementary_Tables_S1-S10.xlsx` (Microsoft Excel workbook). Editable reading interface for Supplementary Tables S1-S10.

**Additional file 2.** `Additional_file_2_Machine_Readable_Supplementary_Data.zip` (ZIP archive). Machine-readable TSV, TSV.GZ and JSON source/derived tables, plus a file-level SHA-256 manifest and README. The archive excludes individual-level genotype, BAM, pseudobulk matrices and licensed source archives.

The two files contain the following prespecified supplementary modules:

{supp_details}
"""
manuscript = re.sub(r"(?ms)^## Supplementary materials\s*$.*\Z", additional_files.rstrip() + "\n", manuscript)

manuscript_path = OUT / "manuscript" / "R7B4A_HumanGenomics_manuscript.md"
manuscript_path.write_text(manuscript, encoding="utf-8")

abstract_text = " ".join([abstract_background, abstract_results, abstract_conclusions])
abstract_words = len(re.findall(r"\b[\w%.-]+\b", abstract_text))
if abstract_words > 350:
    raise RuntimeError(f"Human Genomics abstract exceeds 350 words: {abstract_words}")
if "### Methods" in section(manuscript, "Abstract"):
    raise RuntimeError("Human Genomics abstract must not contain a Methods subheading")


# Copy the editable supplement workbook and build a separate machine-readable bundle.
supp_out = OUT / "supplement"
source_supp = SOURCE / "supplements"
workbook_out = supp_out / "Additional_file_1_Supplementary_Tables_S1-S10.xlsx"
workbook_out.write_bytes((source_supp / "R7B3A_Supplementary_Tables_S1-S10.xlsx").read_bytes())

manifest_rows = list(csv.DictReader((source_supp / "SUPPLEMENT_FINAL_MANIFEST.tsv").open(encoding="utf-8"), delimiter="\t"))
bundle_path = supp_out / "Additional_file_2_Machine_Readable_Supplementary_Data.zip"
bundle_readme = textwrap.dedent(
    """\
    R7B4A machine-readable supplementary data

    This bundle accompanies the Human Genomics submission-interface manuscript.
    It contains public source/derived tables from Supplementary Tables S1-S10.
    The editable XLSX workbook is supplied separately as Additional file 1.
    Individual-level genotype, BAM, pseudobulk matrices and licensed source archives
    are not included. See S1 and SUPPLEMENT_FINAL_MANIFEST.tsv for provenance.
    """
)
with zipfile.ZipFile(bundle_path, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9) as archive:
    archive.writestr("README.txt", bundle_readme)
    archive.write(source_supp / "SUPPLEMENT_FINAL_MANIFEST.tsv", "SUPPLEMENT_FINAL_MANIFEST.tsv")
    for row in manifest_rows:
        rel = row["relative_path"]
        if rel.endswith(".xlsx"):
            continue
        archive.write(source_supp / Path(rel), rel)
with zipfile.ZipFile(bundle_path) as archive:
    if archive.testzip() is not None:
        raise RuntimeError("Additional file 2 failed ZIP CRC")


# Build the mandatory 920 x 300 graphical abstract.
mpl.rcParams.update({"font.family": "Arial", "font.size": 9})
fig = plt.figure(figsize=(9.2, 3.0), dpi=100, facecolor="white")
ax = fig.add_axes([0, 0, 1, 1])
ax.set_xlim(0, 1); ax.set_ylim(0, 1); ax.axis("off")
colors = ["#1F5A85", "#2A9D8F", "#7A5195"]
titles = ["Frozen PBC scope", "Input-matched attribution", "Bounded interpretation"]
bodies = [
    "6,923 screened\n5,460 ABF-eligible\n642 high-information",
    "A0 → A1 → A2 → M\nsupport, disease statistics,\nmulti-signal inference",
    "stable / reclassified pairs\nexternal QTL support\ntissue null boundary",
]
for i, (x, color, title, body) in enumerate(zip([0.025, 0.355, 0.685], colors, titles, bodies)):
    ax.add_patch(FancyBboxPatch((x, 0.24), 0.29, 0.55, boxstyle="round,pad=0.012", fc="#F7F9FB", ec=color, lw=2))
    ax.add_patch(FancyBboxPatch((x, 0.67), 0.29, 0.12, boxstyle="round,pad=0.012", fc=color, ec=color, lw=0))
    ax.text(x + 0.145, 0.73, title, ha="center", va="center", color="white", weight="bold", fontsize=10)
    ax.text(x + 0.145, 0.47, body, ha="center", va="center", color="#22313F", fontsize=9, linespacing=1.35)
    if i < 2:
        ax.annotate("", xy=(x + 0.325, 0.515), xytext=(x + 0.295, 0.515), arrowprops=dict(arrowstyle="-|>", lw=1.8, color="#667785"))
ax.text(0.5, 0.91, "Reliable cell-regulatory attribution requires matched inputs and signal-pair reporting", ha="center", va="center", fontsize=12, weight="bold", color="#22313F")
ax.text(0.5, 0.10, "No claim of general method superiority, complete regulatory cascade, or PBC-specific tissue mediation", ha="center", va="center", fontsize=8.3, color="#C44E52")
ga_png = OUT / "graphical_abstract" / "Graphical_Abstract_HumanGenomics_920x300.png"
ga_svg = OUT / "graphical_abstract" / "Graphical_Abstract_HumanGenomics_920x300.svg"
fig.savefig(ga_png, dpi=100, facecolor="white")
fig.savefig(ga_svg, facecolor="white")
plt.close(fig)
if Image.open(ga_png).size != (920, 300):
    with Image.open(ga_png) as image:
        image.resize((920, 300), Image.Resampling.LANCZOS).save(ga_png)
if Image.open(ga_png).size != (920, 300):
    raise RuntimeError(f"Graphical abstract dimensions are {Image.open(ga_png).size}")


cover_md = f"""# Cover letter — Human Genomics

**Date:** 10 October 2026

**To:** Professor Vasilis Vasiliou, Editor-in-Chief, *Human Genomics*

**From:** [CORRESPONDING AUTHOR NAME, DEGREE, AFFILIATION, POSTAL ADDRESS, EMAIL, TELEPHONE]

**Re:** Research manuscript, “{new_title}”

Dear Professor Vasiliou,

Please consider our Research manuscript, “{new_title},” for publication in *Human Genomics*. The study addresses a recurring problem in post-GWAS interpretation: changes in colocalization labels can arise from variant-support and disease-statistic differences as well as from the move from single-causal to multi-signal inference.

We screened 6,923 PBC locus-gene-cell comparisons and applied source-matched multi-signal analysis to a frozen 642-comparison high-information subset. A four-arm A0/A1/A2/M design separated support restriction, disease-statistic matching and inferential procedure. Input matching explained part of the bidirectional reclassification, but stable H4-to-H3 and H3-to-H4 changes remained. Truth-known simulations exposed a trade-off between fewer false-H4 decisions and lower shared-signal recovery, with only one-third of iterations forming an evaluable signal pair. IL12RB2-NK and FCRL3 B-cell assignments received bounded cross-resource support, while the exact five-versus-five liver analysis supported lineage detectability without corrected PBC-specific enrichment.

The contribution is an evidence-governed analysis of regulatory-assignment reliability rather than a claim of new-gene discovery or universal method superiority. It fits *Human Genomics* through its integration of human genetic epidemiology, regulatory genomics, statistical genetics, single-cell molecular-QTL resources and transparent open-data reproducibility. The manuscript retains negative, model-sensitive and uninformative results and distinguishes statistical signal sharing from mechanism and mediation.

Code, frozen protocols, compact aggregate outputs, figure sources and machine-readable supplementary data are fixed in the public repository at `{REPO}/tree/{TAG}`. Third-party individual-level or licence-restricted files are represented by accession and provenance records rather than redistributed.

Before submission, the corresponding author must replace and confirm the following statements:

- [ALL AUTHORS APPROVE THIS VERSION AND ITS SUBMISSION.]
- [THE MANUSCRIPT IS ORIGINAL AND IS NOT UNDER CONSIDERATION ELSEWHERE.]
- [ALL FINANCIAL AND NON-FINANCIAL COMPETING INTERESTS ARE DISCLOSED.]
- [ETHICS/WAIVER WORDING HAS BEEN CONFIRMED BY THE RESPONSIBLE INSTITUTION.]
- [FUNDING, AUTHOR CONTRIBUTIONS, ACKNOWLEDGEMENTS, APC ROUTE AND LICENCE ARE CONFIRMED.]

Thank you for considering this manuscript.

Sincerely,

[CORRESPONDING AUTHOR NAME AND SIGNATURE BLOCK]
"""
cover_md_path = OUT / "cover_letter" / "R7B4A_HumanGenomics_cover_letter_DRAFT.md"
cover_md_path.write_text(cover_md, encoding="utf-8")


def add_page_field(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    paragraph._p.append(fld)


def cover_to_docx(text: str, output: Path) -> None:
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Cm(1.65); sec.bottom_margin = Cm(1.65); sec.left_margin = Cm(2.1); sec.right_margin = Cm(2.1)
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"; normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman"); normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(3)
    normal.paragraph_format.line_spacing = 1.0
    lines = text.splitlines()
    for raw in lines:
        line = raw.strip()
        if not line:
            continue
        if line.startswith("# "):
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(line[2:].replace("—", "–")); run.bold = True; run.font.name = "Arial"; run.font.size = Pt(13); run.font.color.rgb = RGBColor(31, 90, 133)
            p.paragraph_format.space_after = Pt(4)
        elif line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet"); p.add_run(line[2:])
            p.paragraph_format.left_indent = Cm(0.45); p.paragraph_format.first_line_indent = Cm(-0.25)
            p.paragraph_format.space_after = Pt(0)
            for run in p.runs:
                run.font.size = Pt(9.5)
        else:
            clean = line.replace("**", "")
            clean = re.sub(r"\*([^*]+)\*", r"\1", clean).replace("`", "")
            p = doc.add_paragraph(clean)
            if line.startswith("**"):
                p.runs[0].bold = True
    add_page_field(sec.footer.paragraphs[0])
    doc.core_properties.title = "Cover letter — Human Genomics"
    doc.core_properties.comments = "Draft with explicit author-confirmation fields; not signed."
    doc.save(output)


cover_docx = OUT / "cover_letter" / "R7B4A_HumanGenomics_cover_letter_DRAFT.docx"
cover_to_docx(cover_md, cover_docx)


guideline_rows = [
    {"journal": "Human Genomics", "requirement": "Scope", "observed": "Human genetic epidemiology; GWAS; regulatory genomics; statistical genetics; single-cell and integrative multi-omics", "source_type": "official journal", "url": "https://link.springer.com/journal/40246/aims-and-scope", "accessed": "2026-10-10"},
    {"journal": "Human Genomics", "requirement": "Article type", "observed": "Research; original primary research", "source_type": "official journal", "url": "https://link.springer.com/journal/40246/submission-guidelines/research", "accessed": "2026-10-10"},
    {"journal": "Human Genomics", "requirement": "Abstract", "observed": "Background/Results/Conclusions; <=350 words; no citations", "source_type": "official journal", "url": "https://link.springer.com/journal/40246/submission-guidelines/research", "accessed": "2026-10-10"},
    {"journal": "Human Genomics", "requirement": "Keywords", "observed": "3-10", "source_type": "official journal", "url": "https://link.springer.com/journal/40246/submission-guidelines/research", "accessed": "2026-10-10"},
    {"journal": "Human Genomics", "requirement": "Graphical abstract", "observed": "mandatory; landscape; approximately 920x300 px; JPEG/PNG/SVG", "source_type": "official journal", "url": "https://link.springer.com/journal/40246/submission-guidelines/research", "accessed": "2026-10-10"},
    {"journal": "Human Genomics", "requirement": "Declarations", "observed": "all required headings; ethics statement required even if approval was waived", "source_type": "official journal", "url": "https://link.springer.com/journal/40246/submission-guidelines/research", "accessed": "2026-10-10"},
    {"journal": "Human Genomics", "requirement": "LLM disclosure", "observed": "document supervised LLM use in Methods; LLMs cannot be authors", "source_type": "official journal", "url": "https://link.springer.com/journal/40246/submission-guidelines/research", "accessed": "2026-10-10"},
    {"journal": "Human Genomics", "requirement": "Metric", "observed": "2025 JIF 4.1; median first decision 6 days", "source_type": "official journal", "url": "https://link.springer.com/journal/40246", "accessed": "2026-10-10"},
    {"journal": "Human Genomics", "requirement": "APC", "observed": "GBP 2590 / USD 3390 / EUR 2890; taxes may apply", "source_type": "official journal", "url": "https://link.springer.com/journal/40246/how-to-publish-with-us", "accessed": "2026-10-10"},
    {"journal": "Genetic Epidemiology", "requirement": "Current abstract and keywords", "observed": "200-word abstract; up to 7 keywords; current forauthors page supersedes old 250-word guide", "source_type": "official journal", "url": "https://onlinelibrary.wiley.com/page/journal/10982272/homepage/forauthors.html", "accessed": "2026-10-10"},
    {"journal": "Genetic Epidemiology", "requirement": "Metric", "observed": "JIF 3.4; CiteScore 5.5; acceptance 9%; median first decision 27 days", "source_type": "official journal", "url": "https://onlinelibrary.wiley.com/journal/10982272", "accessed": "2026-10-10"},
    {"journal": "BMC Medical Genomics", "requirement": "Scope and metric", "observed": "human health/disease genomics including immunogenomics and bioinformatics; 2025 JIF 2.6; median first decision 3 days", "source_type": "official journal", "url": "https://link.springer.com/journal/12920", "accessed": "2026-10-10"},
]
write_tsv(OUT / "results" / "R7B4A_official_guideline_audit.tsv", list(guideline_rows[0]), guideline_rows)

journal_rows = [
    {"priority": "PRIMARY_LOCKED", "journal": "Human Genomics", "article_type": "Research", "scope_fit": "HIGH", "publisher_metric": "2025 JIF 4.1", "quartile_evidence": "JCR Q2 in current public JCR-derived 2026 list; SCImago 2025 Q1", "quartile_source_caveat": "Clarivate JCR direct subscriber record was not accessible; institutional confirmation advised immediately before submission", "format_delta": "completed in R7B4A", "main_risk": "editorial novelty; mandatory OA APC", "decision": "GO_AFTER_HUMAN_METADATA"},
    {"priority": "BACKUP_1_COMPLETION_FIRST", "journal": "BMC Medical Genomics", "article_type": "Research", "scope_fit": "HIGH", "publisher_metric": "2025 JIF 2.6", "quartile_evidence": "current public JCR-derived and SCImago sources report Q2", "quartile_source_caveat": "institutional JCR/CAS system must be checked for the user's exact evaluation rule", "format_delta": "minor BMC retargeting", "main_risk": "lower journal priority; OA APC", "decision": "READY_AFTER_RETARGET"},
    {"priority": "BACKUP_2_METHOD_FACING", "journal": "Genetic Epidemiology", "article_type": "Research Article", "scope_fit": "MODERATE_HIGH", "publisher_metric": "current JIF 3.4", "quartile_evidence": "current public JCR-derived 2026 list reports Q2", "quartile_source_caveat": "institutional confirmation advised", "format_delta": "200-word abstract; <=7 keywords; method-facing cover letter; Wiley upload interface", "main_risk": "higher method-innovation threshold; study is an applied reliability analysis", "decision": "HOLD_AS_SECONDARY_BACKUP"},
]
write_tsv(OUT / "results" / "R7B4A_journal_route_matrix.tsv", list(journal_rows[0]), journal_rows)

upload_rows = [
    {"order": 1, "submission_designation": "Main manuscript", "file": "manuscript/R7B4A_HumanGenomics_manuscript.docx", "status": "TECHNICALLY_READY_HUMAN_FIELDS_PENDING", "note": "Replace author/declaration placeholders before upload"},
    {"order": 2, "submission_designation": "Graphical abstract", "file": "graphical_abstract/Graphical_Abstract_HumanGenomics_920x300.png", "status": "READY", "note": "920x300 pixels; SVG also supplied"},
    {"order": 3, "submission_designation": "Figure 1", "file": "figures/Figure1_study_scope_and_model_integrity.png", "status": "READY", "note": "600 dpi source-bound composite"},
    {"order": 4, "submission_designation": "Figure 2", "file": "figures/Figure2_input_matched_reclassification.png", "status": "READY", "note": "semantic label corrected; counts unchanged"},
    {"order": 5, "submission_designation": "Figure 3", "file": "figures/Figure3_simulation_discovery_and_inference.png", "status": "READY", "note": "600 dpi source-bound composite"},
    {"order": 6, "submission_designation": "Figure 4", "file": "figures/Figure4_IL12RB2_external_evidence.png", "status": "READY", "note": "bounded chromatin relationships"},
    {"order": 7, "submission_designation": "Figure 5", "file": "figures/Figure5_FCRL3_support_and_counterexample.png", "status": "READY", "note": "exact A0/A1/A2/M values"},
    {"order": 8, "submission_designation": "Figure 6", "file": "figures/Figure6_liver_tissue_boundary.png", "status": "READY", "note": "deterministic donor display"},
    {"order": 9, "submission_designation": "Additional file 1", "file": "supplement/Additional_file_1_Supplementary_Tables_S1-S10.xlsx", "status": "READY", "note": "editable workbook"},
    {"order": 10, "submission_designation": "Additional file 2", "file": "supplement/Additional_file_2_Machine_Readable_Supplementary_Data.zip", "status": "READY", "note": "public derived tables only"},
    {"order": 11, "submission_designation": "Cover letter", "file": "cover_letter/R7B4A_HumanGenomics_cover_letter_DRAFT.docx", "status": "HUMAN_CONFIRMATION_PENDING", "note": "complete signature and confirmation fields"},
    {"order": 12, "submission_designation": "Reviewer PDF", "file": "manuscript/R7B4A_HumanGenomics_manuscript_WPS.pdf", "status": "READY_AFTER_WPS_EXPORT", "note": "not a substitute for editable manuscript"},
]
write_tsv(OUT / "submission" / "R7B4A_HumanGenomics_upload_map.tsv", list(upload_rows[0]), upload_rows)

author_form = """# Human Genomics author-completion form

Complete every field with verified information. Unknown fields must remain explicit; do not infer them from other projects.

## Title page

| Field | Required verified value |
|---|---|
| Final author order | Full publication names in agreed order |
| Affiliations | Department, institution, city, postal code and country |
| Author-affiliation mapping | Superscript affiliation number(s) for every author |
| Corresponding author | Name, postal address, institutional email and telephone |
| ORCID | Verified ORCID for each author where available |

## Authorship and CRediT

Assign verified initials to: Conceptualization; Data curation; Formal analysis; Funding acquisition; Investigation; Methodology; Project administration; Resources; Software; Supervision; Validation; Visualization; Writing—original draft; Writing—review and editing. Every author must approve the submitted version and accept accountability.

## Declarations

- Ethics/local determination: committee/institution, decision, reference number and exact approved wording, including waiver or exemption if applicable.
- Consent for publication: confirm that no identifiable individual data are presented and approve the final statement.
- Competing interests: obtain a financial and non-financial declaration from every author.
- Funding: official funder name, grant number, recipient initials and funder role.
- Acknowledgements: names and permission from acknowledged contributors, or “Not applicable”.
- AI disclosure: confirm the supervised OpenAI Codex/QiTeng statement in Methods.
- Originality: confirm the manuscript is not published or under review elsewhere.
- Final approval: dated confirmation from all authors.

## Publication choices

- Confirm APC funding, institutional agreement, discount or waiver request before submission.
- Confirm licence selection (CC BY or CC BY-NC-ND) against funder rules.
- Optional reviewer suggestions/exclusions must be verified and conflict-checked.
"""
(OUT / "submission" / "R7B4A_author_metadata_completion_form.md").write_text(author_form, encoding="utf-8")

checklist = f"""# Human Genomics Research submission checklist — verified 10 October 2026

## Completed technical items

- [x] Primary route locked as *Human Genomics*, article type Research.
- [x] Scope matched to human genetic epidemiology, GWAS, regulatory genomics, statistical genetics and integrative single-cell genomics.
- [x] Title contains no abbreviation and describes the analytical design.
- [x] Abstract uses Background, Results and Conclusions only and contains {abstract_words} words (limit 350).
- [x] Eight keywords supplied (allowed range 3–10).
- [x] Main text retains Background, Methods, Results, Discussion and Conclusions.
- [x] All required Declarations headings are present.
- [x] Supervised LLM use is documented in Methods; no LLM is listed as an author.
- [x] Graphical abstract supplied as 920×300 PNG and editable SVG.
- [x] Six source-bound figures supplied separately as PNG, PDF and SVG.
- [x] Figure 2F wording changed to “no qualifying H3 pair”; counts and classifications are unchanged.
- [x] Additional file 1 contains the editable S1–S10 workbook.
- [x] Additional file 2 contains machine-readable public derived tables and a SHA-256 manifest.
- [x] Immutable public repository tag specified as `{TAG}`.
- [x] Cover letter draft and upload map generated.

## Blocking human inputs

- [ ] Replace author, affiliation and corresponding-author placeholders.
- [ ] Complete CRediT contributions and obtain all-author approval.
- [ ] Confirm ethics/waiver wording with the responsible institution and provide reference details if applicable.
- [ ] Complete funding, competing interests and acknowledgements.
- [ ] Confirm consent-for-publication wording.
- [ ] Confirm originality and absence of concurrent submission.
- [ ] Confirm AI-assistance disclosure with all authors.
- [ ] Decide APC funding/waiver/discount route and licence.
- [ ] Add only verified reviewer suggestions or exclusions, if desired.
- [ ] Recheck the institution's required JCR/CAS partition immediately before submission.

## Official pages

- Research instructions: https://link.springer.com/journal/40246/submission-guidelines/research
- Aims and scope: https://link.springer.com/journal/40246/aims-and-scope
- Metrics: https://link.springer.com/journal/40246
- Fees and funding: https://link.springer.com/journal/40246/how-to-publish-with-us
"""
(OUT / "submission" / "R7B4A_HumanGenomics_submission_checklist.md").write_text(checklist, encoding="utf-8")

state = {
    "stage": "R7B4A_TARGET_JOURNAL_AND_SUBMISSION_INTERFACE",
    "date": "2026-10-10",
    "status": "TECHNICAL_INTERFACE_COMPLETE_HUMAN_GATE_PENDING",
    "primary_journal": "Human Genomics",
    "article_type": "Research",
    "primary_route": "LOCKED",
    "Q2_eligibility": "PASS_WITH_SOURCE_CAVEAT_AND_INSTITUTIONAL_RECHECK",
    "new_biological_analysis": "NO",
    "scientific_master_changed": "NO",
    "journal_specific_manuscript": "CREATED",
    "graphical_abstract": "CREATED_920x300",
    "additional_files": 2,
    "abstract_words": abstract_words,
    "author_metadata": "HUMAN_COMPLETION_GATE",
    "next": "R7B4B_AUTHOR_COMPLETION_AND_FINAL_SUBMISSION_QA",
}
(OUT / "results" / "R7B4A_state.json").write_text(json.dumps(state, indent=2), encoding="utf-8")

print(f"MANUSCRIPT={manuscript_path}")
print(f"ABSTRACT_WORDS={abstract_words}")
print(f"GRAPHICAL_ABSTRACT={Image.open(ga_png).size}")
print(f"ADDITIONAL_FILE_2_CRC=PASS")
print(f"STATE={state['status']}")
