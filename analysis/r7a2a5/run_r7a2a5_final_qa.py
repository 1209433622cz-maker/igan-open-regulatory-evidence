from __future__ import annotations

import hashlib
import json
import re
import zipfile
from pathlib import Path

from docx import Document
from openpyxl import load_workbook
from PIL import Image
from pypdf import PdfReader


ROOT = Path(r"H:\SCI2\YR1")
BASE = ROOT / "5_manuscript" / "R7A2A5_HumanGenomics"
MD = BASE / "manuscript" / "R7A2A5_HumanGenomics_manuscript_v3.md"
DOCX = BASE / "manuscript" / "R7A2A5_HumanGenomics_manuscript_v3.docx"
PDF = BASE / "manuscript" / "R7A2A5_HumanGenomics_manuscript_v3_WPS.pdf"
COVER_DOCX = BASE / "cover_letter" / "R7A2A5_HumanGenomics_cover_letter_DRAFT.docx"
COVER_PDF = BASE / "cover_letter" / "R7A2A5_HumanGenomics_cover_letter_DRAFT_WPS.pdf"
XLSX = BASE / "supplement" / "Additional_file_1_R7A2A5_Supplementary_Tables.xlsx"
INPUT_ZIP = Path(r"C:\Users\Administrator\Downloads\CMM_R7A2A4_HostileManuscriptAudit_2026-09-30.zip")
EXPECTED_INPUT_SHA = "c7c50a3637068d13d593f10b182fc0f26c715d8ef53e9d6fdf228c34dd3daaf9"
EXPECTED_SHEETS = [
    "Index", "S1 Sources", "S2 Universe", "S3 ABF screen", "S4 LD QC",
    "S5 Multi-signal", "S6 Credible sets", "S7 TenK replication",
    "S8 Liver donors", "S9 Liver sensitivity", "S10 Claims",
]


checks: list[dict[str, str]] = []


def check(name: str, ok: bool, detail: str, blocking: bool = True) -> None:
    checks.append({
        "check": name,
        "status": "PASS" if ok else ("FAIL" if blocking else "BLOCKED_HUMAN_INPUT"),
        "detail": detail,
    })


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


for path in [MD, DOCX, PDF, COVER_DOCX, COVER_PDF, XLSX]:
    check(f"file exists: {path.name}", path.is_file() and path.stat().st_size > 0,
          f"{path.stat().st_size if path.exists() else 0} bytes")

input_sha = sha256(INPUT_ZIP)
check("R7A2A4 input identity", input_sha == EXPECTED_INPUT_SHA, input_sha)

text = MD.read_text(encoding="utf-8")
title = text.splitlines()[0].removeprefix("# ").strip()
check("title <=20 words", len(title.split()) <= 20, f"{len(title.split())} words")

abstract_match = re.search(r"## Abstract\s+(.*?)\s+## Background", text, re.S)
abstract_text = abstract_match.group(1) if abstract_match else ""
abstract_words = re.findall(r"\b[\w'-]+\b", re.sub(r"^###\s+.*$", "", abstract_text, flags=re.M))
check("structured abstract <=350 words", 0 < len(abstract_words) <= 350, f"{len(abstract_words)} words")
for heading in ["### Background", "### Results", "### Conclusions"]:
    check(f"abstract heading {heading[4:]}", heading in abstract_text, heading)

for heading in ["## Background", "## Methods", "## Results", "## Discussion", "## Conclusions",
                "## List of abbreviations", "## Declarations", "## References", "## Figure legends"]:
    check(f"section {heading[3:]}", heading in text, heading)

for heading in [
    "### Ethics approval and consent to participate", "### Consent for publication",
    "### Availability of data and materials", "### Competing interests", "### Funding",
    "### Authors' contributions", "### Acknowledgements", "### Authors' information",
]:
    check(f"declaration {heading[4:]}", heading in text, heading)

check("claim ceiling retained", "expression mediation, biological mechanism, and therapeutic relevance remain untested" in text,
      "explicit association/signal-sharing ceiling")
check("tissue null boundary retained", "did not support corrected PBC-specific enrichment" in text,
      "detectability separated from enrichment")
check("TenK disease-replication boundary retained", "not independent disease-cohort replication" in text,
      "same PBC GWAS explicitly disclosed")
check("external preregistration not claimed", "preregistered" not in text.lower(), "prespecified/protocol-frozen wording")
check("QiTeng disclosure present", "QiTeng Academic Writing Skill" in text, "supervised writing disclosure")

body_before_refs = text.split("## References", 1)[0]
seen: list[int] = []
for token in re.findall(r"\[([0-9,;\-– ]+)\]", body_before_refs):
    for part in re.split(r"[,;]", token):
        part = part.strip()
        if not part:
            continue
        if re.fullmatch(r"\d+", part):
            nums = [int(part)]
        else:
            m = re.fullmatch(r"(\d+)\s*[-–]\s*(\d+)", part)
            nums = list(range(int(m.group(1)), int(m.group(2)) + 1)) if m else []
        for n in nums:
            if n not in seen:
                seen.append(n)
ref_nums = [int(x) for x in re.findall(r"(?m)^(\d+)\.\s", text.split("## References", 1)[1].split("## Figure legends", 1)[0])]
check("reference continuity", ref_nums == list(range(1, 20)), f"references={ref_nums[:3]}...{ref_nums[-3:] if ref_nums else []}")
check("citation first-appearance order", seen == list(range(1, 20)), f"first appearances={seen}")

with zipfile.ZipFile(DOCX) as zf:
    document_xml = zf.read("word/document.xml").decode("utf-8", errors="replace")
    settings_xml = zf.read("word/settings.xml").decode("utf-8", errors="replace")
    all_word_xml = "\n".join(
        zf.read(name).decode("utf-8", errors="replace")
        for name in zf.namelist() if name.startswith("word/") and name.endswith(".xml")
    )
check("continuous line numbering", "w:lnNumType" in document_xml and 'w:restart="continuous"' in document_xml,
      "DOCX section properties")
check("page numbering field", "PAGE" in all_word_xml, "DOCX footer field")
doc = Document(DOCX)
normal = doc.styles["Normal"]
spacing = normal.paragraph_format.line_spacing
check("double-spaced normal style", spacing == 2.0, f"line spacing={spacing}")

main_pdf = PdfReader(str(PDF))
cover_pdf = PdfReader(str(COVER_PDF))
check("WPS manuscript PDF", len(main_pdf.pages) == 27, f"{len(main_pdf.pages)} pages; creator={main_pdf.metadata.get('/Creator')}")
check("one-page WPS cover letter", len(cover_pdf.pages) == 1, f"{len(cover_pdf.pages)} page; creator={cover_pdf.metadata.get('/Creator')}")
check("PDF text extraction", sum(len(p.extract_text() or "") for p in main_pdf.pages) > 45000,
      f"{sum(len(p.extract_text() or '') for p in main_pdf.pages)} characters")

ga = Image.open(BASE / "graphical_abstract" / "Graphical_Abstract_R7A2A5_HumanGenomics_920x300.png")
check("graphical abstract 920x300", ga.size == (920, 300), f"{ga.size[0]}x{ga.size[1]}")

for i in range(1, 7):
    path = BASE / "figures" / f"Figure{i}.png"
    im = Image.open(path)
    dpi = im.info.get("dpi", (0, 0))[0]
    check(f"Figure {i} technical gate", path.stat().st_size < 10 * 1024 * 1024 and dpi >= 295,
          f"{im.size[0]}x{im.size[1]}, {dpi:.1f} dpi, {path.stat().st_size} bytes")

for i, path in enumerate(sorted((BASE / "supplement").glob("Supplementary_Figure_S*.png")), 1):
    im = Image.open(path)
    dpi = im.info.get("dpi", (0, 0))[0]
    check(f"Supplementary Figure S{i}", dpi >= 295, f"{im.size[0]}x{im.size[1]}, {dpi:.1f} dpi")

wb = load_workbook(XLSX, read_only=True, data_only=False)
check("supplement workbook sheet set", wb.sheetnames == EXPECTED_SHEETS, f"{len(wb.sheetnames)} sheets")
formula_count = 0
formula_errors = 0
for ws in wb.worksheets:
    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and cell.value.startswith("="):
                formula_count += 1
            if isinstance(cell.value, str) and cell.value in {"#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A"}:
                formula_errors += 1
check("supplement formula/error scan", formula_errors == 0, f"formulas={formula_count}; errors={formula_errors}")

placeholder_count = len(re.findall(r"\[[A-Z][A-Z0-9 /&,_-]{4,}\]", text))
check("human metadata complete", placeholder_count == 0, f"{placeholder_count} explicit manuscript placeholders remain", blocking=False)

pass_count = sum(x["status"] == "PASS" for x in checks)
fail_count = sum(x["status"] == "FAIL" for x in checks)
blocked_count = sum(x["status"] == "BLOCKED_HUMAN_INPUT" for x in checks)
summary = {
    "schema": "R7A2A5_FINAL_QA_1.0",
    "date": "2026-09-30",
    "pass": pass_count,
    "fail": fail_count,
    "blocked_human_input": blocked_count,
    "machine_gate": "PASS" if fail_count == 0 else "FAIL",
    "submission_state": "AWAITING_AUTHOR_METADATA" if fail_count == 0 and blocked_count else "READY" if fail_count == 0 else "BLOCKED_TECHNICAL",
    "checks": checks,
}

out_json = BASE / "results" / "R7A2A5_final_QA.json"
out_tsv = BASE / "results" / "R7A2A5_final_QA.tsv"
out_json.write_text(json.dumps(summary, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
out_tsv.write_text(
    "check\tstatus\tdetail\n" + "\n".join(
        f"{c['check']}\t{c['status']}\t{c['detail'].replace(chr(9), ' ')}" for c in checks
    ) + "\n",
    encoding="utf-8",
)
print(json.dumps({k: summary[k] for k in ["pass", "fail", "blocked_human_input", "machine_gate", "submission_state"]}, ensure_ascii=False))
if fail_count:
    raise SystemExit(1)
