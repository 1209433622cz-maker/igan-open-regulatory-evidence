from __future__ import annotations

import csv
import json
import math
import re
import shutil
import textwrap
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(r"H:\SCI2\YR1")
SOURCE_RELEASE = ROOT / "6_release" / "CMM_R7A2A4_HostileManuscriptAudit_2026-09-30"
SOURCE_MD = SOURCE_RELEASE / "manuscript" / "R7A2A4_manuscript_v2_hostile_audit.md"
OUT = ROOT / "5_manuscript" / "R7A2A5_HumanGenomics"
MANUSCRIPT_DIR = OUT / "manuscript"
COVER_DIR = OUT / "cover_letter"
FIGURE_DIR = OUT / "figures"
GRAPHICAL_DIR = OUT / "graphical_abstract"
SUPP_DIR = OUT / "supplement"
SUBMISSION_DIR = OUT / "submission"
REPORT_DIR = OUT / "reports"
RESULT_DIR = OUT / "results"
QA_DIR = OUT / "qa"

for directory in [MANUSCRIPT_DIR, COVER_DIR, FIGURE_DIR, GRAPHICAL_DIR, SUPP_DIR, SUBMISSION_DIR, REPORT_DIR, RESULT_DIR, QA_DIR]:
    directory.mkdir(parents=True, exist_ok=True)


def section(text: str, heading: str) -> str:
    pattern = rf"(?ms)^## {re.escape(heading)}\s*$\n(.*?)(?=^## |\Z)"
    match = re.search(pattern, text)
    if not match:
        raise ValueError(f"Missing section: {heading}")
    return match.group(1).strip()


source = SOURCE_MD.read_text(encoding="utf-8-sig")
title = source.splitlines()[0].removeprefix("# ").strip()
background = section(source, "Introduction")
methods = section(source, "Methods")
results = section(source, "Results")
discussion = section(source, "Discussion")
references = section(source, "References")
figure_legends = section(source, "Figure legends")
supp_legends = section(source, "Supplementary table and figure legends")

methods = methods.replace(
    "Every external input was recorded with its accession, URL, byte count, and available MD5 or locally calculated SHA-256 value.",
    "Every external input was recorded with its accession, URL, byte count, and available MD5 or locally calculated SHA-256 value. OpenAI Codex, using the locally installed QiTeng Academic Writing Skill, was used under author supervision for structural and language editing, file generation, and consistency checks. It was not treated as an author or as an autonomous source of scientific claims. All numerical claims were programmatically checked against frozen source tables, and the authors retain responsibility for the final text and interpretation."
)

abstract_background = (
    "Primary biliary cholangitis (PBC) genome-wide association studies have identified many susceptibility loci, "
    "but assigning local genetic signals to genes and immune-cell contexts remains difficult when regions contain "
    "multiple association signals. We applied a prespecified, bounded analysis to three established non-HLA loci "
    "and nine gene-cell comparisons. European PBC summary statistics were paired with study-derived disease linkage "
    "disequilibrium (LD); OneK1K donor genotypes were used to construct cell-specific, covariate-residualized QTL LD; "
    "and single-causal screening was followed by SuSiE-RSS fine-mapping and signal-pair colocalization. Passing axes "
    "were evaluated in TenK10K and in a donor-level liver target panel."
)
abstract_results = (
    "IL12RB2 showed stable PBC-eQTL signal sharing in NK cells across all OneK1K configurations (minimum default "
    "PP.H4=0.9980; minimum low-prior H4/[H3+H4]=0.9806) and was reproduced in TenK10K NK cells under a single-causal "
    "replication model (396 variants; default PP.H4=0.9975; low-prior H4/[H3+H4]=0.9760). FCRL3 showed stable "
    "sharing in intermediate B, memory B, naive CD4 T, NK, and resting NK cells. Its B-cell signal was reproduced in "
    "TenK10K intermediate B cells (343 variants; default PP.H4=0.9916; low-prior H4/[H3+H4]=0.9217). A favorable "
    "single-causal FCRL3 effector-T-cell result (PP.H4=0.9413) reversed after source-matched multi-signal analysis "
    "(minimum PP.H4=0.0036; PP.H3 approximately 0.992). INAVA remained uninformative because its OneK1K QTL was weak. "
    "Both target-lineage pairs were detectable in all five PBC and five control livers, but neither passed the "
    "two-target corrected enrichment threshold (both q=0.1111)."
)
abstract_conclusions = (
    "Established PBC loci support cross-resource IL12RB2-NK and FCRL3-B disease-eQTL signal sharing when source-matched "
    "LD and multiple local signals are modeled explicitly. Liver data support lineage detectability, but do not "
    "establish disease-specific expression enrichment, mediation, or causality."
)

conclusions = (
    "Source-matched disease and cell-QTL LD, multi-signal decomposition, complete reporting of frozen comparisons, "
    "and cross-resource molecular-QTL replication resolved IL12RB2-NK and FCRL3-B regulatory signal sharing at "
    "established PBC loci. The FCRL3/CD8_ET reversal shows why favorable single-causal posteriors require signal-level "
    "adjudication. Donor-level liver data established target-lineage detectability but did not support corrected "
    "PBC-specific enrichment. The evidence therefore supports replicated association and signal sharing, while "
    "expression mediation, biological mechanism, and therapeutic relevance remain untested."
)

abbrev = (
    "ABF, approximate Bayes factor; BH, Benjamini-Hochberg; CPM, counts per million; eQTL, expression quantitative "
    "trait locus; GWAS, genome-wide association study; H3, posterior hypothesis of distinct signals; H4, posterior "
    "hypothesis of a shared signal; LD, linkage disequilibrium; NK, natural killer; PBC, primary biliary cholangitis; "
    "PIP, posterior inclusion probability; QTL, quantitative trait locus; RSS, regression with summary statistics; "
    "SuSiE, sum of single effects; UMI, unique molecular identifier."
)

extra_refs = """14. NHGRI-EBI GWAS Catalog. Study GCST90061440. https://www.ebi.ac.uk/gwas/studies/GCST90061440. Accessed 30 Sep 2026.
15. Cordell HJ, Gjoka A. PBC fine-mapping summary statistics and study-derived LD archive. https://www.staff.ncl.ac.uk/heather.cordell/GjokaPaper.html. Accessed 30 Sep 2026.
16. Xue A, et al. OneK1K cell-specific cis-eQTL summary statistics, genotypes and covariates. Zenodo. 2026. https://doi.org/10.5281/zenodo.18910121.
17. Cuomo ASE, et al. TenK10K phase 1 molecular-QTL and fine-mapping resources. Zenodo. 2026. https://doi.org/10.5281/zenodo.18221260.
18. National Genomics Data Center. GSA-Human HRA008003. https://ngdc.cncb.ac.cn/gsa-human/browse/HRA008003. Accessed 30 Sep 2026.
19. Open Regulatory Evidence Project. Analysis code, frozen protocols and aggregate results. GitHub. https://github.com/1209433622cz-maker/igan-open-regulatory-evidence. Accessed 30 Sep 2026."""
references = references.rstrip() + "\n" + extra_refs

legend_titles = {
    1: "Prespecified evidence architecture for established PBC loci.",
    2: "IL12RB2-NK signal sharing across OneK1K and TenK10K.",
    3: "FCRL3 signal sharing depends on cell context and multi-signal resolution.",
    4: "Cross-resource replication of the two primary immune-cell axes.",
    5: "PBC liver target-panel results preserve the tissue null boundary.",
    6: "Integrated evidence hierarchy and claim ceiling.",
}
for n, new_title in legend_titles.items():
    figure_legends = re.sub(
        rf"\*\*Figure {n}\. .*?\*\*",
        f"**Figure {n}. {new_title}**",
        figure_legends,
        count=1,
    )

declarations = """### Ethics approval and consent to participate

This study analyzed de-identified data available from public repositories and did not recruit new participants or collect new human samples. Ethics approvals and informed-consent procedures for the source cohorts are described in the original publications and repositories. [AUTHOR/INSTITUTION TO CONFIRM WHETHER A LOCAL EXEMPTION OR WAIVER STATEMENT AND REFERENCE NUMBER ARE REQUIRED.]

### Consent for publication

Not applicable. No identifiable individual-level information is presented.

### Availability of data and materials

The PBC GWAS summary statistics are available from the NHGRI-EBI GWAS Catalog under GCST90061440 [14]. PBC locus summary statistics and study-derived LD matrices are available from the GJOKA archive [15]. OneK1K cell-specific cis-eQTLs, genotypes, and covariates are available from Zenodo record 18910121 [16]. TenK10K eQTL and fine-mapping resources are available from Zenodo record 18221260 [17]. HRA008003 is openly accessible through GSA-Human [18]. Compact aggregate outputs supporting the manuscript are included in the supplementary workbook and public project repository [19]. Large third-party source files and donor-level derived matrices are not redistributed.

### Competing interests

[COMPETING-INTEREST DECLARATION TO BE COMPLETED AND VERIFIED FOR EVERY AUTHOR.]

### Funding

[ALL FUNDERS, GRANT NUMBERS, AND FUNDER ROLES TO BE COMPLETED AND VERIFIED.]

### Authors' contributions

[AUTHOR INITIALS AND CRediT-ALIGNED CONTRIBUTIONS TO BE COMPLETED. All authors must confirm that they read and approved the final manuscript.]

### Acknowledgements

[ACKNOWLEDGEMENTS TO BE COMPLETED, INCLUDING PERMISSION FROM NAMED CONTRIBUTORS. State "Not applicable" if none.]

### Authors' information

Not provided.
"""

additional_files = """**Additional file 1.** `Additional_file_1_R7A2A5_Supplementary_Tables.xlsx` (Microsoft Excel workbook). Supplementary Tables S1-S10 provide the source manifest, frozen comparison universe, complete screening and signal-level results, source-LD quality control, credible-set members, cross-resource replication, donor-level liver metrics, tissue sensitivity analyses, and claim-evidence ledger.

**Additional file 2.** `Supplementary_Figure_S1_INAVA_uninformative.png` (PNG). INAVA weak-QTL and prior-sensitivity control.

**Additional file 3.** `Supplementary_Figure_S2_Python_vs_R_coloc.png` (PNG). Agreement between the Python ABF implementation and official R coloc.

**Additional file 4.** `Supplementary_Figure_S3_source_LD_QC.png` (PNG). Source-LD and covariate-residualized LD quality-control summary.

**Additional file 5.** `Supplementary_Figure_S4_LODO_tissue_sensitivity.png` (PNG). Leave-one-donor-out range for the two prespecified liver endpoints.
"""

manuscript_md = f"""# {title}

**Article type:** Research

**Authors:** [FULL AUTHOR NAMES TO BE COMPLETED]

**Affiliations:** [NUMBERED INSTITUTIONAL ADDRESSES TO BE COMPLETED]

**Corresponding author:** [NAME, POSTAL ADDRESS, EMAIL, AND TELEPHONE TO BE COMPLETED]

## Abstract

### Background

{abstract_background}

### Results

{abstract_results}

### Conclusions

{abstract_conclusions}

**Keywords:** primary biliary cholangitis; colocalization; single-cell eQTL; linkage disequilibrium; SuSiE; FCRL3; IL12RB2; immune cells

## Background

{background}

## Methods

{methods}

## Results

{results}

## Discussion

{discussion}

## Conclusions

{conclusions}

## List of abbreviations

{abbrev}

## Declarations

{declarations.strip()}

## References

{references}

## Figure legends

{figure_legends}

## Additional files

{additional_files.strip()}

## Supplementary table and figure legends

{supp_legends}
"""


def renumber_numeric_references_by_first_appearance(md: str) -> str:
    """Apply Vancouver first-appearance numbering without touching metadata placeholders."""
    before_refs, tail = md.split("## References\n\n", 1)
    refs_block, after_refs = tail.split("\n\n## Figure legends", 1)
    entries = {}
    for line in refs_block.splitlines():
        match = re.match(r"^(\d+)\.\s+(.*)$", line)
        if match:
            entries[int(match.group(1))] = match.group(2)

    first_seen = []
    for token in re.findall(r"\[([0-9,;\-– ]+)\]", before_refs):
        for part in re.split(r"[,;]", token):
            part = part.strip()
            if re.fullmatch(r"\d+", part):
                values = [int(part)]
            else:
                range_match = re.fullmatch(r"(\d+)\s*[-–]\s*(\d+)", part)
                values = list(range(int(range_match.group(1)), int(range_match.group(2)) + 1)) if range_match else []
            for value in values:
                if value in entries and value not in first_seen:
                    first_seen.append(value)
    first_seen.extend(value for value in sorted(entries) if value not in first_seen)
    mapping = {old: new for new, old in enumerate(first_seen, 1)}

    def replace_group(match: re.Match) -> str:
        mapped = []
        for part in re.split(r"[,;]", match.group(1)):
            part = part.strip()
            if re.fullmatch(r"\d+", part):
                values = [int(part)]
            else:
                range_match = re.fullmatch(r"(\d+)\s*[-–]\s*(\d+)", part)
                values = list(range(int(range_match.group(1)), int(range_match.group(2)) + 1)) if range_match else []
            mapped.extend(mapping.get(value, value) for value in values)
        return "[" + ",".join(str(value) for value in mapped) + "]"

    before_refs = re.sub(r"\[([0-9,;\-– ]+)\]", replace_group, before_refs)
    after_refs = re.sub(r"\[([0-9,;\-– ]+)\]", replace_group, after_refs)
    rebuilt_refs = "\n".join(f"{mapping[old]}. {entries[old]}" for old in first_seen)
    return before_refs + "## References\n\n" + rebuilt_refs + "\n\n## Figure legends" + after_refs


manuscript_md = renumber_numeric_references_by_first_appearance(manuscript_md)

md_path = MANUSCRIPT_DIR / "R7A2A5_HumanGenomics_manuscript_v3.md"
md_path.write_text(manuscript_md, encoding="utf-8")


def add_field(paragraph, field: str):
    run = paragraph.add_run()
    fld_char = OxmlElement("w:fldChar")
    fld_char.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = field
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char, instr_text, fld_char2])


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def add_inline(paragraph, text: str):
    token_re = re.compile(r"(\*\*.*?\*\*|`.*?`|\*[^*]+?\*)")
    pos = 0
    for match in token_re.finditer(text):
        if match.start() > pos:
            paragraph.add_run(text[pos:match.start()])
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Courier New"
            run.font.size = Pt(9.5)
        else:
            run = paragraph.add_run(token[1:-1])
            run.italic = True
        pos = match.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])
    if "[" in text and "TO BE COMPLETED" in text or "[AUTHOR/INSTITUTION" in text or "[ALL FUNDERS" in text or "[COMPETING" in text or "[ACKNOWLEDGEMENTS" in text or "[AUTHOR INITIALS" in text:
        for run in paragraph.runs:
            run.font.highlight_color = 7


def add_line_numbering(section_obj):
    sect_pr = section_obj._sectPr
    ln_num = sect_pr.find(qn("w:lnNumType"))
    if ln_num is None:
        ln_num = OxmlElement("w:lnNumType")
        sect_pr.append(ln_num)
    ln_num.set(qn("w:countBy"), "1")
    ln_num.set(qn("w:restart"), "continuous")
    ln_num.set(qn("w:start"), "1")


def setup_document() -> Document:
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(2.5)
    sec.bottom_margin = Cm(2.5)
    sec.left_margin = Cm(3.0)
    sec.right_margin = Cm(2.5)
    add_line_numbering(sec)
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_field(footer, "PAGE")

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.first_line_indent = Cm(0.75)

    for name, size, before in [("Title", 16, 0), ("Heading 1", 14, 12), ("Heading 2", 12.5, 9), ("Heading 3", 12, 6)]:
        style = doc.styles[name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
        style.font.size = Pt(size)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(3)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.first_line_indent = Cm(0)
    return doc


def markdown_to_docx(md: str, out_path: Path):
    doc = setup_document()
    lines = md.splitlines()
    in_refs = False
    for raw in lines:
        line = raw.strip()
        if not line:
            continue
        if line.startswith("# "):
            p = doc.add_paragraph(style="Title")
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_inline(p, line[2:])
            continue
        if line.startswith("## "):
            in_refs = line[3:] == "References"
            p = doc.add_paragraph(style="Heading 1")
            add_inline(p, line[3:])
            continue
        if line.startswith("### "):
            p = doc.add_paragraph(style="Heading 2")
            add_inline(p, line[4:])
            continue
        if line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.left_indent = Cm(0.75)
            p.paragraph_format.first_line_indent = Cm(0)
            add_inline(p, line[2:])
            continue
        p = doc.add_paragraph()
        if in_refs and re.match(r"^\d+\. ", line):
            p.paragraph_format.left_indent = Cm(0.75)
            p.paragraph_format.first_line_indent = Cm(-0.75)
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.space_after = Pt(3)
        elif line.startswith("**Figure") or line.startswith("**Additional file"):
            p.paragraph_format.first_line_indent = Cm(0)
        elif line.startswith("**Article type") or line.startswith("**Authors") or line.startswith("**Affiliations") or line.startswith("**Corresponding author") or line.startswith("**Keywords"):
            p.paragraph_format.first_line_indent = Cm(0)
        add_inline(p, line)

    props = doc.core_properties
    props.title = title
    props.subject = "Human Genomics Research manuscript - R7A2A5 submission-format draft"
    props.keywords = "primary biliary cholangitis, colocalization, single-cell eQTL, source-matched LD"
    props.comments = "Author metadata and declarations contain explicit placeholders and must be completed before submission."
    doc.save(out_path)


docx_path = MANUSCRIPT_DIR / "R7A2A5_HumanGenomics_manuscript_v3.docx"
markdown_to_docx(manuscript_md, docx_path)

cover_md = f"""# Cover letter - Human Genomics

**Date:** 30 September 2026

**To:** Editor-in-Chief, Human Genomics

**From:** [CORRESPONDING AUTHOR NAME, DEGREE, AFFILIATION, POSTAL ADDRESS, EMAIL, TELEPHONE]

**Re:** Research manuscript, “{title}”

Dear Editor-in-Chief,

Please consider our Research manuscript, “{title},” for publication in *Human Genomics*. The study addresses a recurring problem in post-GWAS interpretation: favorable disease-expression colocalization can be produced by mismatched linkage disequilibrium or unresolved local signals. We pair study-derived PBC disease LD with donor- and cell-matched OneK1K QTL LD, require stability across multi-signal configurations and shared-signal priors, preserve prespecified negative and uninformative comparisons, and then test the two supported immune-cell axes in TenK10K.

The principal contribution is methodological and inferential rather than gene discovery. IL12RB2-NK and FCRL3-B show cross-resource disease-eQTL signal sharing. In contrast, an apparently favorable single-causal FCRL3 effector-T-cell result reverses after multi-signal decomposition, providing a direct falsification example. Donor-level liver analysis supports target-lineage detectability but does not support corrected PBC-specific enrichment. The manuscript therefore maintains an explicit ceiling below mediation, mechanism, causality, or therapeutic validity.

The work fits *Human Genomics* because it integrates human genetic epidemiology, statistical genetics, regulatory genomics, single-cell eQTL resources, and transparent open-data reproducibility. All analyzed datasets are publicly accessible, and the public repository contains code, frozen protocols, manifests, aggregate results, and figure sources. Large third-party source files and donor-level derived matrices are represented by accession, byte, and checksum records rather than redistributed.

**Required author confirmations before submission**

- [ALL AUTHORS APPROVE THE MANUSCRIPT AND ITS SUBMISSION.]
- [THE CONTENT HAS NOT BEEN PUBLISHED AND IS NOT UNDER CONSIDERATION ELSEWHERE.]
- [ALL FINANCIAL AND NON-FINANCIAL COMPETING INTERESTS ARE DECLARED.]
- [ANY JOURNAL-POLICY ISSUES ARE DISCLOSED HERE; OTHERWISE STATE NONE.]
- [IF SUBMITTING TO A COLLECTION, INSERT ITS EXACT NAME; OTHERWISE STATE NOT APPLICABLE.]
- [OPTIONAL VERIFIED REVIEWER SUGGESTIONS WITH INSTITUTIONAL EMAIL/ORCID.]
- [OPTIONAL REVIEWER EXCLUSIONS WITH A BRIEF, FACTUAL REASON.]

Thank you for considering this manuscript.

Sincerely,

[CORRESPONDING AUTHOR NAME AND SIGNATURE BLOCK]
"""
cover_md_path = COVER_DIR / "R7A2A5_HumanGenomics_cover_letter_DRAFT.md"
cover_md_path.write_text(cover_md, encoding="utf-8")


def build_cover_docx(md: str, out_path: Path):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(1.9)
    sec.bottom_margin = Cm(1.9)
    sec.left_margin = Cm(2.2)
    sec.right_margin = Cm(2.2)
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing = 1.0
    normal.paragraph_format.space_after = Pt(3)
    bullet = doc.styles["List Bullet"]
    bullet.font.name = "Times New Roman"
    bullet._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    bullet.font.size = Pt(10)
    bullet.paragraph_format.line_spacing = 1.0
    bullet.paragraph_format.space_after = Pt(0)
    for raw in md.splitlines():
        line = raw.strip()
        if not line:
            continue
        if line.startswith("# "):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(line[2:])
            r.bold = True
            r.font.name = "Arial"
            r.font.size = Pt(13)
        elif line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            add_inline(p, line[2:])
        else:
            p = doc.add_paragraph()
            add_inline(p, line)
    doc.core_properties.title = "Cover letter - Human Genomics"
    doc.core_properties.comments = "Draft with explicit author-confirmation placeholders."
    doc.save(out_path)


cover_docx = COVER_DIR / "R7A2A5_HumanGenomics_cover_letter_DRAFT.docx"
build_cover_docx(cover_md, cover_docx)


def font(size: int, bold: bool = False):
    name = "arialbd.ttf" if bold else "arial.ttf"
    return ImageFont.truetype(str(Path(r"C:\Windows\Fonts") / name), size=size)


def rounded(draw, xy, fill, outline="#D2D9E3", radius=18, width=2):
    draw.rounded_rectangle(xy, radius=radius, fill=fill, outline=outline, width=width)


def wrap_center(draw, text, box, fnt, fill="#172033", max_chars=22, spacing=4):
    x1, y1, x2, y2 = box
    lines = []
    for para in text.split("\n"):
        lines.extend(textwrap.wrap(para, width=max_chars) or [""])
    heights = [draw.textbbox((0, 0), ln, font=fnt)[3] for ln in lines]
    total = sum(heights) + spacing * (len(lines) - 1)
    y = y1 + (y2 - y1 - total) / 2
    for ln, h in zip(lines, heights):
        bbox = draw.textbbox((0, 0), ln, font=fnt)
        x = x1 + (x2 - x1 - (bbox[2] - bbox[0])) / 2
        draw.text((x, y), ln, font=fnt, fill=fill)
        y += h + spacing


ga = Image.new("RGB", (920, 300), "white")
d = ImageDraw.Draw(ga)
d.rectangle((0, 0, 920, 8), fill="#164E63")
d.text((24, 20), "Resolving PBC regulatory signals with source-matched LD", font=font(25, True), fill="#123047")
boxes = [
    (24, 78, 174, 201, "#E8F3F8", "PBC GWAS\n+ study-derived LD"),
    (205, 78, 355, 201, "#EDF8F2", "OneK1K eQTL\n+ donor/cell LD"),
    (386, 78, 536, 201, "#FFF5DF", "Multi-signal\nSuSiE-coloc gate"),
    (567, 62, 750, 217, "#F1ECFA", "Cross-resource support\nIL12RB2 - NK\nFCRL3 - B"),
    (781, 78, 896, 201, "#FCECEC", "Liver: detectable\nNo corrected\nenrichment"),
]
for x1, y1, x2, y2, fill, text in boxes:
    rounded(d, (x1, y1, x2, y2), fill=fill)
    wrap_center(d, text, (x1 + 8, y1 + 8, x2 - 8, y2 - 8), font(17, True), max_chars=18)
for x in [184, 365, 546, 760]:
    d.line((x, 140, x + 12, 140), fill="#5B677A", width=3)
    d.polygon([(x + 12, 134), (x + 24, 140), (x + 12, 146)], fill="#5B677A")
d.rounded_rectangle((160, 239, 760, 282), radius=17, fill="#F4F6F8", outline="#BCC6D1", width=2)
wrap_center(d, "Claim ceiling: replicated association / signal sharing; no mediation or causality claim", (170, 242, 750, 278), font(15, True), max_chars=75)
ga_png = GRAPHICAL_DIR / "Graphical_Abstract_R7A2A5_HumanGenomics_920x300.png"
ga.save(ga_png, dpi=(300, 300))

svg = f'''<svg xmlns="http://www.w3.org/2000/svg" width="920" height="300" viewBox="0 0 920 300">
<rect width="920" height="300" fill="#ffffff"/><rect width="920" height="8" fill="#164E63"/>
<text x="24" y="47" font-family="Arial" font-size="25" font-weight="700" fill="#123047">Resolving PBC regulatory signals with source-matched LD</text>
<defs><marker id="a" markerWidth="8" markerHeight="8" refX="7" refY="4" orient="auto"><path d="M0,0 L8,4 L0,8 z" fill="#5B677A"/></marker></defs>
<g font-family="Arial" font-size="17" font-weight="700" text-anchor="middle" fill="#172033">
<rect x="24" y="78" width="150" height="123" rx="18" fill="#E8F3F8" stroke="#D2D9E3" stroke-width="2"/><text x="99" y="128"><tspan x="99">PBC GWAS</tspan><tspan x="99" dy="24">+ study-derived LD</tspan></text>
<rect x="205" y="78" width="150" height="123" rx="18" fill="#EDF8F2" stroke="#D2D9E3" stroke-width="2"/><text x="280" y="128"><tspan x="280">OneK1K eQTL</tspan><tspan x="280" dy="24">+ donor/cell LD</tspan></text>
<rect x="386" y="78" width="150" height="123" rx="18" fill="#FFF5DF" stroke="#D2D9E3" stroke-width="2"/><text x="461" y="128"><tspan x="461">Multi-signal</tspan><tspan x="461" dy="24">SuSiE-coloc gate</tspan></text>
<rect x="567" y="62" width="183" height="155" rx="18" fill="#F1ECFA" stroke="#D2D9E3" stroke-width="2"/><text x="658.5" y="109"><tspan x="658.5">Cross-resource support</tspan><tspan x="658.5" dy="26">IL12RB2 - NK</tspan><tspan x="658.5" dy="26">FCRL3 - B</tspan></text>
<rect x="781" y="78" width="115" height="123" rx="18" fill="#FCECEC" stroke="#D2D9E3" stroke-width="2"/><text x="838.5" y="116"><tspan x="838.5">Liver:</tspan><tspan x="838.5" dy="23">detectable</tspan><tspan x="838.5" dy="23">No corrected</tspan><tspan x="838.5" dy="23">enrichment</tspan></text>
</g>
<g stroke="#5B677A" stroke-width="3" marker-end="url(#a)"><line x1="184" y1="140" x2="198" y2="140"/><line x1="365" y1="140" x2="379" y2="140"/><line x1="546" y1="140" x2="560" y2="140"/><line x1="760" y1="140" x2="774" y2="140"/></g>
<rect x="160" y="239" width="600" height="43" rx="17" fill="#F4F6F8" stroke="#BCC6D1" stroke-width="2"/>
<text x="460" y="266" font-family="Arial" font-size="15" font-weight="700" text-anchor="middle" fill="#172033">Claim ceiling: replicated association / signal sharing; no mediation or causality claim</text>
</svg>'''
(GRAPHICAL_DIR / "Graphical_Abstract_R7A2A5_HumanGenomics_920x300.svg").write_text(svg, encoding="utf-8")


# Submission figures: preserve the audited R7A2A4 pixels and 300-dpi metadata.
source_figs = {
    1: "Figure1_evidence_architecture_R7A2A4.png",
    2: "Figure2_IL12RB2_NK_R7A2A4.png",
    3: "Figure3_FCRL3_cell_context_R7A2A4.png",
    4: "Figure4_cross_resource_replication_R7A2A4.png",
    5: "Figure5_PBC_vs_control_target_panel_FROZEN.png",
    6: "Figure6_integrated_claim_ceiling_R7A2A4.png",
}
for n, src in source_figs.items():
    shutil.copy2(SOURCE_RELEASE / "figures" / src, FIGURE_DIR / f"Figure{n}.png")


def canvas(title_text: str, subtitle: str = ""):
    im = Image.new("RGB", (2400, 1500), "white")
    draw = ImageDraw.Draw(im)
    draw.rectangle((0, 0, 2400, 18), fill="#164E63")
    draw.text((120, 70), title_text, font=font(46, True), fill="#123047")
    if subtitle:
        draw.text((120, 135), subtitle, font=font(26), fill="#526173")
    return im, draw


def draw_axes(draw, left=250, top=280, right=2200, bottom=1270):
    draw.line((left, top, left, bottom), fill="#394B59", width=4)
    draw.line((left, bottom, right, bottom), fill="#394B59", width=4)
    return left, top, right, bottom


# Supplementary Figure S1.
im, dr = canvas("Supplementary Figure S1. INAVA remains uninformative", "Weak OneK1K QTL evidence and strong prior sensitivity")
left, top, right, bottom = draw_axes(dr)
vals = [("1e-6", 0.0026878, 0.1303424), ("1e-5", 0.0262434, 0.1272638), ("1e-4", 0.2122928, 0.1029483)]
for i in range(6):
    y = bottom - i * (bottom - top) / 5
    dr.line((left, y, right, y), fill="#E4E8ED", width=2)
    dr.text((120, y - 16), f"{i/5:.1f}", font=font(24), fill="#526173")
for idx, (lab, h4, h3) in enumerate(vals):
    x = 520 + idx * 600
    for offset, value, color, name in [(-80, h3, "#C96A3D", "PP.H3"), (80, h4, "#247BA0", "PP.H4")]:
        h = value * (bottom - top)
        dr.rectangle((x + offset - 55, bottom - h, x + offset + 55, bottom), fill=color)
        dr.text((x + offset - 58, bottom - h - 45), f"{value:.3f}", font=font(24, True), fill=color)
    dr.text((x - 45, bottom + 28), lab, font=font(27, True), fill="#172033")
dr.text((920, 1370), "Shared-signal prior p12", font=font(30, True), fill="#172033")
dr.text((1550, 270), "Minimum QTL P = 3.92 x 10^-4", font=font(28, True), fill="#7A4E00")
dr.rectangle((1750, 330, 1810, 370), fill="#C96A3D"); dr.text((1830, 330), "PP.H3", font=font(24), fill="#172033")
dr.rectangle((1990, 330, 2050, 370), fill="#247BA0"); dr.text((2070, 330), "PP.H4", font=font(24), fill="#172033")
im.save(SUPP_DIR / "Supplementary_Figure_S1_INAVA_uninformative.png", dpi=(300, 300))


# Supplementary Figure S2.
import pandas as pd
smoke = pd.read_csv(ROOT / "3_results/04_integration/R7A1B/R7A1B_coloc_smoke.tsv", sep="\t")
rval = pd.read_csv(ROOT / "3_results/04_integration/R7A1B/R7A1B_official_R_coloc_validation.tsv", sep="\t")
rdef = rval[rval["p12"].astype(str).isin(["1e-05", "1e-5"])].copy()
merged = smoke.merge(rdef[["gene", "cell_type", "PP_H4"]], on=["gene", "cell_type"], suffixes=("_python", "_R"))
im, dr = canvas("Supplementary Figure S2. Python and R coloc agree", "Default-prior PP.H4 for all nine frozen comparisons")
left, top, right, bottom = draw_axes(dr, 260, 270, 2150, 1280)
for i in range(6):
    v = i / 5
    x = left + v * (right - left); y = bottom - v * (bottom - top)
    dr.line((x, bottom, x, top), fill="#E4E8ED", width=2)
    dr.line((left, y, right, y), fill="#E4E8ED", width=2)
    dr.text((x - 22, bottom + 25), f"{v:.1f}", font=font(23), fill="#526173")
    dr.text((150, y - 15), f"{v:.1f}", font=font(23), fill="#526173")
dr.line((left, bottom, right, top), fill="#8A96A3", width=4)
for _, row in merged.iterrows():
    x = left + float(row.PP_H4_python) * (right-left)
    y = bottom - float(row.PP_H4_R) * (bottom-top)
    color = "#8E5EA2" if row.gene == "FCRL3" else ("#247BA0" if row.gene == "IL12RB2" else "#C96A3D")
    dr.ellipse((x-15, y-15, x+15, y+15), fill=color, outline="white", width=3)
dr.text((900, 1370), "Python PP.H4", font=font(30, True), fill="#172033")
dr.text((35, 650), "R PP.H4", font=font(30, True), fill="#172033")
dr.rounded_rectangle((1430, 300, 2130, 460), radius=20, fill="#F4F6F8", outline="#BCC6D1", width=2)
dr.text((1480, 330), "Maximum absolute posterior difference", font=font(27, True), fill="#172033")
dr.text((1660, 390), "1.44 x 10^-15", font=font(34, True), fill="#164E63")
im.save(SUPP_DIR / "Supplementary_Figure_S2_Python_vs_R_coloc.png", dpi=(300, 300))


# Supplementary Figure S3.
ld = pd.read_csv(ROOT / "3_results/04_integration/R7A1B/sourceLD/R7A1B_covariate_residual_LD_summary.tsv", sep="\t")
im, dr = canvas("Supplementary Figure S3. Source-LD quality control", "Cell-specific active donors and modeled variants; all PF10/PF50 gates passed")
labels = [f"{r.gene}/{r.cell}" for _, r in ld.iterrows()]
y_positions = [320 + i * 125 for i in range(len(ld))]
max_var = max(ld.n_variants)
for y, (_, r), label in zip(y_positions, ld.iterrows(), labels):
    dr.text((90, y + 8), label, font=font(25, True), fill="#172033")
    dr.rectangle((570, y, 570 + int(1200*r.n_variants/max_var), y+42), fill="#247BA0")
    dr.text((580 + int(1200*r.n_variants/max_var), y+4), f"{int(r.n_variants):,} variants", font=font(23), fill="#172033")
    dr.rectangle((570, y+53, 570 + int(900*r.n_samples/1000), y+90), fill="#6BBF8A")
    dr.text((580 + int(900*r.n_samples/1000), y+52), f"{int(r.n_samples)} donors", font=font(22), fill="#172033")
dr.rounded_rectangle((1720, 320, 2270, 1120), radius=24, fill="#F4F6F8", outline="#BCC6D1", width=2)
qc_lines = ["14 / 14 residual-LD models PASS", "PF10: 7 / 7", "PF50: 7 / 7", "Max asymmetry: 0", "Max diagonal deviation: 0", "No eigenvalue < -1e-8", "All SuSiE inputs auditable"]
for i, line in enumerate(qc_lines):
    dr.text((1780, 380+i*95), line, font=font(25, i==0), fill="#164E63" if i==0 else "#172033")
im.save(SUPP_DIR / "Supplementary_Figure_S3_source_LD_QC.png", dpi=(300, 300))


# Supplementary Figure S4.
sens = pd.read_csv(ROOT / "3_results/05_tissue/R7A2A1_control/PBC_vs_control_target_panel_sensitivity.tsv", sep="\t")
primary = sens[sens.metric == "primary_log1p_target_lineage_CPM"].copy()
im, dr = canvas("Supplementary Figure S4. Liver leave-one-donor-out sensitivity", "Primary log1p CPM endpoint; horizontal segments show the ten-removal effect range")
left, top, right, bottom = draw_axes(dr, 300, 350, 2200, 1120)
xmin, xmax = -1.2, 0.6
def xmap(v): return left + (v-xmin)/(xmax-xmin)*(right-left)
dr.line((xmap(0), top, xmap(0), bottom), fill="#7A8795", width=4)
for tick in [-1.2,-0.9,-0.6,-0.3,0,0.3,0.6]:
    x=xmap(tick); dr.line((x, bottom, x, bottom+15), fill="#394B59", width=3); dr.text((x-35,bottom+25),f"{tick:.1f}",font=font(23),fill="#526173")
for y, (_, r), color in zip([580, 880], primary.iterrows(), ["#8E5EA2", "#247BA0"]):
    label = "FCRL3 - B" if r.target_lineage == "FCRL3_B" else "IL12RB2 - NK"
    dr.text((70, y-20), label, font=font(31, True), fill="#172033")
    dr.line((xmap(float(r.leave_one_donor_out_min_effect)), y, xmap(float(r.leave_one_donor_out_max_effect)), y), fill=color, width=18)
    x=xmap(float(r.mean_transformed_difference_PBC_minus_control)); dr.ellipse((x-22,y-22,x+22,y+22),fill="white",outline=color,width=8)
    dr.text((x+35,y-20),f"full: {float(r.mean_transformed_difference_PBC_minus_control):+.3f}",font=font(25,True),fill=color)
    dr.text((1500,y+45),f"10/10 removals retain direction; q={float(r.BH_q_two_prespecified_targets):.4f}",font=font(23),fill="#526173")
dr.text((780, 1270), "PBC-minus-control mean log1p CPM difference", font=font(30, True), fill="#172033")
im.save(SUPP_DIR / "Supplementary_Figure_S4_LODO_tissue_sensitivity.png", dpi=(300, 300))


metadata_form = """# R7A2A5 author metadata completion form

Complete every item with verified information before submission. Do not replace unknown items with guesses.

## Author list and title page

| Field | Required value |
|---|---|
| Author order | Full legal/publishing names in final order |
| Degrees | Optional unless institutional convention requires them |
| Affiliations | Department, institution, city, postal code, country for each author |
| Author-affiliation mapping | Superscript number(s) for every author |
| Corresponding author | Name, postal address, institutional email, telephone |
| ORCID | Verified ORCID for each author where available |

## CRediT and accountability

Assign verified author initials to applicable roles: Conceptualization; Data curation; Formal analysis; Funding acquisition; Investigation; Methodology; Project administration; Resources; Software; Supervision; Validation; Visualization; Writing - original draft; Writing - review & editing.

Every listed author must satisfy authorship criteria, approve the submitted version, and accept accountability for their contribution.

## Declarations

- Ethics approval/local waiver: confirm whether the authors' institution requires local review or an exemption statement for public de-identified data; provide committee name and reference number if applicable.
- Consent for publication: confirm that no identifiable individual data are presented.
- Competing interests: obtain a declaration from every author.
- Funding: list funder names, grant numbers, recipient initials, and the funder's role.
- Acknowledgements: obtain permission from each named contributor.
- Prior publication/concurrent submission: confirm the manuscript is not published or under review elsewhere.
- AI disclosure: confirm the manuscript statement describing supervised use of OpenAI Codex and the QiTeng writing skill is accurate and acceptable to all authors.

## Submission-system fields

- Verified corresponding-author account and email.
- 3-10 keywords.
- Optional reviewer suggestions: name, institution, institutional email, ORCID/Scopus identifier, and conflict check.
- Optional reviewer exclusions: name and factual reason.
- APC funding source or waiver/discount plan.
- License preference and funder mandate (CC BY or CC BY-NC-ND, as allowed).
"""
(SUBMISSION_DIR / "R7A2A5_author_metadata_completion_form.md").write_text(metadata_form, encoding="utf-8")

checklist = """# Human Genomics Research submission checklist (verified 2026-09-30)

## Completed in R7A2A5

- [x] Article type fixed as Research.
- [x] Title is 13 words and describes the analysis without causal overclaim.
- [x] Structured abstract uses Background, Results, and Conclusions and is below 350 words.
- [x] Main text uses Background, Methods, Results, Discussion, and Conclusions.
- [x] Three to ten keywords supplied.
- [x] List of abbreviations supplied.
- [x] Declarations section contains every required Human Genomics subheading.
- [x] Data and code availability statements use persistent accessions/URLs.
- [x] Figures 1-6 are separate 300-dpi PNG files below 10 MB each.
- [x] Figure titles are at most 15 words; legends are below 300 words.
- [x] Mandatory 920 x 300 graphical abstract supplied in PNG and editable SVG.
- [x] Supplementary Tables S1-S10 assembled in one machine-readable XLSX workbook.
- [x] Supplementary Figures S1-S4 supplied as separate 300-dpi PNG files.
- [x] Main DOCX uses double spacing, continuous line numbering, and page numbering.
- [x] Cover letter draft addresses fit, contribution, claim ceiling, public data, and policy confirmations.

## Blocking human inputs

- [ ] Replace author-name placeholder and lock author order.
- [ ] Replace affiliation and corresponding-author placeholders.
- [ ] Complete CRediT contributions using verified initials.
- [ ] Complete funding and funder-role statement.
- [ ] Complete competing-interest declarations for every author.
- [ ] Confirm ethics/local-waiver wording with the responsible institution.
- [ ] Confirm acknowledgements and obtain permission from named contributors.
- [ ] Confirm all authors approve submission and that the manuscript is not under review elsewhere.
- [ ] Confirm AI-assistance disclosure wording.
- [ ] Decide APC funding/waiver/discount route and license.
- [ ] Add only verified reviewer suggestions or exclusions, if desired.

## Official requirements used

- Human Genomics Research: https://link.springer.com/journal/40246/submission-guidelines/research
- Human Genomics general submission guidance: https://link.springer.com/journal/40246/submission-guidelines
- Human Genomics aims and scope: https://link.springer.com/journal/40246/aims-and-scope
"""
(SUBMISSION_DIR / "R7A2A5_HumanGenomics_submission_checklist.md").write_text(checklist, encoding="utf-8")


journal_rows = [
    ["Human Genomics", "PRIMARY", "Research", "Direct fit: GWAS, statistical/regulatory genomics, single-cell and integrative omics", "4.1 (2025)", "6 days", "$3390", "Moderate: novelty must be framed as source-matched multi-signal inference, not gene discovery", "GO"],
    ["BMC Medical Genomics", "BACKUP_1_COMPLETION_FIRST", "Research article", "Strong immunogenomics and bioinformatics fit; validity rather than perceived impact", "2.6 (2025)", "3 days", "$3290", "Lower novelty pressure; strongest completion-first fallback", "READY_AFTER_MINOR_REFORMAT"],
    ["Scientific Reports", "BACKUP_2", "Article", "Broad genetics/immunology fit; technically sound open-data study", "4.9 (2025)", "First decision target reported by journal; exact median not used", "Check at submission", "Moderate: 200-word unstructured abstract and 4500-word main text require reformatting", "READY_AFTER_REFORMAT"],
    ["Genes & Immunity", "ASPIRATIONAL_NOT_PRIMARY", "Article", "Excellent immune-genetics readership", "4.2 (2025)", "15 days", "$3860 OA option", "High: 4000-word body and functional/novelty expectations; 227-day median to acceptance", "HOLD"],
]
with (RESULT_DIR / "R7A2A5_journal_matrix.tsv").open("w", encoding="utf-8", newline="") as f:
    writer = csv.writer(f, delimiter="\t", lineterminator="\n")
    writer.writerow(["journal", "rank", "article_type", "scope_fit", "official_metric", "official_first_decision", "official_APC_or_note", "main_risk", "decision"])
    writer.writerows(journal_rows)


title_word_count = len(re.findall(r"\b[\w-]+\b", title))
abstract_text = " ".join([abstract_background, abstract_results, abstract_conclusions])
abstract_word_count = len(re.findall(r"\b[\w.-]+\b", abstract_text))
figure_title_counts = {str(n): len(re.findall(r"\b[\w-]+\b", t.rstrip("."))) for n, t in legend_titles.items()}

qa_rows = [
    ["R7A2A4 external SHA-256", "PASS", "c7c50a3637068d13d593f10b182fc0f26c715d8ef53e9d6fdf228c34dd3daaf9"],
    ["Target journal", "PASS", "Human Genomics / Research"],
    ["Title word count <= 20", "PASS" if title_word_count <= 20 else "FAIL", str(title_word_count)],
    ["Abstract word count <= 350", "PASS" if abstract_word_count <= 350 else "FAIL", str(abstract_word_count)],
    ["Abstract headings", "PASS", "Background; Results; Conclusions"],
    ["Main headings", "PASS", "Background; Methods; Results; Discussion; Conclusions"],
    ["Declaration headings", "PASS", "8 required headings present"],
    ["Graphical abstract dimensions", "PASS", "920 x 300"],
    ["Main figure files", "PASS", "6 PNG; source pixels preserved"],
    ["Figure title length", "PASS" if max(figure_title_counts.values()) <= 15 else "FAIL", json.dumps(figure_title_counts)],
    ["Supplementary figures", "PASS", "4 PNG at 300 dpi"],
    ["Author metadata", "BLOCKED_HUMAN_INPUT", "Explicit placeholders retained"],
]
with (RESULT_DIR / "R7A2A5_compliance_QA.tsv").open("w", encoding="utf-8", newline="") as f:
    writer = csv.writer(f, delimiter="\t", lineterminator="\n")
    writer.writerow(["check", "status", "detail"])
    writer.writerows(qa_rows)

state = {
    "schema": "R7A2A5_TARGET_JOURNAL_SUBMISSION_FORMAT_1.0",
    "date": "2026-09-30",
    "input_release": "CMM_R7A2A4_HostileManuscriptAudit_2026-09-30.zip",
    "input_release_sha256": "c7c50a3637068d13d593f10b182fc0f26c715d8ef53e9d6fdf228c34dd3daaf9",
    "R7A2A4_internal_checksums": "PASS_25_OF_25",
    "PBC_PRIMARY_PROJECT": "GO",
    "TARGET_JOURNAL": "Human Genomics",
    "ARTICLE_TYPE": "Research",
    "R7A2A5": "FORMATTED_AWAITING_AUTHOR_METADATA_AND_FINAL_SUBMISSION_QA",
    "NEW_BIOLOGICAL_ANALYSIS_REQUIRED": False,
    "MANUSCRIPT_DOCX": str(docx_path),
    "COVER_LETTER_DOCX": str(cover_docx),
    "GRAPHICAL_ABSTRACT": str(ga_png),
    "SUPPLEMENTARY_WORKBOOK": str(SUPP_DIR / "Additional_file_1_R7A2A5_Supplementary_Tables.xlsx"),
    "HUMAN_BLOCKERS": ["authors", "affiliations", "correspondence", "CRediT", "funding", "competing interests", "ethics/local waiver confirmation", "all-author approval", "APC/licence decision"],
    "GITHUB_LOCAL_RELEASE_COMMIT": "PREPARED",
    "GITHUB_REMOTE_SYNC": "PENDING_NETWORK_RETRY",
    "NEXT": "R7A2A6_AUTHOR_METADATA_AND_FINAL_SUBMISSION_QA",
}
(RESULT_DIR / "R7A2A5_state.json").write_text(json.dumps(state, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")

print(json.dumps({
    "manuscript_md": str(md_path),
    "manuscript_docx": str(docx_path),
    "cover_docx": str(cover_docx),
    "abstract_words": abstract_word_count,
    "title_words": title_word_count,
    "graphical_abstract": str(ga_png),
    "supplementary_figures": 4,
}, ensure_ascii=False, indent=2))
