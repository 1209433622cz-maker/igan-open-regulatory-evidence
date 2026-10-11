#!/usr/bin/env python3
"""Build the R7C1 Human Genomics manuscript fork from the frozen R7B4B2 baseline."""

from __future__ import annotations

import hashlib
import json
import re
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(r"H:\SCI2\YR1")
BASELINE = ROOT / "5_manuscript/R7B4B_FinalSubmissionCandidate/manuscript/R7B4B_HumanGenomics_final.md"
WORK = ROOT / "5_manuscript/R7C1_IL12RB2_ExternalValidation_ManuscriptFork"
OUT_MD = WORK / "manuscript/R7C1_HumanGenomics_IL12RB2_external_validation_fork.md"
OUT_DOCX = WORK / "manuscript/R7C1_HumanGenomics_IL12RB2_external_validation_fork.docx"
TITLE = "Input-matched multi-signal colocalization clarifies immune-cell regulatory assignments at primary biliary cholangitis risk loci"


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


def replace_section(text: str, start: str, end: str, replacement: str) -> str:
    a = text.find(start)
    if a < 0:
        raise RuntimeError(f"section start missing: {start}")
    b = text.find(end, a + len(start))
    if b < 0:
        raise RuntimeError(f"section end missing: {end}")
    return text[:a] + replacement.rstrip() + "\n\n" + text[b:]


def build_markdown() -> str:
    text = BASELINE.read_text(encoding="utf-8")

    old = (
        "Support restriction and disease-statistic matching changed 15 and 19 categorical assignments, respectively. "
        "With matched inputs, 79 of 113 approximate-Bayes-factor H4 comparisons retained shared-pair support, whereas eight supported distinct pairs; "
        "ten of 452 H3 comparisons acquired shared-pair support. Final PF10 states comprised 92 H4-supported, 428 H3-supported, 59 model-sensitive and 63 uninformative comparisons. "
        "Thirty-two H4-supported comparisons also contained a qualifying distinct pair. Across the frozen simulation grid, mean false-H4 decisions in distinct-signal scenarios decreased from 12.41% to 8.56%, "
        "while shared-signal recovery decreased from 32.99% to 26.66%; only 160,782 iterations (33.08%) formed an evaluable signal pair. IL12RB2-NK and FCRL3 B-cell assignments received bounded cross-resource support. "
        "Both target-lineage pairs were detected in all five PBC and five control livers without corrected PBC-specific enrichment (both q=0.111)."
    )
    new = (
        "Support restriction and disease-statistic matching changed 15 and 19 categorical assignments, respectively. "
        "With matched inputs, 79 of 113 approximate-Bayes-factor H4 comparisons retained shared-pair support, whereas eight supported distinct pairs; ten of 452 H3 comparisons acquired shared-pair support. "
        "Final PF10 states comprised 92 H4-supported, 428 H3-supported, 59 model-sensitive and 63 uninformative comparisons. Across the frozen simulation grid, mean false-H4 decisions in distinct-signal scenarios decreased from 12.41% to 8.56%, while shared-signal recovery decreased from 32.99% to 26.66%. "
        "In a different disease study system, FinnGen R13 × OneK1K NK/IL12RB2 analysis used 1,175 variants and gave PP.H4=0.9977 at the default prior and 0.9772 at p12=10^-6; both OneK QTL credible-set variants occurred in the R13 disease credible set. "
        "A linked NK-cell peak supported PBC-caQTL sharing (PP.H4=0.9695) and a promoter-proximal peak-IL12RB2 link, while the direct eQTL-caQTL posterior remained unavailable. Both target-lineage pairs were detected in all five PBC and five control livers without corrected PBC-specific enrichment (both q=0.111)."
    )
    if old not in text:
        raise RuntimeError("abstract Results anchor changed")
    text = text.replace(old, new)
    text = text.replace(
        "Input matching explained part, but not all, of the differences in PBC regulatory assignments. Shared and distinct signal pairs, covariate-model sensitivity and discovery limitations should be reported separately. The results do not establish causal truth, general method superiority or tissue mediation.",
        "Input matching explained part, but not all, of the differences in PBC regulatory assignments. Different-study-system validation strengthened the IL12RB2-NK association, but did not establish person-level independence, chromatin mediation or causal mechanism. Shared and distinct signal pairs, covariate-model sensitivity and discovery limitations should be reported separately.",
    )
    text = text.replace(
        "Finally, we examined representative IL12RB2-NK and FCRL3 B-cell assignments across independent molecular-QTL resources and bounded tissue evidence.",
        "Finally, we examined representative IL12RB2-NK and FCRL3 B-cell assignments across molecular-QTL resources, evaluated IL12RB2-NK in a different disease study system and separated chromatin signal coherence from direct mediation.",
    )

    methods = """### External molecular-QTL evaluation

External evaluation was bounded to the representative IL12RB2-NK and FCRL3 B-cell axes; it was not used to select new genes. TenK10K common-variant cis-eQTL statistics within +/-100 kb were intersected with the same PBC disease summaries [4,14]. ABF colocalization used the frozen prior series, and source TenK10K credible-set membership was reviewed where available. The PBC GWAS was shared across the OneK1K and TenK10K analyses, and donor-level TenK10K LD was not reconstructed. These results were therefore termed cross-resource molecular-QTL replication, not independent disease replication or a second source-matched multi-signal proof.

FinnGen R12 PBC and molecular-QTL colocalization records were queried through the public outputs of the population-scale immune multiome resource [15] for the two representative genes. Returned IL12RB2 records were summarized by source cell stratum. Multiple strata from the same resource were not counted as independent cohorts. Absence of a returned PBC-FCRL3 pair was treated as a public-output coverage boundary rather than a powered negative result.

For allele-direction analysis, the PBC risk allele was mapped to the QTL effect allele in each compatible resource. Directional agreement was reported without comparing effect magnitudes across different expression scales. A consistent expression direction was not interpreted as mediation or as a therapeutic direction.

### Different-disease-study validation and source credible-set audit

FinnGen R12 and R13 `CHIRBIL_PRIM` metadata, regional R13 summary statistics and official disease fine-mapping records were obtained from the public release interfaces [16]. The release-stability audit compared R12 and R13 disease credible-set membership and probabilities in the IL12RB2 region. The R13 disease signal was bridged to the frozen OneK1K NK/IL12RB2 QTL by rsID and allele identity across GRCh38 and GRCh37 representations. After harmonization, 1,175 shared variants entered single-causal ABF colocalization with p1=p2=10^-4 and p12=10^-6, 10^-5 or 10^-4. OneK1K QTL credible-set members were then checked directly against the source FinnGen R13 disease credible set.

FinnGen R13 and the discovery PBC GWAS constitute different study systems, but person-level non-overlap could not be demonstrated from public metadata. This module was therefore treated as external disease-study validation with a source-CS audit, not as proof of an entirely independent cohort and not as a second source-LD multi-signal fit.

### Source-integrated chromatin evidence

Before querying the full CASCADE outputs, we fixed the 92 stable-H4 comparisons, comprising 27 genes and their corresponding broad cell strata. For those genes, we recorded relevant-cell eQTL availability and returned PBC-eQTL colocalization. The IL12RB2-linked peak `chr1-67307618-67308670` was prespecified from the earlier source record. We extracted the PBC-eQTL and PBC-caQTL colocalization records in l1.NK, source credible-set membership and the peak-IL12RB2 link [15]. The disease-eQTL and disease-caQTL edges were evaluated separately. A direct eQTL-caQTL posterior was unavailable and was not inferred from overlapping credible sets. The integrated result was classified as triangular signal coherence rather than chromatin mediation.

### Descriptive spatial-transcriptomic boundary

The public OMIX001122 processed spatial archive was checked at byte level and related to its original PBC study [17,18]. The released archive contained one control matrix and one PBC matrix without tissue images, spot coordinates or donor-replicated case-control structure. Matrix dimensions, total counts, positive spots and IL12RB2/FCRL3 detectability were recorded. No differential-expression or spatial-localization inference was attempted, and this layer was not counted as independent tissue replication.

"""
    text = replace_section(text, "### External molecular-QTL evaluation", "### Donor-level liver target-panel analysis", methods)
    text = text.replace(
        "HRA008003 is an open-access single-cell study of liver and blood in PBC [16,17].",
        "HRA008003 is an open-access single-cell study of liver and blood in PBC [19,20].",
    )
    text = text.replace(
        "Code, compact aggregate results, manifests and evidence ledgers were retained in the project repository [18].",
        "Code, compact aggregate results, manifests and evidence ledgers were retained in the project repository [21].",
    )

    results = """### IL12RB2-NK gained external disease-study and chromatin signal support

IL12RB2-NK was a representative stable-H4 comparison in OneK1K. Best PF10 signal-pair PP.H4 values were 0.9980 across L=5, 10 and 20; PF50 L=10 yielded PP.H4=0.9965. The comparison also contained a distinct signal pair, illustrating why stable H4 did not impose a single global mechanism label on the region.

TenK10K NK-cell eQTL data reproduced the PBC-IL12RB2 sharing pattern under the bounded single-causal replication model (396 overlapping variants; default PP.H4=0.9975; low-prior H4/[H3+H4]=0.9760). FinnGen immune multiome outputs returned three PBC-IL12RB2 molecular-QTL records, for l1.NK, l1.PBMC and l2.NK strata. PP.H4 ranged from 0.9698 to 0.9706, with credible-set overlap of five to nine variants. These strata came from one resource and were not counted as three independent cohorts.

The FinnGen disease signal was stable between public releases. The R12 disease credible set contained 19 variants and the R13 set contained 17; all 17 R13 members occurred in R12 (Jaccard=0.895), the same variant led both sets, and member probabilities correlated at r=0.833 (Figure 4A). FinnGen R13 contained 795 PBC cases and 368,758 controls. At the exact bridged disease variant rs6679356, the aligned C-allele effect was positive in both the discovery PBC GWAS and FinnGen R13.

FinnGen R13 × OneK1K NK/IL12RB2 colocalization used 1,175 harmonized variants. PP.H4 was 0.9772 at p12=10^-6, 0.9977 at p12=10^-5 and 0.9998 at p12=10^-4; H4/(H3+H4) was 0.9776, 0.9977 and 0.9998, respectively (Figure 4B). The shared top variant was rs6702599, with concordant effect direction, and both members of the OneK QTL credible set occurred in the FinnGen R13 disease credible set. This supported the regulatory association in a different disease study system. Public metadata did not establish person-level zero overlap, and the external analysis remained single-causal rather than source-LD multi-signal.

The prespecified CASCADE audit further constrained the chromatin interpretation. Of 27 fixed genes, 22 had relevant broad-cell eQTL coverage, but only IL12RB2 returned a PBC-eQTL colocalization record (Figure 4C). At the linked l1.NK peak, PBC-IL12RB2 eQTL sharing gave PP.H4=0.9703 and PBC-caQTL sharing gave PP.H4=0.9695. All nine eQTL credible-set members and all 11 caQTL credible-set members occurred within the 19-member disease credible set. The peak was 239 bp from the IL12RB2 transcription start site and linked to the gene with coefficient 0.141 (Figure 4D). These three observed edges supported triangular signal coherence. They did not supply the unavailable direct eQTL-caQTL posterior or establish chromatin-to-expression mediation.

The aligned PBC risk allele was associated with higher IL12RB2 expression across compatible OneK1K, TenK10K and FinnGen records. Effect magnitudes were not compared because the expression transformations and models differed. In the OMIX001122 processed archive, IL12RB2 and FCRL3 were detectable in both the single PBC and single control spatial matrices. The absence of donor replication, tissue-image coordinates and spot annotations restricted this observation to descriptive detectability rather than disease enrichment or spatial validation.

"""
    text = replace_section(
        text,
        "### IL12RB2-NK showed cross-resource support with an incomplete chromatin chain",
        "### FCRL3 combined reproducible B-cell sharing with a model-sensitive T-cell counterexample",
        results,
    )

    text = text.replace(
        "The external axes demonstrate how a reliability framework can strengthen biological interpretation without inflating mechanism. IL12RB2-NK was stable across OneK1K configurations, reproduced in TenK10K and supported by FinnGen molecular-QTL strata. The risk allele was consistently associated with higher expression. Prior PBC studies have also prioritized IL12RB2 through integrative genetic analyses [19,20], so the contribution here is not first nomination of the gene. It is the source-aware, cell-context-specific and signal-pair-bounded evaluation. The chromatin data remained positional because the eQTL anchor was absent from the peak caQTL credible set and disease-caQTL colocalization was not established.",
        "The external axes demonstrate how a reliability framework can strengthen biological interpretation without inflating mechanism. IL12RB2-NK was stable across OneK1K configurations, reproduced in TenK10K and supported by FinnGen molecular-QTL strata. FinnGen R13 then supported the same OneK NK/IL12RB2 association in a different disease study system across the complete prior series and at the source-CS level. Prior PBC studies have also prioritized IL12RB2 through integrative genetic analyses [22,23], so the contribution is not first nomination of the gene. It is a source-aware, cell-context-specific chain that joins multi-signal discovery evidence to bounded external disease-study validation. The added chromatin layer strengthens biological coherence because the disease signal shared with both the eQTL and a linked caQTL. The missing direct eQTL-caQTL posterior prevents this triangle from becoming evidence of causal mediation.",
    )
    text = text.replace(
        "Independent 2026 work has reported PBC-related FCRL3 evidence in liver B cells and plasma-protein colocalization [21,22].",
        "Independent 2026 work has reported PBC-related FCRL3 evidence in liver B cells and plasma-protein colocalization [24,25].",
    )
    text = text.replace(
        "Several limitations define the claim ceiling. First, only the high-information subset received multi-signal analysis; the study is PBC-wide screening followed by bounded reclassification, not a multi-signal analysis of all 6,923 comparisons. Second, real data contain no causal truth, so H4-to-H3 and H3-to-H4 changes are reclassifications rather than measured error corrections. Third, the simulation used two empirical 128-variant LD templates and a finite parameter grid. It describes single-configuration operating characteristics over that grid, not end-to-end calibration of the composite real-data gate or performance at every disease architecture, sample size or LD mismatch. Fourth, credible-set discovery limited pair-evaluable iterations; larger or more diverse simulations would not by themselves solve weak information in a given dataset. Fifth, OneK1K and TenK10K analyses reused the same PBC GWAS. Sixth, the liver study contained five donors per group, used non-lesion surgical controls and applied a bounded target panel rather than full-transcriptome state modeling. Finally, no perturbational experiment tested whether changing IL12RB2 or FCRL3 expression alters PBC-relevant immune or biliary phenotypes.",
        "Several limitations define the claim ceiling. First, only the high-information subset received multi-signal analysis; the study is PBC-wide screening followed by bounded reclassification, not a multi-signal analysis of all 6,923 comparisons. Second, real data contain no causal truth, so H4-to-H3 and H3-to-H4 changes are reclassifications rather than measured error corrections. Third, the simulation used two empirical 128-variant LD templates and a finite parameter grid. It describes operating characteristics over that grid rather than end-to-end calibration at every disease architecture, sample size or LD mismatch. Fourth, credible-set discovery limited pair-evaluable iterations. Fifth, OneK1K and TenK10K analyses reused the same discovery PBC GWAS. FinnGen R13 supplied a different disease study system, but had 795 cases, public metadata could not establish person-level non-overlap and the external analysis was single-causal rather than a second source-LD multi-signal fit. Sixth, the chromatin triangle lacked a direct eQTL-caQTL posterior, and the OMIX archive contained only one matrix per condition without spatial coordinates. Seventh, the liver study contained five donors per group, used non-lesion surgical controls and applied a bounded target panel rather than full-transcriptome state modeling. Finally, no perturbational experiment tested whether changing IL12RB2 or FCRL3 expression alters PBC-relevant immune or biliary phenotypes.",
    )
    text = text.replace(
        "Three next tests follow directly from these limits. The first priority is external disease replication: repeat signal-pair analysis when an independent, well-powered PBC GWAS with compatible ancestry, allele definitions and regional summaries becomes available. The second is model resolution: test PF-sensitive assignments using independent QTL cohorts with donor-level genotypes and prespecified covariate models. The third is mechanism: perturb IL12RB2 in NK cells and FCRL3 in B-cell subsets, then measure defined immune phenotypes and interactions with cholangiocyte-relevant systems. Those experiments would distinguish a reproducible regulatory association from direct biological mediation.",
        "Three next tests follow directly from these limits. The first priority is to repeat source-matched multi-signal analysis in a larger external PBC GWAS with documented participant separation and compatible ancestry. The second is to estimate the eQTL-caQTL relationship directly with common variant support and source-matched LD, rather than inferring it from two disease-sharing edges. The third is to perturb IL12RB2 in NK cells and FCRL3 in B-cell subsets, then measure defined immune phenotypes and interactions with cholangiocyte-relevant systems. These tests would determine whether the reproducible regulatory associations represent robust cross-study signal identity, chromatin-linked expression control or direct biological mediation.",
    )
    text = text.replace(
        "IL12RB2-NK and FCRL3 B-cell assignments received bounded cross-resource support, whereas liver data supported lineage detectability without corrected disease enrichment. The resulting framework improves interpretability without treating statistical signal sharing as causal mechanism.",
        "IL12RB2-NK additionally received bounded support in a different disease study system and a disease-linked chromatin triangle, whereas FCRL3 B-cell evidence remained cross-QTL-resource and liver data supported lineage detectability without corrected disease enrichment. The resulting framework improves interpretability without treating statistical signal sharing as causal mechanism.",
    )

    availability_old = (
        "PBC summary statistics are available through GWAS Catalog study GCST90061440. Study-derived PBC regional statistics and LD are available from the GJOKA archive. "
        "The OneK1K cohort, updated release and TenK10K resource are identified separately in their cited publications and Zenodo records. HRA008003 is available through GSA-Human. "
        "Code, compact aggregate outputs, manifests, frozen protocols, source-bound figures and submission-interface assets are versioned in the public repository at `https://github.com/1209433622cz-maker/igan-open-regulatory-evidence/tree/r7b4a-human-genomics-interface-2026-10-10` [18]. "
        "Additional files 1 and 2 contain the released supplementary workbook and machine-readable derived tables. Individual-level genotype, BAM, pseudobulk matrices and source archives governed by third-party licences are not redistributed; their source accessions and conditions are recorded in Supplementary Table S1."
    )
    availability_new = (
        "PBC summary statistics are available through GWAS Catalog study GCST90061440. Study-derived PBC regional statistics and LD are available from the GJOKA archive. "
        "The OneK1K cohort, updated release and TenK10K resource are identified separately in their cited publications and Zenodo records. FinnGen R13 CHIRBIL_PRIM summary statistics and fine-mapping records are publicly distributed by FinnGen [16]; CASCADE records are accessible through the FinnGen immune multiome portal [15]. "
        "HRA008003 is available through GSA-Human, and OMIX001122 through the National Genomics Data Center [17]. Code, compact aggregate outputs, manifests, frozen protocols, source-bound figures and submission assets are versioned in the public repository at `https://github.com/1209433622cz-maker/igan-open-regulatory-evidence/tree/r7c1-il12rb2-external-validation-2026-10-11` [21]. "
        "Additional files 1 and 2 contain the released supplementary workbook and primary machine-readable tables; Additional file 3 contains the R7C1 external-validation module. Individual-level genotype, BAM, pseudobulk matrices and third-party source archives are not redistributed."
    )
    if availability_old not in text:
        raise RuntimeError("availability anchor changed")
    text = text.replace(availability_old, availability_new)

    references = """## References

1. Gulamhusein AF, Hirschfield GM. Primary biliary cholangitis: pathogenesis and therapeutic opportunities. *Nat Rev Gastroenterol Hepatol.* 2020;17:93-110. doi:10.1038/s41575-019-0226-7.
2. Cordell HJ, Fryett JJ, Ueno K, et al. An international genome-wide meta-analysis of primary biliary cholangitis: novel risk loci and candidate drugs. *J Hepatol.* 2021;75:572-581. doi:10.1016/j.jhep.2021.04.055.
3. Yazar S, Alquicira-Hernandez J, Wing K, et al. Single-cell eQTL mapping identifies cell type-specific genetic control of autoimmune disease. *Science.* 2022;376:eabf3041. doi:10.1126/science.abf3041.
4. Cuomo ASE, Spenceley E, Tanudisastro HA, et al. Impact of rare and common genetic variation on cell type-specific gene expression in human blood. *medRxiv.* 2025. doi:10.1101/2025.03.20.25324352. Dataset: Zenodo record 18221260.
5. Giambartolomei C, Vukcevic D, Schadt EE, et al. Bayesian test for colocalisation between pairs of genetic association studies using summary statistics. *PLoS Genet.* 2014;10:e1004383. doi:10.1371/journal.pgen.1004383.
6. Wang G, Sarkar A, Carbonetto P, Stephens M. A simple new approach to variable selection in regression, with application to genetic fine mapping. *J R Stat Soc Series B Stat Methodol.* 2020;82:1273-1300. doi:10.1111/rssb.12388.
7. Wallace C. A more accurate method for colocalisation analysis allowing for multiple causal variants. *PLoS Genet.* 2021;17:e1009440. doi:10.1371/journal.pgen.1009440.
8. NHGRI-EBI GWAS Catalog. Study GCST90061440. https://www.ebi.ac.uk/gwas/studies/GCST90061440. Accessed 10 Oct 2026.
9. Gjoka A, Cordell HJ. Fine-mapping the results from genome-wide association studies of primary biliary cholangitis using SuSiE and h2-D2. *Genet Epidemiol.* 2025;49:e22592. doi:10.1002/gepi.22592.
10. Cordell HJ, Gjoka A. PBC fine-mapping summary statistics and study-derived LD archive. https://www.staff.ncl.ac.uk/heather.cordell/GjokaPaper.html. Accessed 10 Oct 2026.
11. Xue A, Yazar S, Alquicira-Hernández J, et al. Genetic variants associated with cell-type-specific intra-individual gene expression variability reveal mechanisms of genome regulation. *Nat Commun.* 2026. Published online September 21, 2026. doi:10.1038/s41467-026-77453-9.
12. Xue A, et al. OneK1K cell-specific cis-eQTL summary statistics, genotypes and covariates. Zenodo. 2026. https://doi.org/10.5281/zenodo.18910121.
13. Zou Y, Carbonetto P, Wang G, Stephens M. Fine-mapping from summary data with the \"Sum of Single Effects\" model. *PLoS Genet.* 2022;18(7):e1010299. doi:10.1371/journal.pgen.1010299.
14. Cuomo ASE, et al. TenK10K phase 1 molecular-QTL and fine-mapping resources. Zenodo. 2026. https://doi.org/10.5281/zenodo.18221260.
15. Kanai M, Delorey TM, Honkanen J, et al. Population-scale immune multiome atlas reveals regulatory disease mechanisms. *Nature.* 2026. Published online September 30, 2026. doi:10.1038/s41586-026-11078-2.
16. FinnGen. Release 13 summary statistics and fine-mapping records for CHIRBIL_PRIM. https://r13.finngen.fi/pheno/CHIRBIL_PRIM. Accessed 11 Oct 2026.
17. National Genomics Data Center. OMIX001122: single-cell RNA profiling and spatial transcriptomics of primary biliary cholangitis. https://ngdc.cncb.ac.cn/omix/release/OMIX001122. Accessed 11 Oct 2026.
18. Li X, Li Y, Xiao J, et al. Unique DUOX2+ACE2+ small cholangiocytes are pathogenic targets for primary biliary cholangitis. *Nat Commun.* 2023;14:29. doi:10.1038/s41467-022-34606-w.
19. Jin C, Jiang P, Zhang Z, et al. Single-cell RNA sequencing reveals the pro-inflammatory roles of liver-resident Th1-like cells in primary biliary cholangitis. *Nat Commun.* 2024;15:8690. doi:10.1038/s41467-024-53104-9.
20. National Genomics Data Center. GSA-Human HRA008003. https://ngdc.cncb.ac.cn/gsa-human/browse/HRA008003. Accessed 10 Oct 2026.
21. Open Regulatory Evidence Project. Analysis code, frozen protocols and aggregate results. GitHub. https://github.com/1209433622cz-maker/igan-open-regulatory-evidence. Accessed 11 Oct 2026.
22. Wang Q, Yan H, Wang N, Yuan J. Causal genetics prioritizes therapeutic targets and blood signatures in primary biliary cholangitis. *npj Syst Biol Appl.* 2026. doi:10.1038/s41540-026-00817-w.
23. Lin Y, Xu X, Chen H, Tang H, Hua Y, Mei J. Genetic mechanisms of primary biliary cholangitis and its association with immune cells using Mendelian randomization and biological annotation. *Medicine (Baltimore).* 2026;105(22):e49146. doi:10.1097/MD.0000000000049146.
24. Han Z, Ran Y, Liu R, et al. FCRL3 as a potential link between Benzo[a]pyrene exposure and primary biliary cholangitis: insights from comparative toxicogenomics and multi-omics analysis. *BMC Gastroenterol.* 2026;26:101. doi:10.1186/s12876-026-04614-x.
25. Xi J, Wang S, Chen J, et al. Exploration and discovery of treatment targets for primary biliary cholangitis based on plasma and cerebrospinal fluid proteomics: a multicenter Mendelian randomization study. *PLoS One.* 2026;21(2):e0340166. doi:10.1371/journal.pone.0340166.

"""
    text = replace_section(text, "## References", "## Figure legends", references)

    fig4_old = re.search(r"\*\*Figure 4\..*?(?=\n\n\*\*Figure 5\.)", text, flags=re.S)
    if not fig4_old:
        raise RuntimeError("Figure 4 legend not found")
    fig4_new = (
        "**Figure 4. IL12RB2-NK gains external disease-study and chromatin signal support within bounded claims.** "
        "(A) FinnGen R12/R13 IL12RB2-region disease credible-set stability. (B) Single-causal FinnGen R13 × OneK1K NK/IL12RB2 posterior across the frozen p12 series; the different disease study system does not prove person-level zero overlap. "
        "(C) Prespecified CASCADE coverage among the 27 genes represented by the 92 stable-H4 comparisons. (D) Separate PBC-eQTL sharing, PBC-caQTL sharing and peak-IL12RB2 linkage. The dashed edge marks the unavailable direct eQTL-caQTL posterior. OMIX001122 contained one PBC and one control processed matrix and was retained as descriptive evidence only.  \n"
        "**Abbreviations:** caQTL, chromatin-accessibility quantitative-trait locus; CS, credible set; eQTL, expression quantitative-trait locus; H4, shared-signal posterior hypothesis; PBC, primary biliary cholangitis; PIP, posterior inclusion probability."
    )
    text = text[:fig4_old.start()] + fig4_new + text[fig4_old.end():]

    text = text.replace(
        "**Additional file 2.** `Additional_file_2_Machine_Readable_Supplementary_Data.zip` (ZIP archive). Machine-readable TSV, TSV.GZ and JSON source/derived tables, plus a file-level SHA-256 manifest and README. The archive excludes individual-level genotype, BAM, pseudobulk matrices and licensed source archives.\n\nThe two files contain the following prespecified supplementary modules:",
        "**Additional file 2.** `Additional_file_2_Machine_Readable_Supplementary_Data.zip` (ZIP archive). Machine-readable TSV, TSV.GZ and JSON source/derived tables, plus a file-level SHA-256 manifest and README. The archive excludes individual-level genotype, BAM, pseudobulk matrices and licensed source archives.\n\n**Additional file 3.** `Additional_file_3_R7C1_External_Validation_Data.zip` (ZIP archive). Compact FinnGen R12/R13 disease-CS, FinnGen R13 × OneK1K, source-CS, fixed CASCADE, chromatin-triangle, OMIX descriptive-QC and source-ledger tables with an internal SHA-256 manifest.\n\nThe three files contain the following prespecified supplementary modules:",
    )
    text += "\n\n**Supplementary Table S11.** R7C1 IL12RB2-NK external disease-study validation, source credible-set audit, CASCADE coverage, chromatin-triangle adjudication and OMIX descriptive boundary. Delivered in Additional file 3.\n"
    return "\n".join(line.rstrip() for line in text.splitlines()) + "\n"


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    paragraph._p.append(field)


def add_line_numbering(section):
    sect_pr = section._sectPr
    node = sect_pr.find(qn("w:lnNumType"))
    if node is None:
        node = OxmlElement("w:lnNumType")
        sect_pr.append(node)
    node.set(qn("w:countBy"), "1")
    node.set(qn("w:restart"), "continuous")
    node.set(qn("w:distance"), "360")


def add_rich_text(paragraph, text: str):
    parts = re.split(r"(\*\*.*?\*\*|`.*?`|\*.*?\*)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Courier New"
            run.font.size = Pt(9)
        elif part.startswith("*") and part.endswith("*"):
            paragraph.add_run(part[1:-1]).italic = True
        else:
            paragraph.add_run(part)


def configure_styles(doc: Document):
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.first_line_indent = Cm(0.74)
    normal.paragraph_format.widow_control = True
    for name, size in (("Title", 17), ("Heading 1", 14), ("Heading 2", 12), ("Heading 3", 11)):
        style = doc.styles[name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(12 if name != "Title" else 0)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.line_spacing = 1.05
    caption = doc.styles["Figure Caption"] if "Figure Caption" in doc.styles else doc.styles.add_style("Figure Caption", WD_STYLE_TYPE.PARAGRAPH)
    caption.font.name = "Times New Roman"
    caption.font.size = Pt(9)
    caption.paragraph_format.line_spacing = 1.0
    caption.paragraph_format.space_after = Pt(6)
    caption.paragraph_format.first_line_indent = Cm(0)


def figure_path(number: int) -> Path:
    names = {
        1: "Figure1_study_scope_and_model_integrity.png",
        2: "Figure2_input_matched_reclassification.png",
        3: "Figure3_simulation_discovery_and_inference.png",
        4: "Figure4_IL12RB2_external_evidence.png",
        5: "Figure5_FCRL3_support_and_counterexample.png",
        6: "Figure6_liver_tissue_boundary.png",
    }
    return WORK / "figures" / names[number]


def build_docx(md_text: str) -> None:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.2)
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.header_distance = Cm(1)
    section.footer_distance = Cm(1)
    add_line_numbering(section)
    add_page_number(section.footer.paragraphs[0])
    configure_styles(doc)
    doc.core_properties.title = TITLE
    doc.core_properties.subject = "R7C1 external-validation manuscript fork"
    doc.core_properties.author = "Zhi Chen; Teng Qi"
    doc.core_properties.keywords = "PBC; colocalization; eQTL; caQTL; source-matched LD; IL12RB2"

    in_legends = False
    for raw in md_text.splitlines():
        line = raw.rstrip()
        if not line:
            continue
        if line.startswith("# "):
            paragraph = doc.add_paragraph(style="Title")
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.first_line_indent = Cm(0)
            add_rich_text(paragraph, line[2:])
        elif line == "## Figure legends":
            doc.add_page_break()
            doc.add_heading("Figure legends", level=1)
            in_legends = True
        elif line.startswith("## "):
            doc.add_heading(line[3:], level=1)
        elif line.startswith("### "):
            doc.add_heading(line[4:], level=2)
        elif in_legends and line.startswith("**Figure "):
            match = re.match(r"\*\*Figure (\d+)\.", line)
            if not match:
                raise RuntimeError(f"bad figure legend: {line[:60]}")
            number = int(match.group(1))
            if number > 1:
                doc.add_page_break()
            p_img = doc.add_paragraph()
            p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_img.paragraph_format.first_line_indent = Cm(0)
            p_img.add_run().add_picture(str(figure_path(number)), width=Inches(6.45))
            p_cap = doc.add_paragraph(style="Figure Caption")
            add_rich_text(p_cap, line)
        elif in_legends and line.startswith("**Abbreviations:**"):
            p = doc.add_paragraph(style="Figure Caption")
            add_rich_text(p, line)
        else:
            p = doc.add_paragraph()
            if line.startswith(("**Authors:", "**Affiliations:", "**Corresponding author:", "**Article type:")):
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.first_line_indent = Cm(0)
                p.paragraph_format.line_spacing = 1.15
            if line.startswith("G* =") or line.startswith("and R_QTL"):
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.first_line_indent = Cm(0)
            add_rich_text(p, line.rstrip("  "))
    OUT_DOCX.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT_DOCX)


def main() -> None:
    text = build_markdown()
    OUT_MD.parent.mkdir(parents=True, exist_ok=True)
    OUT_MD.write_text(text, encoding="utf-8")
    build_docx(text)
    state = {
        "stage": "R7C1_IL12RB2_EXTERNAL_VALIDATION_MANUSCRIPT_FORK",
        "baseline_md_sha256": sha256(BASELINE),
        "output_md": str(OUT_MD),
        "output_md_sha256": sha256(OUT_MD),
        "output_docx": str(OUT_DOCX),
        "output_docx_sha256": sha256(OUT_DOCX),
        "figure4_sha256": sha256(WORK / "figures/Figure4_IL12RB2_external_evidence.png"),
        "additional_file_3_sha256": sha256(WORK / "supplement/Additional_file_3_R7C1_External_Validation_Data.zip"),
        "claim_ceiling": "E2_ROBUST_ASSOCIATION; NO_DIRECT_MEDIATION",
    }
    (WORK / "results/R7C1_manuscript_build_state.json").write_text(
        json.dumps(state, indent=2, ensure_ascii=False), encoding="utf-8"
    )
    print(json.dumps(state, indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()
