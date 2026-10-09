#!/usr/bin/env python3
"""Create the author-review DOCX from the frozen R7B2 Markdown manuscript."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(r"H:\SCI2\YR1")
BASE = ROOT / "5_manuscript" / "R7B2_ManuscriptV1"
MD = BASE / "manuscript" / "R7B2_full_english_manuscript_v1.md"
OUT = BASE / "manuscript" / "R7B2_full_english_manuscript_v1.docx"
FIG = BASE / "figures"


def set_cell_margins(cell, top=90, start=90, bottom=90, end=90):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tcMar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tcMar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    run._r.addnext(fld)


def add_line_numbering(section):
    sectPr = section._sectPr
    ln = sectPr.find(qn("w:lnNumType"))
    if ln is None:
        ln = OxmlElement("w:lnNumType")
        sectPr.append(ln)
    ln.set(qn("w:countBy"), "1")
    ln.set(qn("w:restart"), "continuous")
    ln.set(qn("w:distance"), "360")


def add_rich_text(paragraph, text: str):
    # Minimal Markdown support for bold and inline code; scientific content remains editable.
    parts = re.split(r"(\*\*.*?\*\*|`.*?`)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            r = paragraph.add_run(part[2:-2]); r.bold = True
        elif part.startswith("`") and part.endswith("`"):
            r = paragraph.add_run(part[1:-1]); r.font.name = "Courier New"; r.font.size = Pt(9)
        else:
            paragraph.add_run(part.replace("*", ""))


def figure_path(n: int) -> Path:
    names = {
        1: "Figure1_study_scope_and_model_integrity.png",
        2: "Figure2_input_matched_reclassification.png",
        3: "Figure3_simulation_discovery_and_inference.png",
        4: "Figure4_IL12RB2_external_evidence.png",
        5: "Figure5_FCRL3_support_and_counterexample.png",
        6: "Figure6_liver_tissue_boundary.png",
    }
    return FIG / names[n]


def configure_styles(doc: Document):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.first_line_indent = Cm(0.74)
    normal.paragraph_format.widow_control = True

    for name, size, color in [("Title", 17, "17365D"), ("Heading 1", 14, "17365D"), ("Heading 2", 12, "1F4E79"), ("Heading 3", 11, "1F4E79")]:
        s = styles[name]
        s.font.name = "Arial"
        s._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
        s.font.size = Pt(size)
        s.font.bold = True
        s.font.color.rgb = RGBColor.from_string(color)
        s.paragraph_format.keep_with_next = True
        s.paragraph_format.space_before = Pt(12 if name != "Title" else 0)
        s.paragraph_format.space_after = Pt(6)
        s.paragraph_format.line_spacing = 1.05

    if "Figure Caption" not in styles:
        cap = styles.add_style("Figure Caption", WD_STYLE_TYPE.PARAGRAPH)
    else:
        cap = styles["Figure Caption"]
    cap.font.name = "Times New Roman"; cap.font.size = Pt(9)
    cap.paragraph_format.line_spacing = 1.0; cap.paragraph_format.space_after = Pt(6)
    cap.paragraph_format.first_line_indent = Cm(0)


def main():
    text = MD.read_text(encoding="utf-8")
    lines = text.splitlines()
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Cm(2.2); sec.bottom_margin = Cm(2.2)
    sec.left_margin = Cm(2.5); sec.right_margin = Cm(2.2)
    sec.page_width = Cm(21.0); sec.page_height = Cm(29.7)
    sec.header_distance = Cm(1.0); sec.footer_distance = Cm(1.0)
    add_line_numbering(sec)
    add_page_number(sec.footer.paragraphs[0])
    configure_styles(doc)

    props = doc.core_properties
    props.title = "Input-matched multi-signal analysis clarifies immune-cell regulatory assignments at PBC risk loci"
    props.subject = "R7B2 full English manuscript v1"
    props.author = "[AUTHORS TO BE COMPLETED]"
    props.keywords = "PBC; colocalization; eQTL; SuSiE; source-matched LD"

    in_figure_legends = False
    last_was_blank = False
    for line in lines:
        raw = line.rstrip()
        if not raw:
            last_was_blank = True
            continue
        if raw.startswith("# "):
            p = doc.add_paragraph(style="Title")
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.first_line_indent = Cm(0)
            add_rich_text(p, raw[2:])
            continue
        if raw == "## Figure legends":
            doc.add_page_break()
            doc.add_heading("Figure legends", level=1)
            in_figure_legends = True
            continue
        if raw.startswith("## "):
            doc.add_heading(raw[3:], level=1)
            continue
        if raw.startswith("### "):
            doc.add_heading(raw[4:], level=2)
            continue
        if in_figure_legends and raw.startswith("**Figure "):
            m = re.match(r"\*\*Figure (\d+)\.", raw)
            if m:
                n = int(m.group(1))
                if n > 1:
                    doc.add_page_break()
                pimg = doc.add_paragraph()
                pimg.alignment = WD_ALIGN_PARAGRAPH.CENTER
                pimg.paragraph_format.first_line_indent = Cm(0)
                pimg.add_run().add_picture(str(figure_path(n)), width=Inches(6.45))
                p = doc.add_paragraph(style="Figure Caption")
                add_rich_text(p, raw)
                continue
        if in_figure_legends and raw.startswith("**Abbreviations:**"):
            p = doc.add_paragraph(style="Figure Caption")
            add_rich_text(p, raw)
            continue
        p = doc.add_paragraph()
        if raw.startswith("**Authors:") or raw.startswith("**Affiliations:") or raw.startswith("**Corresponding author:") or raw.startswith("**Article type:"):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.line_spacing = 1.15
        if raw.startswith("G* =") or raw.startswith("and R_QTL"):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.first_line_indent = Cm(0)
        add_rich_text(p, raw.rstrip("  "))
        last_was_blank = False

    # Add a small completion table for the author metadata that remains genuinely unknown.
    doc.add_page_break()
    doc.add_heading("Author metadata completion gate", level=1)
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    hdr[0].text = "Field"; hdr[1].text = "Required verified information"
    for field, need in [
        ("Authors and order", "Full names, order, equal-contribution marks"),
        ("Affiliations", "Numbered official institutional addresses"),
        ("Correspondence", "Name, address, email, telephone"),
        ("CRediT", "Author-specific contribution roles"),
        ("Funding", "Exact funder and grant numbers, or verified no-funding statement"),
        ("Competing interests", "Verified author declaration"),
        ("Ethics", "Local secondary-analysis determination or waiver wording"),
    ]:
        c = table.add_row().cells; c[0].text = field; c[1].text = need
        for x in c: set_cell_margins(x)
    for c in hdr:
        set_cell_margins(c)
        for r in c.paragraphs[0].runs: r.bold = True

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
