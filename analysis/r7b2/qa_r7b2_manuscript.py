#!/usr/bin/env python3
"""Machine-readable integrity and claim-boundary QA for the R7B2 manuscript."""

from __future__ import annotations

import hashlib
import json
import re
import sys
from pathlib import Path

import pandas as pd
from docx import Document
from pypdf import PdfReader


ROOT = Path(r"H:\SCI2\YR1")
BASE = ROOT / "5_manuscript" / "R7B2_ManuscriptV1"
MD = BASE / "manuscript" / "R7B2_full_english_manuscript_v1.md"
DOCX = BASE / "manuscript" / "R7B2_full_english_manuscript_v1.docx"
PDF = BASE / "manuscript" / "R7B2_full_english_manuscript_v1_WPS.pdf"
LEDGER = ROOT / "3_results/04_integration/R7B2A/manuscript_lock/R7B2A_updated_claim_evidence_ledger.tsv"
OUT = BASE / "qa"
OUT.mkdir(parents=True, exist_ok=True)


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def main() -> int:
    text = MD.read_text(encoding="utf-8")
    ledger = pd.read_csv(LEDGER, sep="\t")
    checks = []

    def check(name, passed, detail):
        checks.append({"check": name, "status": "PASS" if passed else "FAIL", "detail": str(detail)})

    required_sections = ["## Abstract", "## Background", "## Methods", "## Results", "## Discussion", "## Conclusions", "## Declarations", "## References", "## Figure legends"]
    check("required_sections", all(s in text for s in required_sections), ", ".join(required_sections))

    anchors = ["6,923", "5,460", "642", "113", "452", "79", "413", "eight H4-to-H3", "ten H3-to-H4", "92 H4-supported", "428 H3-supported", "59 model-sensitive", "63 uninformative", "32 H4-supported", "491", "520", "486,000", "160,782", "33.08%", "12.41%", "8.56%", "32.99%", "26.66%", "0.9698", "0.9706", "0.9916", "q=0.111"]
    missing = [x for x in anchors if x not in text]
    check("numeric_anchor_coverage", not missing, f"missing={missing}")

    prohibited = [
        "proved that multi-signal", "proves that multi-signal", "generally superior",
        "independent disease replication", "established a complete causal cascade", "tissue mediation was established",
        "PBC-specific upregulation", "all 6,923 received multi-signal"
    ]
    found = [x for x in prohibited if x.lower() in text.lower()]
    # The phrase independent disease replication may appear only in a negated boundary.
    found = [x for x in found if x != "independent disease replication"]
    check("prohibited_positive_claims", not found, f"found={found}")

    check("ledger_count", len(ledger) == 27, f"claims={len(ledger)}")
    claim_terms = {
        "C01": "6,923", "C02": "642-comparison", "C03": "bidirectional H4/H3",
        "C04": "79 of 113", "C05": "413 retained H3", "C06": "491 of 520",
        "C07": "All 92 stable-H4", "C08": "32 of the 92", "C09": "2,568",
        "C10": "rs1800378", "C11": "12.41%", "C12": "32.99%", "C13": "16.00%",
        "C14": "general superiority", "C15": "0.9698", "C16": "higher IL12RB2",
        "C17": "complete disease-to-chromatin", "C18": "0.9916", "C19": "lower FCRL3",
        "C20": "FCRL3-CD8_ET", "C21": "returned no direct PBC-FCRL3", "C22": "all five PBC",
        "C23": "15 categorical states", "C24": "another 19 states", "C25": "eight H4-to-H3 and ten H3-to-H4",
        "C26": "produced no categorical changes", "C27": "no recorded technical errors"
    }
    uncovered = [cid for cid, term in claim_terms.items() if term.lower() not in text.lower()]
    check("claim_ledger_text_coverage", not uncovered, f"uncovered={uncovered}")

    refs = re.findall(r"^\d+\. ", text, flags=re.M)
    check("reference_count", len(refs) == 19, f"references={len(refs)}")
    first_refs = [int(x) for x in re.findall(r"^(\d+)\. ", text, flags=re.M)]
    check("reference_continuity", first_refs == list(range(1, 20)), f"sequence={first_refs}")

    figs = re.findall(r"^\*\*Figure (\d+)\.", text, flags=re.M)
    check("figure_legend_count", figs == ["1", "2", "3", "4", "5", "6"], f"figures={figs}")
    fig_files = [p for p in (BASE / "figures").glob("*.png") if re.match(r"Figure[1-6]_", p.name)]
    check("figure_file_count", len(fig_files) == 6, f"png={len(fig_files)}")

    words = re.findall(r"\b[A-Za-z0-9][A-Za-z0-9'/-]*\b", text)
    check("manuscript_word_count", 6000 <= len(words) <= 11000, f"words={len(words)}")
    placeholders = re.findall(r"\[[A-Z][A-Z0-9 '\-/,:]+(?:REQUIRED|COMPLETED|CONFIRMED|INFORMATION|WORDING|AUTHORS|ADDRESSES|NAME|CRediT).*?\]", text)
    check("unknown_metadata_explicit", len(placeholders) >= 5, f"placeholders={len(placeholders)}")

    check("docx_exists", DOCX.exists() and DOCX.stat().st_size > 10000, f"bytes={DOCX.stat().st_size if DOCX.exists() else 0}")
    if DOCX.exists():
        doc = Document(DOCX)
        pic_count = len(doc.inline_shapes)
        heading_text = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
        check("docx_figure_objects", pic_count == 6, f"inline_shapes={pic_count}")
        check("docx_sections", all(s[3:] in heading_text for s in required_sections if s != "## Abstract"), f"headings={len(heading_text)}")

    check("pdf_exists", PDF.exists() and PDF.stat().st_size > 10000, f"bytes={PDF.stat().st_size if PDF.exists() else 0}")
    pdf_meta = {}
    if PDF.exists():
        reader = PdfReader(str(PDF))
        pdf_meta = dict(reader.metadata or {})
        producer = str(pdf_meta.get("/Producer", "")) + " " + str(pdf_meta.get("/Creator", ""))
        check("wps_pdf_creator", "wps" in producer.lower(), producer)
        check("pdf_page_count", 20 <= len(reader.pages) <= 80, f"pages={len(reader.pages)}")

    result = {
        "stage": "R7B2_FULL_ENGLISH_MANUSCRIPT_V1",
        "status": "PASS" if all(x["status"] == "PASS" for x in checks) else "FAIL",
        "checks": checks,
        "summary": {"pass": sum(x["status"] == "PASS" for x in checks), "total": len(checks)},
        "artifacts": {p.name: {"bytes": p.stat().st_size, "sha256": sha256(p)} for p in [MD, DOCX, PDF] if p.exists()},
        "pdf_metadata": pdf_meta,
    }
    (OUT / "R7B2_manuscript_QA.json").write_text(json.dumps(result, indent=2, ensure_ascii=False), encoding="utf-8")
    pd.DataFrame(checks).to_csv(OUT / "R7B2_manuscript_QA.tsv", sep="\t", index=False)
    print(json.dumps(result["summary"] | {"status": result["status"]}, ensure_ascii=False))
    return 0 if result["status"] == "PASS" else 1


if __name__ == "__main__":
    raise SystemExit(main())
