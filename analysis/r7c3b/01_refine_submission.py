"""R7C3B bounded QiTeng refinement. No scientific analyses are rerun."""
from pathlib import Path
from copy import deepcopy
import csv, hashlib, io, json, re, shutil, zipfile
import pymupdf as fitz
from PIL import Image, ImageChops, ImageDraw
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(r'H:\SCI2\YR1')
SRC = ROOT/'0_admin/intake/R7C3A_20261011'
OUT = ROOT/'5_manuscript/R7C3B_QiTeng_WPS_SubmissionRefinement_20261011'
for n in ['01_Manuscript','02_Figures','03_Supplement','04_Cover','05_Source_Data','06_QA','07_Reports','08_Protocol','09_Archive_Original']:
    (OUT/n).mkdir(parents=True, exist_ok=True)
def h(b): return hashlib.sha256(b).hexdigest()
def dump(n,obj): (OUT/'06_QA'/n).write_text(json.dumps(obj,ensure_ascii=False,indent=2),encoding='utf-8')
def tsv(n,rows):
    with (OUT/'06_QA'/n).open('w',encoding='utf-8-sig',newline='') as f:
        w=csv.DictWriter(f,fieldnames=list(rows[0]),delimiter='\t');w.writeheader();w.writerows(rows)

# Archive the exact received release rather than overwrite any approved version.
original_zip=Path(r'C:\Users\Administrator\Downloads\CMM_R7C3A_QiTeng_Refined_Manuscript_Figures_SubmissionReview_20261011.zip')
shutil.copy2(original_zip,OUT/'09_Archive_Original'/original_zip.name)
shutil.copy2(SRC/'input_integrity_independent.json',OUT/'06_QA/Input_R7C3A_36_checksums.json')
shutil.copy2(SRC/'crossref_reference_metadata.json',OUT/'06_QA/Crossref_Metadata_Receipts.json')

# Production derivatives: replace panel letters only. Frozen source PDFs remain in input/archive.
figrows=[]
for num in range(1,7):
    f=SRC/f'02_Figures/Figure_{num}_SOURCE_FROZEN.pdf'
    d=fitz.open(f);p=d[0];old_pix=p.get_pixmap(matrix=fitz.Matrix(2,2),alpha=False)
    labels=[]
    for b in p.get_text('rawdict')['blocks']:
        for ln in b.get('lines',[]):
            for sp in ln['spans']:
                text=''.join(c['c'] for c in sp['chars'])
                if re.match(r'^[A-F](?:\s|$)',text) and sp['size']>=10:
                    char=sp['chars'][0];labels.append((char['c'],fitz.Rect(char['bbox']),fitz.Point(char['origin']),sp['size'],sp['color']))
    expected=6 if num==2 else (3 if num==6 else 4)
    assert len(labels)==expected,(num,labels)
    for label,r,origin,size,color in labels:
        p.add_redact_annot(r,fill=(1,1,1))
    p.apply_redactions(images=0,graphics=0,text=0)
    for label,r,origin,size,color in labels:
        rgb=tuple(((color>>s)&255)/255 for s in [16,8,0])
        p.insert_text(origin,label.lower()+')',fontname='hebo',fontsize=size,color=rgb)
    dest=OUT/f'02_Figures/Figure_{num}.pdf';d.save(dest,garbage=4,deflate=True)
    d.close();d=fitz.open(dest);p=d[0]
    new_pix=p.get_pixmap(matrix=fitz.Matrix(2,2),alpha=False)
    a=Image.frombytes('RGB',[old_pix.width,old_pix.height],old_pix.samples)
    b=Image.frombytes('RGB',[new_pix.width,new_pix.height],new_pix.samples)
    diff=ImageChops.difference(a,b)
    draw=ImageDraw.Draw(diff)
    label_masks=[]
    for label,r,origin,size,_ in labels:
        # Explicit allowance for new glyph antialiasing around each isolated panel letter.
        new_r=fitz.Rect(r.x0,r.y0,origin.x+fitz.get_text_length(label.lower()+')',fontname='hebo',fontsize=size),r.y1)
        mask=r | new_r;label_masks.append(list(mask))
        draw.rectangle([(int((mask.x0-2)*2),int((mask.y0-2)*2)),(int((mask.x1+2)*2),int((mask.y1+2)*2))],fill=(0,0,0))
    assert diff.getbbox() is None,('unexpected figure change',num,diff.getbbox())
    p.get_pixmap(matrix=fitz.Matrix(300/72,300/72),alpha=False).save(OUT/f'02_Figures/Figure_{num}_embed.png')
    figrows.append({'figure':num,'source_sha256':h(f.read_bytes()),'derivative_sha256':h(dest.read_bytes()),'changed_object':'PANEL_LABELS_ONLY_LOWERCASE_WITH_CLOSING_PARENTHESIS','letters':len(labels),'label_masks_pdf_points':label_masks,'outside_label_raster_delta':'ZERO','data_replotted':False})
dump('Figure_Production_Derivative_Identity.json',figrows)

base=Document(SRC/'01_Manuscript/R7C3A_npj_Manuscript_FigureLegends_Refined.docx')
orig={i:p.text for i,p in enumerate(base.paragraphs)}
edits=[]
def replace(i,new,reason,operation='REWRITE_SAME_JOB'):
    p=base.paragraphs[i];old=p.text
    assert not p._p.xpath('.//w:drawing'),('attempt to overwrite figure',i)
    p.text=new;edits.append({'source_paragraph':i,'operation':operation,'reason':reason,'before':old,'after':new,'numeric_source_changed':'NO'})

replace(6,
 'Regulatory assignments can change when colocalization analyses use different inputs or causal assumptions. We screened 5,460 of 6,923 primary biliary cholangitis locus-gene-cell comparisons and applied source-matched multi-signal analysis to a prespecified 642-comparison subset. Matching variant support and disease statistics changed 15 and 19 classifications, respectively. Eight shared-signal (H4) comparisons became distinct-signal (H3), and ten H3 comparisons became H4. Across 486,000 simulations, multi-signal inference reduced false sharing from 12.41% to 8.56% but reduced shared-signal recovery from 32.99% to 26.66%. IL12RB2-NK sharing received external disease-study and linked chromatin support, without establishing participant independence or mediation. Liver expression was detectable without corrected disease enrichment. These findings identify input matching, signal-pair identity, covariate sensitivity and discovery adequacy as complementary requirements for interpreting regulatory assignments. Statistical sharing supports regulatory association while direct mechanism remains untested.',
 'Define H3/H4 on first use; propagate simulation trade-off and tissue null; retain evidence ceiling within the 150-word limit.')
replace(6,base.paragraphs[6].text.replace('IL12RB2-NK sharing','IL12RB2 sharing in natural killer (NK) cells'),
 'Define NK at its first abstract appearance.')
replace(9,orig[9].replace('Primary biliary cholangitis is','Primary biliary cholangitis (PBC) is'),
 'Define PBC at first appearance in the main text.')
replace(11,orig[11].replace('Bayesian ABF colocalization','Approximate Bayes factor (ABF) colocalization'),
 'Define ABF at first appearance rather than rely on its legend definition.')
replace(13,
 'We used PBC to assess the reliability of immune-cell regulatory assignment through four linked tasks. First, we screened PBC-wide comparisons and performed multi-signal analysis in a prespecified high-information subset. Second, we reconstructed matched-input ABF arms to distinguish the effects of variant support and disease statistics from residual differences between single-causal and multi-signal procedures. Third, truth-known simulations separated credible-set discovery from inference conditional on an available signal pair. Finally, representative IL12RB2-NK and FCRL3 B-cell assignments were examined across molecular-QTL resources, a different disease study system and bounded chromatin and liver-expression evidence. Together, these analyses asked which assignments remained stable, which depended on inputs or model specification, and what additional evidence would be needed to infer mechanism.',
 'Replace repeated statements of what the paper does not claim with a positive, question-led research contract; boundaries remain in evidence-specific Results and Discussion.', 'COMPRESS_NAMED_DUPLICATE_ROADMAP')
replace(18,
 'The analysis comprised source validation, harmonization, ABF screening, high-information follow-up, cell-specific LD construction, multi-signal inference, diagnostic review, matched-input comparison, simulations and bounded external and tissue evaluation. The liver analysis preceded the PBC-wide design; signal-semantics and matched-input comparisons were post-result audits. Negative, model-sensitive, coverage-limited and uninformative results were retained.',
 'Keep actual timing and prespecification distinctions while replacing editorial architecture commentary with factual study design.')
replace(64,
 'Molecule representatives were counted from records carrying the 10x xf bit 8 flag, a corrected cell barcode and a unique gene assignment; unmapped, secondary and supplementary alignments were excluded. Cell calling used the library-specific order-of-magnitude rule with an expected recovered-cell count of 6,000: the UMI threshold was the ceiling of one tenth of the 99th-percentile count among the 6,000 highest-count barcodes. Called-cell counts had to fall between 2,500 and 7,500, with median total UMI count at least 500. B markers were CD79A, CD79B, MS4A1, CD37, CD74, HLA-DRA, CD19 and CD22; NK markers were NKG7, GNLY, KLRD1, PRF1, GZMB, XCL1 and XCL2. T markers were CD3D, CD3E, TRBC1, TRBC2 and IL7R, and myeloid markers were LST1, TYROBP, FCER1G, CTSS, AIF1 and LYZ. Both lineage rules required at least two detected panel genes and three panel UMIs. B-panel UMIs had to exceed NK, T and myeloid panel counts; NK-panel UMIs had to exceed B and myeloid counts and be at least the T-panel count. Each donor-lineage unit required at least ten cells. Target detectability required at least three target UMIs and two target-positive lineage cells. Target CPM was target UMIs divided by total lineage UMIs multiplied by 10^6.',
 'Expand reproducibility-critical cell-calling and lineage definitions from the frozen, actually executed v2 counting code and v3.2 wrapper; no threshold is changed.', 'EXPAND_REPRODUCIBILITY')
replace(67,
 'External inputs were recorded with accession, URL, byte count and available MD5 or locally calculated SHA-256 values. Large third-party files and donor-level genotype data were not mirrored publicly. Code, aggregate results and source manifests are available in the project repository [21]. Source tables provide the reported estimates and analysis identities, including uncertain and reviewed diagnostic outcomes.',
 'Remove editorial claim-ledger workflow from author methodology while retain source traceability and reproducibility.')
replace(73,
 'The source-identity audit verified 14 comparison-model records from seven targeted comparisons under PF10 and PF50. This counting unit differs from the 14 cell types in the expanded registry. Kriging diagnostics were reviewed for all 2,568 comparison-model-trait units. One unique variant satisfied the stated diagnostic flag rule; it lay outside all recorded credible sets and was retained without a diagnostic-triggered refit (Figure 1D).',
 'Report the actual audit object and reviewed anomaly rather than treating a passed gate as absence of anomaly.')
replace(76,orig[76].replace('changed another 19 states','changed 19 states').replace('Input coordination therefore mattered, but it did not erase the high-information shared-signal screen.','The two change counts can involve the same comparisons and are not additive.'),
 'Prevent the adjacent-arm counts from being read as 34 unique changes; preserve exact 15/19 counts.')
replace(78,orig[78].replace('This prevented preservation of a historical narrative from overriding the matched-input definition.','Strict reversal counts therefore follow the matched-input starting state.'),
 'Replace internal narrative-governance comment with the scientific counting rule.')
replace(79,orig[79].replace('PF50 concordance was therefore high among definite PF10 states, but it was not 491 of all 642 comparisons.','The 94.42% agreement is conditional on a definite PF10 state.'),
 'State the denominator directly and avoid repeated defensive contrast.')
replace(103,orig[103].replace('the effect of representing several causal components.','the effect of representing several causal components [5-7,13].'),
 'Anchor the methodological interpretation to existing primary colocalization/fine-mapping references.')
replace(110,orig[110].replace('Genetic regulatory effects need not produce a large average case-control expression difference in an affected organ, particularly across heterogeneous cell states and disease stages.','A genotype-expression association within a cell context and an average expression contrast between case and control donors estimate different relationships. The latter also combines cell-state, disease-stage and surgical-control differences.'),
 'Explain why genetic sharing and a non-significant tissue group contrast need not conflict, without implying an untested disease mechanism.')
replace(113,
 'Three tests would resolve distinct remaining uncertainties. First, source-matched multi-signal analysis in a larger external PBC GWAS with documented participant separation would test whether the same disease-QTL signal pairs persist under matched ancestry and variant support. Second, a direct eQTL-caQTL analysis with common variants and source-matched LD would test whether the chromatin and expression associations share a signal beyond their separate disease-sharing edges. Third, perturbing IL12RB2 in NK cells and FCRL3 in B-cell subsets, with non-targeting controls, would test whether altered expression changes immune activation and cholangiocyte injury in a defined coculture model. Reproduced signal identity would strengthen cross-study regulatory association; a perturbation response would address mechanism, and neither would alone establish clinical efficacy.',
 'Compile uncertainties into specific study/comparator/readout/decision tests while maintain hypothetical status.', 'REWRITE_SAME_JOB')
replace(122,'This research received no specific grant from any funding agency in the public, commercial or not-for-profit sectors.',
 'Remove redundant no-funder-role sentence; retain confirmed no-funding status.', 'COMPRESS_NAMED_REDUNDANCY')
replace(124,orig[124].replace('Both authors interpreted the results and approved the final manuscript.','Both authors interpreted the results.'),
 'Do not assert fresh approval of this edited version before the authors have reconfirmed it.')
caption_additions={
 157:'; PBC, primary biliary cholangitis; PF50, model with 50 expression factors',
 161:'; L, maximum number of SuSiE effects; PF10/PF50, models with ten/50 expression factors',
 165:'; QTL, quantitative-trait locus',
 169:'; NK, natural killer cells; PP.H4, posterior probability of the shared-signal hypothesis',
 173:'; A0, historical ABF; A1, support-set-matched ABF; A2, support- and disease-statistic-matched ABF; M, multi-signal analysis; PF10/PF50, models with ten/50 expression factors'}
for i,suffix in caption_additions.items():
    replace(i,orig[i].rstrip('.')+suffix+'.','Define abbreviations actually present in the associated figure, for independent legend interpretability.', 'EXPAND_LEGEND')
replace(168,orig[168]+' Arrows denote reported association or peak-gene links, not the direction of a causal effect.',
 'Make the chromatin triangle visual semantics explicit without claiming mediation.', 'EXPAND_LEGEND')
replace(176,orig[176]+' Boxes show the interquartile range, centre lines the median, and whiskers extend to the most extreme observations within 1.5 times the interquartile range. All donor observations are overlaid.',
 'Define the plotted box/whisker geometry from the frozen seaborn boxplot code; these are distribution summaries, not confidence intervals.', 'EXPAND_LEGEND')
for i,label in {71:'Source Data 1, S2/',77:'Source Data 1, S3/',85:'Source Data 1, S7/',90:'Source Data 2, data/',99:'Supplementary Data 1, worksheet S9; Source Data 1, S9/'}.items():
    replace(i,base.paragraphs[i].text+f' Complete records are provided in {label}.',
            'Add a local path to exhaustive supporting records without duplicating the table contents.', 'INSERT_SOURCE_POINTER')

# Website/source portals become parenthetical URLs; scholarly and DOI datasets remain numbered.
resource_urls={8:'https://www.ebi.ac.uk/gwas/studies/GCST90061440',10:'https://www.staff.ncl.ac.uk/heather.cordell/GjokaPaper.html',16:'https://r13.finngen.fi/pheno/CHIRBIL_PRIM',17:'https://ngdc.cncb.ac.cn/omix/release/OMIX001122',20:'https://ngdc.cncb.ac.cn/gsa-human/browse/HRA008003',21:'https://github.com/1209433622cz-maker/igan-open-regulatory-evidence/tree/r7c2-q1-journal-audit-dual-track-2026-10-11'}
def nums(token):
    result=[]
    for bit in token.split(','):
        if re.search(r'[-–]',bit):
            a,b=map(int,re.split(r'[-–]',bit));result.extend(range(a,b+1))
        else:result.append(int(bit))
    return result
def move_web(m):
    ids=nums(m.group(1));kept=[x for x in ids if x not in resource_urls];urls=[resource_urls[x] for x in ids if x in resource_urls]
    return (('['+','.join(map(str,kept))+']') if kept else '')+((' ('+'; '.join(urls)+')') if urls else '')
refs={int(re.match(r'^(\d+)\.',p.text)[1]):p for p in base.paragraphs if re.match(r'^\d+\. ',p.text)}
for p in base.paragraphs:
    if p not in refs.values() and re.search(r'\[\d',p.text):
        p.text=re.sub(r'\[([\d,\s\-–]+)\]',move_web,p.text)
        for url in resource_urls.values():
            p.text=re.sub(re.escape(url)+r'\s+\('+re.escape(url)+r'\)',lambda m:url,p.text)

# Keep science architecture; only move the intact Methods block into npj Article order.
methods=base.paragraphs[14:69]
anchor=base.paragraphs[115]._p
for p in methods:anchor.addprevious(p._p)
edits.append({'source_paragraph':'14-68','operation':'MOVE_INTACT','reason':'Official npj Article ordering: Introduction, Results, Discussion, Methods. Methods content retained.','before':'Introduction / Methods / Results / Discussion','after':'Introduction / Results / Discussion / Methods','numeric_source_changed':'NO'})

body=[]
for p in base.paragraphs:
    if p.text=='References':break
    body.append(p)
order=[]
for p in body:
    for m in re.finditer(r'\[([\d,\s\-–]+)\]',p.text):
        for i in nums(m[1]):
            if i not in order:order.append(i)
assert set(order)==set(refs)-set(resource_urls),(order,set(refs)-set(resource_urls))
mapping={old:new for new,old in enumerate(order,1)}
for p in body:
    p.text=re.sub(r'\[([\d,\s\-–]+)\]',lambda m:'['+','.join(map(str,sorted(mapping[x] for x in nums(m[1]))))+']',p.text)

metadata=json.loads((SRC/'crossref_reference_metadata.json').read_text('utf-8'))
bydoi={x['doi']:x['metadata'] for x in metadata if x['status']=='PASS'}
reference_rows=[]
journal_abbreviations={1:'Nat. Rev. Gastroenterol. Hepatol.',2:'J. Hepatol.',3:'Science',5:'PLoS Genet.',6:'J. R. Stat. Soc. B',7:'PLoS Genet.',9:'Genet. Epidemiol.',11:'Nat. Commun.',13:'PLoS Genet.',15:'Nature',18:'Nat. Commun.',19:'Nat. Commun.',22:'npj Syst. Biol. Appl.',23:'Medicine',24:'BMC Gastroenterol.',25:'PLoS One'}
def author_format(auth):
    initials=' '.join(v[0]+'.' for v in re.findall(r'[A-Za-z]+',auth.get('given','')))
    return auth.get('family','')+', '+initials
def nature_reference(old):
    text=orig[127+old]
    if old==12:return 'Xue, A. et al. OneK1K cell-specific cis-eQTL summary statistics, genotypes and covariates. Zenodo https://doi.org/10.5281/zenodo.18910121 (2026).'
    if old==14:return 'Cuomo, A. S. E. et al. TenK10K phase 1 molecular-QTL and fine-mapping resources. Zenodo https://doi.org/10.5281/zenodo.18221260 (2026).'
    dm=re.search(r'doi:([^\s]+)',text)
    if dm:
        doi=dm[1].rstrip('.');meta=bydoi.get(doi)
        if meta:
            au=meta['author'];authors=author_format(au[0])+' et al.' if len(au)>5 else (' & '.join([', '.join(author_format(x) for x in au[:-1]),author_format(au[-1])]) if len(au)>1 else author_format(au[0]))
            title=meta['title'][0].rstrip('.')
            journal=journal_abbreviations.get(old,(meta.get('short-container-title') or meta.get('container-title') or [''])[0])
            # Use the volume's print year when present, rather than its earlier online date.
            date=meta.get('published-print') or meta.get('published-online') or meta.get('published') or meta.get('issued')
            year=date['date-parts'][0][0]
            if '10.1101/' in doi:return f'{authors} {title}. Preprint at https://doi.org/{doi} ({year}).'
            volume=meta.get('volume','');page=meta.get('page',meta.get('article-number',''))
            return f'{authors} {title}. {journal} '+(f'{volume}, ' if volume else '')+(f'{page} ' if page else '')+f'({year}). https://doi.org/{doi}.'
        # Metadata retrieval failed: preserve verified source fields rather than invent authors.
        if old==3:return 'Yazar, S. et al. Single-cell eQTL mapping identifies cell type-specific genetic control of autoimmune disease. Science 376, eabf3041 (2022). https://doi.org/10.1126/science.abf3041.'
        if old==18:return 'Li, X. et al. Unique DUOX2+ACE2+ small cholangiocytes are pathogenic targets for primary biliary cholangitis. Nat. Commun. 14, 29 (2023). https://doi.org/10.1038/s41467-022-34606-w.'
        return re.sub(r'^\d+\. ','',text)
    return re.sub(r'^\d+\. ','',text)
refhead=next(p for p in base.paragraphs if p.text=='References')
for p in refs.values():p._p.getparent().remove(p._p)
refanchor=refhead._p
for old in order:
    p=OxmlElement('w:p');refanchor.addnext(p);refanchor=p
    from docx.text.paragraph import Paragraph
    par=Paragraph(p,base._body);par.style=base.styles['Normal'];par.text=f'{mapping[old]}. '+nature_reference(old)
    reference_rows.append({'old_number':old,'new_number':mapping[old],'action':'RENORMALIZED_NATURE_STYLE','identity':orig[127+old]})
for old,url in resource_urls.items():reference_rows.append({'old_number':old,'new_number':'URL','action':'RELOCATED_PARENTHESES_NO_SOURCE_DELETION','identity':url})
tsv('Citation_Source_Conservation.tsv',reference_rows)

# Lowercase panel citations match the production derivatives.
for p in base.paragraphs:
    if not p._p.xpath('.//w:drawing'):
        p.text=re.sub(r'Figure (\d)([A-F])',lambda m:f'Figure {m[1]}{m[2].lower()}',p.text)

# Main submission only points to SI; full table dictionary is retained in SI.
tail=next(i for i,p in enumerate(base.paragraphs) if p.text=='Supplementary information')
for p in list(base.paragraphs)[tail+1:]:p._p.getparent().remove(p._p)
base.add_paragraph('Supplementary Information provides the data dictionary and file map. Supplementary Data 1 contains ten worksheet views (S1-S10). Source Data 1 supplies the complete primary machine-readable tables, including the full screening, signal-pair and simulation records; Source Data 2 supplies the external disease-study and chromatin evidence family (S11). These files contain aggregate, shareable outputs and do not redistribute donor genotypes or source BAMs.')
edits.append({'source_paragraph':'179-193','operation':'MOVE_DUPLICATE_DICTIONARY_TO_EXISTING_SI','reason':'Table legends duplicated the separate SI with the same job; retain all S1-S11 definitions and source paths there, with a concise main-file pointer.','before':'Three file descriptions and eleven full table legends','after':'One file-map paragraph; complete SI dictionary retained','numeric_source_changed':'NO'})

# Update embedded media without changing drawing sizes or object count.
pics=[p for p in base.paragraphs if p._p.xpath('.//w:drawing')]
assert len(pics)==6
for num,p in enumerate(pics,1):
    blip=p._p.xpath('.//a:blip')[0];rid=blip.get(qn('r:embed'));base.part.related_parts[rid]._blob=(OUT/f'02_Figures/Figure_{num}_embed.png').read_bytes()
    docpr=p._p.xpath('.//wp:docPr')[0];docpr.set('descr',f'Figure {num}: source-derived statistical panels; see following legend for units, methods and boundaries.')
    p.paragraph_format.keep_with_next=True
    p.paragraph_format.page_break_before=True
    p.paragraph_format.keep_together=True

# Caption surface and controlled format. Title/house style not applied to scientific truth.
for p in base.paragraphs:
    p.paragraph_format.first_line_indent=Pt(0)
    if p.style.name=='Normal':p.alignment=WD_ALIGN_PARAGRAPH.LEFT
    if p.style.name=='Figure Caption':
        p.paragraph_format.keep_with_next=not p.text.startswith('Abbreviations:')
        if p.text.startswith('Figure '):
            text=p.text;p.text='';split=text.find('. ',text.find('. ')+2)+1
            p.add_run(text[:split]).bold=True;p.add_run(text[split:])
    if p.style.name in ['Heading 1','Heading 2','Title']:
        p.style.font.color.rgb=None
    if p.text=='Ethics statement':p.style=base.styles['Heading 2']

# Bibliography typography is applied after text replacements, so no later p.text reset erases it.
for p in base.paragraphs:
    m=re.match(r'^(\d+)\. ',p.text)
    if not m:continue
    old=next(k for k,v in mapping.items() if v==int(m[1]))
    journal=journal_abbreviations.get(old)
    if journal and journal in p.text:
        before,after=p.text.split(journal,1);p.text='';p.add_run(before)
        p.add_run(journal).italic=True
        vm=re.match(r' (\d+)(,.*)',after)
        if vm:p.add_run(' ');p.add_run(vm[1]).bold=True;p.add_run(vm[2])
        else:p.add_run(after)

figure_heading=next(p for p in base.paragraphs if p.text=='Figure legends')
figure_heading.paragraph_format.page_break_before=True
figure_heading.paragraph_format.keep_with_next=True
pics[0].paragraph_format.page_break_before=False

# Two interface copies share text, references and legends; only figure placement differs.
end=OUT/'01_Manuscript/R7C3B_npj_Manuscript.docx'
base.save(end)
inline=Document(end)
blocks=[]
paras=inline.paragraphs
for i,p in enumerate(paras):
    if p._p.xpath('.//w:drawing'):
        blocks.append([p._p,paras[i+1]._p,paras[i+2]._p])
assert len(blocks)==6
landings=['The only unique', 'H3 support identified', 'We therefore did not generalize S6','descriptive detectability rather than disease enrichment','not evidence that FCRL3 has no role','The tissue layer therefore constrained']
# Identify section-end landings by stable content, not page indices.
landings=[
    next(p for p in inline.paragraphs if p.text.startswith('The source-identity audit verified')),
    next(p for p in inline.paragraphs if p.text.startswith('Signal-pair semantics further')),
    next(p for p in inline.paragraphs if p.text.startswith('The decrease in false H4')),
    next(p for p in inline.paragraphs if p.text.startswith('The aligned PBC risk allele')),
    next(p for p in inline.paragraphs if p.text.startswith('FCRL3-CD8_ET provided')),
    next(p for p in inline.paragraphs if p.text.startswith('These results established'))]
for block,landing in zip(blocks,landings):
    prev=landing._p
    for e in block:prev.addnext(e);prev=e
for p in list(inline.paragraphs):
    if p.text=='Figure legends':p._p.getparent().remove(p._p)
    elif p._p.xpath('.//w:drawing'):p.paragraph_format.page_break_before=True
# Remove six empty figure-block spacer paragraphs left behind after moving the objects.
after_refs=False
for p in list(inline.paragraphs):
    if p.text=='References':after_refs=True
    elif after_refs and not p.text.strip() and not p._p.xpath('.//w:drawing | .//w:sectPr | .//w:fldChar'):
        p._p.getparent().remove(p._p)
inline.save(OUT/'01_Manuscript/R7C3B_AuthorReview_InlineFigures.docx')

# SI gains precise file ownership, not new scientific tables.
si=Document(SRC/'03_Supplement/R7C3A_npj_Supplementary_Information_Refined.docx')
si.paragraphs[0].text='Supplementary Information'
for p in si.paragraphs:
    p.text=p.text.replace('Source_Data_1_R7C3A_Editorial_Crosswalk.zip','Source_Data_1_Primary_Machine_Readable_Data.zip').replace('Source_Data_2_R7C3A_Editorial_Crosswalk.zip','Source_Data_2_External_Regulatory_Evidence.zip')
    p.text=p.text.replace('Supplementary Table S','Evidence family S').replace('Keep the original screening labels separate from matched-input classifications.','Original screening labels and matched-input classifications are recorded separately.')
si.paragraphs[3].text='This document defines the units, provenance and interpretation limits for the S1-S11 evidence families. Supplementary Data 1 contains ten worksheet views (S1-S10); complete row-level tables are supplied in Source Data 1. The S11 external-evidence family is supplied in Source Data 2, rather than an eleventh worksheet. Worksheet and archive labels are cross-referenced below.'
for num in range(1,11):
    i=10+(num-1)*2
    si.paragraphs[i].text += f' File map: Supplementary Data 1, worksheet S{num}; Source Data 1, S{num}/.'
si.paragraphs[30].text += ' File map: Source Data 2, data/; the constituent files retain source-stage names for provenance.'
si.save(OUT/'03_Supplement/R7C3B_Supplementary_Information.docx')
shutil.copy2(SRC/'03_Supplement/Supplementary_Data_1_Tables_S1-S10.xlsx',OUT/'03_Supplement/Supplementary_Data_1_Tables_S1-S10.xlsx')

# Source Data result members retain exact bytes; README/checksum are production metadata.
source_ledger=list(csv.DictReader((SRC/'06_QA/Scientific_Source_Data_Identity_Ledger.tsv').open(encoding='utf-8-sig'),delimiter='\t'))
sciencerows=[]
for num in [1,2]:
    src=SRC/f'05_Source_Data/Source_Data_{num}_R7C3A_Editorial_Crosswalk.zip'
    dst=OUT/'05_Source_Data'/('Source_Data_1_Primary_Machine_Readable_Data.zip' if num==1 else 'Source_Data_2_External_Regulatory_Evidence.zip')
    with zipfile.ZipFile(src) as z:members={n:z.read(n) for n in z.namelist()}
    # Preserve old crosswalk as source provenance, but make the active readme stage-neutral.
    readme='README.txt' if num==1 else 'README.md'
    members[readme]=(
      f'# Source Data {num} for npj Genomic Medicine\n\n'
      + ('Complete primary aggregate evidence in S1/ through S10/. The separate editable workbook contains ten worksheet views; it is not a substitute for the complete large tables.\n' if num==1 else 'Bounded external disease-study, CASCADE, chromatin and descriptive OMIX evidence (S11 family) in data/.\n')
      + '\nScientific data members retain the frozen source-release bytes. Historical filenames and source-state files document provenance rather than additional analyses. See the submitted Supplementary Information for units and interpretation limits.\n'
      + '\nNo source BAMs, donor genotypes or source pseudobulk matrices are redistributed. Shared genetic signals do not establish mediation or clinical utility.\n').encode()
    # Replace only the publication-level checksum metadata; nested data manifests remain evidence.
    for n in list(members):
        if n.startswith('checksums'):del members[n]
    members['checksums.sha256']=''.join(f'{h(v)}  {n}\n' for n,v in sorted(members.items())).encode()
    with zipfile.ZipFile(dst,'w',zipfile.ZIP_DEFLATED,compresslevel=6) as z:
        for n,b in sorted(members.items()):z.writestr(n,b)
    for r in source_ledger:
        if (num==1 and r['archive']=='S1-10') or (num==2 and r['archive']=='S11-external'):
            assert h(members[r['member']])==r['science_sha256'],r
            sciencerows.append({**r,'new_archive':dst.name,'verified_sha256':h(members[r['member']]),'independent_status':'PASS'})
assert len(sciencerows)==48,len(sciencerows)
tsv('Scientific_Source_Data_48_Identity.tsv',sciencerows)

cover=Document(SRC/'04_Cover/01_Cover_Letter.docx')
cover.paragraphs[10].text=cover.paragraphs[10].text.replace('All authors approved the manuscript and its submission. ','')
cover.paragraphs[8].text=cover.paragraphs[8].text.replace('robust regulatory association','regulatory association supported across the tested resources and configurations')
cover.styles['Normal'].font.size=Pt(11)
closing=next(i for i,p in enumerate(cover.paragraphs) if p.text=='Sincerely,')
for p in list(cover.paragraphs)[closing+1:]:p._p.getparent().remove(p._p)
cover.add_paragraph('Zhi Chen (first author); zhichen1@link.cuhk.edu.cn; ORCID: 0009-0001-0072-5576')
cover.add_paragraph('Teng Qi (corresponding author); tengqi@link.cuhk.edu.cn; ORCID: 0009-0007-7648-4776')
cover.add_paragraph('School of Medicine, The Chinese University of Hong Kong, Shenzhen')
cover.save(OUT/'04_Cover/R7C3B_Cover_Letter_Candidate.docx')
edits.append({'source_paragraph':'publication interface','operation':'FORMAT_ONLY','reason':'Nature bibliography typography and volume-year convention; duplicate URL removed; empty inline figure spacers removed; cover-letter body enlarged from 9.5 to 11 pt.','before':'Inherited production formatting','after':'Purpose-matched review and submission formatting','numeric_source_changed':'NO'})
tsv('Bounded_Edit_Ledger.tsv',edits)
dump('Reference_Number_Mapping.json',mapping)
dump('Format_Contract.json',{
 'skill':'QiTeng v0.3.24.2','purpose_primary':'TARGET_JOURNAL_SUBMISSION_CANDIDATE','purpose_inline':'AUTHOR_REVIEW',
 'journal':'npj Genomic Medicine','article_type':'Article','scientific_architecture':'retained; Methods relocated intact into journal Article order',
 'body':'Times New Roman 11 pt; double spacing; left aligned; inherited A4 margins',
 'captions':'Times New Roman 9 pt; title sentence bold; abbreviations separate; same-page figure/legend',
 'citations':'numbered bracket identifiers, globally renumbered first use; Nature-formatted bibliography; resource URLs parenthetical',
 'panels':'lowercase production derivatives; plots outside label masks pixel-identical to source',
 'renderer':'WPS kwps.Application, hidden, read-only DOCX export','global_shortening_target':False,
 'source_tables_changed':False,'author_approval':'PENDING_FOR_R7C3B','publisher_proof':'PENDING'})
print('BUILT',OUT,'abstract words',len(base.paragraphs[6].text.split()),'references',len(order),'scientific identities',len(sciencerows))
