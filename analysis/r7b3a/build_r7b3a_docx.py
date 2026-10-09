#!/usr/bin/env python3
"""Create the source-corrected R7B3A author-review DOCX."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(r"H:\SCI2\YR1")
BASE = ROOT / "5_manuscript" / "R7B3A_ManuscriptV2"
MD = BASE / "manuscript" / "R7B3A_full_english_manuscript_v2.md"
OUT = BASE / "manuscript" / "R7B3A_full_english_manuscript_v2.docx"
FIG = BASE / "figures"


def set_cell_margins(cell, top=90, start=90, bottom=90, end=90):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar"); tc_pr.append(tc_mar)
    for key, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{key}"))
        if node is None:
            node = OxmlElement(f"w:{key}"); tc_mar.append(node)
        node.set(qn("w:w"), str(value)); node.set(qn("w:type"), "dxa")


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(); field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE"); run._r.addnext(field)


def add_line_numbering(section):
    sect_pr = section._sectPr
    node = sect_pr.find(qn("w:lnNumType"))
    if node is None:
        node = OxmlElement("w:lnNumType"); sect_pr.append(node)
    node.set(qn("w:countBy"), "1"); node.set(qn("w:restart"), "continuous"); node.set(qn("w:distance"), "360")


def add_rich_text(paragraph, text: str):
    parts = re.split(r"(\*\*.*?\*\*|`.*?`|\*.*?\*)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2]); run.bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1]); run.font.name = "Courier New"; run.font.size = Pt(9)
        elif part.startswith("*") and part.endswith("*"):
            run = paragraph.add_run(part[1:-1]); run.italic = True
        else:
            paragraph.add_run(part)


def figure_path(number: int) -> Path:
    names = {
        1: "Figure1_study_scope_and_model_integrity.png",
        2: "Figure2_input_matched_reclassification.png",
        3: "Figure3_simulation_discovery_and_inference.png",
        4: "Figure4_IL12RB2_external_evidence.png",
        5: "Figure5_FCRL3_support_and_counterexample.png",
        6: "Figure6_liver_tissue_boundary.png",
    }
    path = FIG / names[number]
    if not path.exists():
        raise FileNotFoundError(path)
    return path


def configure_styles(doc: Document):
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"; normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(11); normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    normal.paragraph_format.space_after = Pt(0); normal.paragraph_format.first_line_indent = Cm(.74)
    normal.paragraph_format.widow_control = True
    for name, size, color in [("Title", 17, "17365D"), ("Heading 1", 14, "17365D"), ("Heading 2", 12, "1F4E79"), ("Heading 3", 11, "1F4E79")]:
        style = doc.styles[name]
        style.font.name = "Arial"; style._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
        style.font.size = Pt(size); style.font.bold = True; style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.keep_with_next = True; style.paragraph_format.space_before = Pt(12 if name != "Title" else 0)
        style.paragraph_format.space_after = Pt(6); style.paragraph_format.line_spacing = 1.05
    if "Figure Caption" not in doc.styles:
        caption = doc.styles.add_style("Figure Caption", WD_STYLE_TYPE.PARAGRAPH)
    else:
        caption = doc.styles["Figure Caption"]
    caption.font.name = "Times New Roman"; caption.font.size = Pt(9)
    caption.paragraph_format.line_spacing = 1.0; caption.paragraph_format.space_after = Pt(6)
    caption.paragraph_format.first_line_indent = Cm(0)


def main() -> None:
    text = MD.read_text(encoding="utf-8")
    doc = Document(); section = doc.sections[0]
    section.top_margin = Cm(2.2); section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.5); section.right_margin = Cm(2.2)
    section.page_width = Cm(21); section.page_height = Cm(29.7)
    section.header_distance = Cm(1); section.footer_distance = Cm(1)
    add_line_numbering(section); add_page_number(section.footer.paragraphs[0]); configure_styles(doc)
    props = doc.core_properties
    props.title = "Input-matched multi-signal analysis clarifies immune-cell regulatory assignments at PBC risk loci"
    props.subject = "R7B3A source-corrected manuscript v2 with frozen figures and supplements"
    props.author = "[AUTHORS TO BE COMPLETED]"
    props.keywords = "PBC; colocalization; eQTL; SuSiE; source-matched LD; multi-signal"

    in_figure_legends = False
    for line in text.splitlines():
        raw = line.rstrip()
        if not raw:
            continue
        if raw.startswith("# "):
            paragraph = doc.add_paragraph(style="Title"); paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.first_line_indent = Cm(0); add_rich_text(paragraph, raw[2:]); continue
        if raw == "## Figure legends":
            doc.add_page_break(); doc.add_heading("Figure legends", level=1); in_figure_legends = True; continue
        if raw.startswith("## "):
            doc.add_heading(raw[3:], level=1); continue
        if raw.startswith("### "):
            doc.add_heading(raw[4:], level=2); continue
        if in_figure_legends and raw.startswith("**Figure "):
            match = re.match(r"\*\*Figure (\d+)\.", raw)
            if match:
                number = int(match.group(1))
                if number > 1: doc.add_page_break()
                image_paragraph = doc.add_paragraph(); image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                image_paragraph.paragraph_format.first_line_indent = Cm(0)
                image_paragraph.add_run().add_picture(str(figure_path(number)), width=Inches(6.45))
                caption = doc.add_paragraph(style="Figure Caption"); add_rich_text(caption, raw); continue
        if in_figure_legends and raw.startswith("**Abbreviations:**"):
            paragraph = doc.add_paragraph(style="Figure Caption"); add_rich_text(paragraph, raw); continue
        paragraph = doc.add_paragraph()
        if raw.startswith(("**Authors:", "**Affiliations:", "**Corresponding author:", "**Article type:")):
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER; paragraph.paragraph_format.first_line_indent = Cm(0)
            paragraph.paragraph_format.line_spacing = 1.15
        if raw.startswith("G* =") or raw.startswith("and R_QTL"):
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER; paragraph.paragraph_format.first_line_indent = Cm(0)
        add_rich_text(paragraph, raw.rstrip("  "))

    doc.add_page_break(); doc.add_heading("Author metadata completion gate", level=1)
    table = doc.add_table(rows=1, cols=2); table.style = "Table Grid"
    table.rows[0].cells[0].text = "Field"; table.rows[0].cells[1].text = "Required verified information"
    for field, need in [
        ("Authors and order", "Full names, order and equal-contribution marks"),
        ("Affiliations", "Numbered official institutional addresses"),
        ("Correspondence", "Name, address, email and telephone"),
        ("CRediT", "Author-specific contribution roles"),
        ("Funding", "Exact funder and grant numbers, or verified no-funding statement"),
        ("Competing interests", "Verified author declaration"),
        ("Ethics", "Local secondary-analysis determination or waiver wording"),
    ]:
        cells = table.add_row().cells; cells[0].text = field; cells[1].text = need
        for cell in cells: set_cell_margins(cell)
    for cell in table.rows[0].cells:
        set_cell_margins(cell)
        for run in cell.paragraphs[0].runs: run.bold = True

    OUT.parent.mkdir(parents=True, exist_ok=True); doc.save(OUT); print(OUT)


if __name__ == "__main__":
    main()
