#!/usr/bin/env python3
"""Final machine QA for the R7B3A manuscript, figures and supplements."""

from __future__ import annotations

import hashlib
import json
import re
import sys
from pathlib import Path

import pandas as pd
from docx import Document
from pypdf import PdfReader


ROOT = Path(r"H:\SCI2\YR1")
BASE = ROOT / "5_manuscript/R7B3A_ManuscriptV2"
MD = BASE / "manuscript/R7B3A_full_english_manuscript_v2.md"
DOCX = BASE / "manuscript/R7B3A_full_english_manuscript_v2.docx"
PDF = BASE / "manuscript/R7B3A_full_english_manuscript_v2_WPS.pdf"
FIG = BASE / "figures"
SUP = BASE / "supplements"
LEDGER = ROOT / "3_results/04_integration/R7B2A/manuscript_lock/R7B2A_updated_claim_evidence_ledger.tsv"
TRAJ = ROOT / "3_results/04_integration/R7B2A/adjudication/R7B2A_642_four_arm_trajectories.tsv"
MULTI = ROOT / "3_results/04_integration/R7B1B_v2/adjudication/R7B1B_v2_642_bidirectional_reclassification.tsv"
OUT = BASE / "qa"; OUT.mkdir(parents=True, exist_ok=True)


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def main() -> int:
    text = MD.read_text(encoding="utf-8")
    checks: list[dict[str, str]] = []

    def check(name: str, passed: bool, detail) -> None:
        checks.append({"check": name, "status": "PASS" if passed else "FAIL", "detail": str(detail)})

    required_sections = ["## Abstract", "## Background", "## Methods", "## Results", "## Discussion", "## Conclusions", "## Declarations", "## References", "## Figure legends", "## Supplementary materials"]
    check("required_sections", all(s in text for s in required_sections), required_sections)

    anchors = ["6,923", "5,460", "642", "113", "452", "79", "413", "eight H4-to-H3", "ten H3-to-H4", "92 H4-supported", "428 H3-supported", "59 model-sensitive", "63 uninformative", "32 H4-supported", "491", "520", "486,000", "160,782", "33.08%", "12.41%", "8.56%", "32.99%", "26.66%", "0.9698", "0.9706", "0.9916", "q=0.111"]
    missing = [value for value in anchors if value not in text]
    check("numeric_anchor_coverage", not missing, f"missing={missing}")

    prohibited = ["proved that multi-signal", "proves that multi-signal", "generally superior", "established a complete causal cascade", "tissue mediation was established", "PBC-specific upregulation", "all 6,923 received multi-signal"]
    found = [phrase for phrase in prohibited if phrase.lower() in text.lower()]
    check("prohibited_positive_claims", not found, f"found={found}")
    check("independent_disease_replication_bounded", "not independent disease replication" in text.lower() or "rather than independent disease replication" in text.lower(), "explicitly bounded")

    ledger = pd.read_csv(LEDGER, sep="\t")
    check("claim_ledger_count", len(ledger) == 27, len(ledger))
    claim_terms = {"C01":"6,923", "C02":"642-comparison", "C04":"79 of 113", "C05":"413 retained H3", "C06":"491 of 520", "C07":"all 92", "C08":"32 of the 92", "C09":"2,568", "C10":"rs1800378", "C11":"12.41%", "C12":"32.99%", "C15":"0.9698", "C18":"0.9916", "C20":"FCRL3-CD8_ET", "C22":"all five PBC", "C27":"no recorded technical errors"}
    uncovered = [cid for cid, term in claim_terms.items() if term.lower() not in text.lower()]
    check("claim_text_coverage", not uncovered, uncovered)

    refs = [int(x) for x in re.findall(r"(?m)^(\d+)\. ", text)]
    check("reference_count_and_continuity", refs == list(range(1, 23)), refs)
    ref_state = json.loads((OUT / "R7B3A_reference_bibliographic_verification_state.json").read_text(encoding="utf-8"))
    check("bibliographic_verification", ref_state.get("status") == "PASS" and ref_state.get("references") == 22, ref_state)

    legends = re.findall(r"(?m)^\*\*Figure (\d+)\.", text)
    check("figure_legends", legends == ["1", "2", "3", "4", "5", "6"], legends)
    fig_files = [p for p in FIG.glob("Figure[1-6]_*") if p.suffix in {".png", ".pdf", ".svg"}]
    check("figure_triplets", len(fig_files) == 18, len(fig_files))
    hashes = json.loads((FIG / "FIGURE_HASHES.json").read_text(encoding="utf-8"))
    check("figure_hash_manifest", len(hashes) == 18 and all(sha256(FIG / name) == digest for name, digest in hashes.items()), len(hashes))

    code = (ROOT / "2_code/11_manuscript/R7B3A/build_r7b3a_figures.py").read_text(encoding="utf-8")
    check("no_figure2_fallback", "no fallback is permitted" in code and "pd.Series([19" not in code, "fail-closed source check")
    check("deterministic_tissue_points", "deterministic_offsets" in code and "stripplot" not in code, "fixed offsets")
    traj = pd.read_csv(TRAJ, sep="\t"); multi = pd.read_csv(MULTI, sep="\t")
    row = traj.loc[traj.comparison_id == "R7B1_000414"].iloc[0]; row_m = multi.loc[multi.comparison_id == "R7B1_000414"].iloc[0]
    actual = [row.A0_PP_H4, row.A1_PP_H4, row.A2_PP_H4, row.pf10_l10_best_h4, row_m.pf50_l10_best_h4]
    expected = [0.941334682686132, 0.9413349988264992, 0.949260546232612, 0.0036361957364122, 0.991894292309108]
    check("figure5c_exact_values", all(abs(a-b) <= 1e-12 for a,b in zip(actual, expected)), actual)

    state = json.loads((SUP / "SUPPLEMENT_STATE.json").read_text(encoding="utf-8"))
    check("supplement_state", state.get("status") == "PASS" and state.get("supplements") == 10, state)
    dirs = [SUP / f"S{i}" for i in range(1, 11)]
    check("supplement_directories", all(p.is_dir() for p in dirs), len([p for p in dirs if p.is_dir()]))
    controls = {2:6923, 3:642, 4:642, 5:3210, 6:11184, 7:540, 8:21, 9:28, 10:37}
    row_details = {}
    for i, expected_rows in controls.items():
        view = pd.read_csv(SUP / f"S{i}" / f"S{i}_workbook_view.tsv", sep="\t", low_memory=False)
        row_details[f"S{i}"] = len(view)
        check(f"S{i}_workbook_rows", len(view) == expected_rows, len(view))
    xlsx = SUP / "R7B3A_Supplementary_Tables_S1-S10.xlsx"
    check("supplement_workbook", xlsx.exists() and xlsx.stat().st_size > 100000, xlsx.stat().st_size if xlsx.exists() else 0)
    preview_count = len(list((SUP / "workbook_previews").glob("S*.png")))
    check("supplement_sheet_visual_review", preview_count == 10, preview_count)

    words = re.findall(r"\b[A-Za-z0-9][A-Za-z0-9'/-]*\b", text)
    check("manuscript_word_count", 6500 <= len(words) <= 11000, len(words))
    placeholders = re.findall(r"\[[A-Z][A-Z0-9 '\-/,.:]+(?:REQUIRED|COMPLETED|CONFIRMED|INFORMATION|WORDING|AUTHORS|ADDRESSES|NAME|CRediT).*?\]", text)
    check("unknown_metadata_explicit", len(placeholders) >= 5, len(placeholders))

    check("docx_exists", DOCX.exists() and DOCX.stat().st_size > 10000, DOCX.stat().st_size if DOCX.exists() else 0)
    if DOCX.exists():
        doc = Document(DOCX)
        check("docx_figure_objects", len(doc.inline_shapes) == 6, len(doc.inline_shapes))
        heading_text = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
        check("docx_section_parity", all(section[3:] in heading_text for section in required_sections if section not in {"## Abstract"}), len(heading_text))

    check("pdf_exists", PDF.exists() and PDF.stat().st_size > 10000, PDF.stat().st_size if PDF.exists() else 0)
    pdf_meta = {}; pages = 0
    if PDF.exists():
        reader = PdfReader(str(PDF)); pages = len(reader.pages); pdf_meta = dict(reader.metadata or {})
        producer = str(pdf_meta.get("/Producer", "")) + " " + str(pdf_meta.get("/Creator", ""))
        check("wps_pdf_creator", "wps" in producer.lower(), producer)
        check("pdf_page_count", 25 <= pages <= 50, pages)
    rendered = list((OUT / "pdf_pages_wps").glob("page-*.png"))
    check("pdf_page_render_parity", len(rendered) == pages and pages > 0, f"rendered={len(rendered)} pages={pages}")

    status = "PASS" if all(x["status"] == "PASS" for x in checks) else "FAIL"
    result = {
        "stage": "R7B3A_SOURCE_BOUND_COMPOSITES_SUPPLEMENT_AND_RENDER_PARITY",
        "status": status, "checks": checks,
        "summary": {"pass": sum(x["status"] == "PASS" for x in checks), "total": len(checks)},
        "artifacts": {p.name: {"bytes": p.stat().st_size, "sha256": sha256(p)} for p in [MD, DOCX, PDF, xlsx] if p.exists()},
        "pdf_metadata": pdf_meta,
    }
    (OUT / "R7B3A_release_QA.json").write_text(json.dumps(result, indent=2, ensure_ascii=False), encoding="utf-8")
    pd.DataFrame(checks).to_csv(OUT / "R7B3A_release_QA.tsv", sep="\t", index=False)
    print(json.dumps({**result["summary"], "status": status}, ensure_ascii=False))
    return 0 if status == "PASS" else 1


if __name__ == "__main__":
    raise SystemExit(main())
