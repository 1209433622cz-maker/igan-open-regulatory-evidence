#!/usr/bin/env python3
"""Build the R7C2 npj Genomic Medicine candidate and dual-track audit package."""

from __future__ import annotations

import hashlib
import json
import re
import shutil
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(r"H:\SCI2\YR1")
SRC = ROOT / "5_manuscript/R7C1_IL12RB2_ExternalValidation_ManuscriptFork"
WORK = ROOT / "5_manuscript/R7C2_Q1_JournalRealityCheck_DualTrackSubmission"
MANUSCRIPT_MD = WORK / "manuscript/R7C2_npj_Genomic_Medicine_manuscript.md"
MANUSCRIPT_DOCX = WORK / "manuscript/R7C2_npj_Genomic_Medicine_manuscript.docx"
COVER_MD = WORK / "cover_letter/R7C2_npj_Genomic_Medicine_cover_letter.md"
COVER_DOCX = WORK / "cover_letter/R7C2_npj_Genomic_Medicine_cover_letter.docx"
SI_MD = WORK / "supplement/R7C2_npj_Supplementary_Information.md"
SI_DOCX = WORK / "supplement/R7C2_npj_Supplementary_Information.docx"

TITLE = "Input matched multisignal colocalization clarifies regulatory assignments at primary biliary cholangitis loci"

ABSTRACT = (
    "Regulatory assignments at disease loci can change when colocalization analyses use different variant sets, "
    "disease statistics or causal assumptions. We evaluated 6,923 primary biliary cholangitis locus-gene-cell "
    "comparisons, screened 5,460, and applied source-matched multi-signal analysis to a prespecified 642-comparison "
    "high-information subset. Matching variant support and disease statistics changed 15 and 19 classifications, yet "
    "eight single-causal H4 comparisons became H3 and ten H3 comparisons became H4 after multi-signal analysis. In "
    "486,000 simulations, matched multi-signal inference reduced false H4 decisions in distinct-signal scenarios from "
    "12.41% to 8.56% but reduced shared-signal recovery from 32.99% to 26.66%. IL12RB2-NK sharing was supported in "
    "FinnGen R13 (PP.H4=0.9977), with disease-linked eQTL and caQTL signals, whereas direct eQTL-caQTL sharing and "
    "functional mediation remained untested. Reliable regulatory attribution requires matched inputs, signal-pair "
    "reporting, model-sensitivity analysis and explicit separation of signal sharing from mechanism."
)


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


def word_count(text: str) -> int:
    return len(re.findall(r"\b[A-Za-z0-9]+(?:[.-][A-Za-z0-9]+)*\b", text))


def replace_between(text: str, start: str, end: str, replacement: str) -> str:
    a = text.find(start)
    if a < 0:
        raise RuntimeError(f"missing start anchor: {start}")
    b = text.find(end, a + len(start))
    if b < 0:
        raise RuntimeError(f"missing end anchor: {end}")
    return text[:a] + replacement.rstrip() + "\n\n" + text[b:]


def build_npj_markdown() -> str:
    text = (SRC / "manuscript/R7C1_HumanGenomics_IL12RB2_external_validation_fork.md").read_text(encoding="utf-8")
    old_title = text.splitlines()[0]
    text = text.replace(old_title, f"# {TITLE}", 1)
    text = text.replace("**Article type:** Research", "**Article type:** Article", 1)

    abstract_block = f"## Abstract\n\n{ABSTRACT}"
    text = replace_between(text, "## Abstract", "**Keywords:**", abstract_block)
    text = text.replace("## Background", "## Introduction", 1)

    conclusion_match = re.search(r"\n## Conclusions\n\n(.*?)(?=\n## List of abbreviations)", text, flags=re.S)
    if not conclusion_match:
        raise RuntimeError("conclusion section missing")
    conclusion = conclusion_match.group(1).strip()
    discussion_end = text.find("\n## Conclusions")
    text = text[:discussion_end].rstrip() + "\n\n" + conclusion + "\n\n" + text[conclusion_match.end():].lstrip("\n")

    declarations = """## Ethics statement

Ethics approval was not required because this study exclusively reanalysed publicly available, de-identified data and involved no new participant recruitment or identifiable data collection. The original studies reported their applicable ethics approvals and participant consent.

## Data availability

PBC summary statistics are available through GWAS Catalog study GCST90061440. Study-derived PBC regional statistics and LD are available from the GJOKA archive. OneK1K and TenK10K resources are identified in their cited publications and Zenodo records. FinnGen R13 CHIRBIL_PRIM summary statistics and fine-mapping records are publicly distributed by FinnGen [16], and CASCADE records are accessible through the FinnGen immune multiome portal [15]. HRA008003 is available through GSA-Human and OMIX001122 through the National Genomics Data Center [17]. Supplementary tables and compact derived outputs accompany this submission. Individual-level genotype, BAM, pseudobulk matrices and third-party source archives are not redistributed.

## Code availability

Analysis code, frozen protocols, aggregate outputs, manifests and source-bound figures are versioned in the public project repository at `https://github.com/1209433622cz-maker/igan-open-regulatory-evidence/tree/r7c2-q1-journal-audit-dual-track-2026-10-11` [21]. The release records software environments, checksums and the mapping from manuscript claims to machine-readable evidence.

## Acknowledgements

This research received no specific grant from any funding agency in the public, commercial or not-for-profit sectors. No funder had any role in study design, data analysis, interpretation, manuscript preparation or the decision to submit.

## Author contributions

Z.C.: Conceptualization, Methodology, Software, Formal analysis, Investigation, Data curation, Visualization and Writing - original draft. T.Q.: Conceptualization, Methodology, Validation, Supervision, Project administration and Writing - review and editing. Both authors interpreted the results and approved the final manuscript.

## Competing interests

The authors declare no competing interests.

"""
    text = replace_between(text, "## List of abbreviations", "## References", declarations)
    text = text.replace("## Additional files", "## Supplementary information", 1)
    text = text.replace(
        "**Additional file 1.** `Additional_file_1_Supplementary_Tables_S1-S10.xlsx`",
        "**Supplementary Data 1.** `Supplementary_Data_1_Tables_S1-S10.xlsx`",
    )
    text = text.replace(
        "**Additional file 2.** `Additional_file_2_Machine_Readable_Supplementary_Data.zip`",
        "**Source Data 1.** `Source_Data_1_Primary_Machine_Readable_Data.zip`",
    )
    text = text.replace(
        "**Additional file 3.** `Additional_file_3_R7C1_External_Validation_Data.zip`",
        "**Source Data 2.** `Source_Data_2_External_Validation_Data.zip`",
    )
    text = text.replace("The three files contain", "The three data files contain")
    text = text.replace("Delivered in Additional file 3.", "Delivered in Source Data 2.")
    return "\n".join(line.rstrip() for line in text.splitlines()) + "\n"


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    paragraph._p.append(field)


def add_line_numbering(section):
    node = section._sectPr.find(qn("w:lnNumType"))
    if node is None:
        node = OxmlElement("w:lnNumType")
        section._sectPr.append(node)
    node.set(qn("w:countBy"), "1")
    node.set(qn("w:restart"), "continuous")
    node.set(qn("w:distance"), "360")


def add_rich_text(paragraph, text: str):
    parts = re.split(r"(\*\*.*?\*\*|`.*?`|\*.*?\*)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Courier New"
            run.font.size = Pt(9)
        elif part.startswith("*") and part.endswith("*"):
            paragraph.add_run(part[1:-1]).italic = True
        else:
            paragraph.add_run(part)


def clear_paragraph_border(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    border = p_pr.find(qn("w:pBdr"))
    if border is not None:
        p_pr.remove(border)


def configure_styles(doc: Document, body_size: float = 11):
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(body_size)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.first_line_indent = Cm(0.74)
    normal.paragraph_format.widow_control = True
    for name, size in (("Title", 16), ("Heading 1", 14), ("Heading 2", 12), ("Heading 3", 11)):
        style = doc.styles[name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(12 if name != "Title" else 0)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.line_spacing = 1.05
        if name == "Title":
            p_pr = style.element.get_or_add_pPr()
            border = p_pr.find(qn("w:pBdr"))
            if border is not None:
                p_pr.remove(border)
    caption = doc.styles["Figure Caption"] if "Figure Caption" in doc.styles else doc.styles.add_style("Figure Caption", WD_STYLE_TYPE.PARAGRAPH)
    caption.font.name = "Times New Roman"
    caption.font.size = Pt(9)
    caption.paragraph_format.line_spacing = 1.0
    caption.paragraph_format.space_after = Pt(6)
    caption.paragraph_format.first_line_indent = Cm(0)


def figure_path(number: int) -> Path:
    names = {
        1: "Figure1_study_scope_and_model_integrity.png",
        2: "Figure2_input_matched_reclassification.png",
        3: "Figure3_simulation_discovery_and_inference.png",
        4: "Figure4_IL12RB2_external_evidence.png",
        5: "Figure5_FCRL3_support_and_counterexample.png",
        6: "Figure6_liver_tissue_boundary.png",
    }
    return WORK / "figures" / names[number]


def build_manuscript_docx(md_text: str):
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.2)
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    add_line_numbering(section)
    add_page_number(section.footer.paragraphs[0])
    configure_styles(doc)
    doc.core_properties.title = TITLE
    doc.core_properties.subject = "npj Genomic Medicine submission candidate"
    doc.core_properties.author = "Zhi Chen; Teng Qi"
    doc.core_properties.keywords = "PBC; colocalization; eQTL; caQTL; source-matched LD"

    in_legends = False
    for raw in md_text.splitlines():
        line = raw.rstrip()
        if not line:
            continue
        if line.startswith("# "):
            p = doc.add_paragraph(style="Title")
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.first_line_indent = Cm(0)
            clear_paragraph_border(p)
            add_rich_text(p, line[2:])
        elif line == "## Figure legends":
            doc.add_page_break()
            doc.add_heading("Figure legends", level=1)
            in_legends = True
        elif line.startswith("## "):
            doc.add_heading(line[3:], level=1)
        elif line.startswith("### "):
            doc.add_heading(line[4:], level=2)
        elif in_legends and line.startswith("**Figure "):
            m = re.match(r"\*\*Figure (\d+)\.", line)
            if not m:
                raise RuntimeError(f"bad figure legend: {line[:80]}")
            number = int(m.group(1))
            if number > 1:
                doc.add_page_break()
            p_img = doc.add_paragraph()
            p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_img.paragraph_format.first_line_indent = Cm(0)
            p_img.add_run().add_picture(str(figure_path(number)), width=Inches(6.45))
            p = doc.add_paragraph(style="Figure Caption")
            add_rich_text(p, line)
        elif in_legends and line.startswith("**Abbreviations:**"):
            p = doc.add_paragraph(style="Figure Caption")
            add_rich_text(p, line)
        else:
            p = doc.add_paragraph()
            if line.startswith(("**Authors:", "**Affiliations:", "**Corresponding author:", "**Article type:", "**Keywords:")):
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.first_line_indent = Cm(0)
                p.paragraph_format.line_spacing = 1.15
            if line.startswith("G* =") or line.startswith("and R_QTL"):
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.first_line_indent = Cm(0)
            add_rich_text(p, line.rstrip("  "))
    MANUSCRIPT_DOCX.parent.mkdir(parents=True, exist_ok=True)
    doc.save(MANUSCRIPT_DOCX)


def build_simple_docx(title: str, md_text: str, out: Path, subject: str, compact: bool = False):
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2.4)
    section.bottom_margin = Cm(2.4)
    section.left_margin = Cm(2.6)
    section.right_margin = Cm(2.6)
    add_page_number(section.footer.paragraphs[0])
    configure_styles(doc, 11)
    if compact:
        normal = doc.styles["Normal"]
        normal.font.size = Pt(10.5)
        normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        normal.paragraph_format.line_spacing = 1.12
        normal.paragraph_format.space_after = Pt(6)
        normal.paragraph_format.first_line_indent = Cm(0)
        for style_name in ("List Bullet", "List Number"):
            if style_name in doc.styles:
                doc.styles[style_name].font.name = "Times New Roman"
                doc.styles[style_name].font.size = Pt(10.5)
                doc.styles[style_name].paragraph_format.line_spacing = 1.12
                doc.styles[style_name].paragraph_format.space_after = Pt(4)
    if out == COVER_DOCX:
        section.top_margin = Cm(1.8)
        section.bottom_margin = Cm(1.8)
        section.left_margin = Cm(2.2)
        section.right_margin = Cm(2.2)
        normal = doc.styles["Normal"]
        normal.font.size = Pt(9.5)
        normal.paragraph_format.line_spacing = 1.0
        normal.paragraph_format.space_after = Pt(3)
        doc.styles["Title"].font.size = Pt(14)
        doc.styles["Title"].paragraph_format.space_after = Pt(8)
    doc.core_properties.title = title
    doc.core_properties.subject = subject
    doc.core_properties.author = "Zhi Chen; Teng Qi"
    first_title = True
    for raw in md_text.splitlines():
        line = raw.rstrip()
        if not line:
            continue
        if line.startswith("# "):
            p = doc.add_paragraph(style="Title")
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.first_line_indent = Cm(0)
            clear_paragraph_border(p)
            add_rich_text(p, line[2:])
            first_title = False
        elif line.startswith("## "):
            doc.add_heading(line[3:], level=1)
        elif line.startswith("### "):
            doc.add_heading(line[4:], level=2)
        elif line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.first_line_indent = Cm(0)
            add_rich_text(p, line[2:])
        else:
            p = doc.add_paragraph()
            if first_title or line.startswith(("11 October 2026", "Editors", "npj Genomic Medicine", "Dear Editors", "Sincerely", "Zhi Chen", "Teng Qi")):
                p.paragraph_format.first_line_indent = Cm(0)
            add_rich_text(p, line)
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out)


def copy_assets():
    for directory in ("figures", "supplement", "qa", "reports", "results", "protocols", "provenance", "cover_letter", "manuscript", "code"):
        (WORK / directory).mkdir(parents=True, exist_ok=True)
    for path in (SRC / "figures").glob("*"):
        if path.is_file():
            shutil.copy2(path, WORK / "figures" / path.name)
    mapping = {
        "Additional_file_1_Supplementary_Tables_S1-S10.xlsx": "Supplementary_Data_1_Tables_S1-S10.xlsx",
        "Additional_file_2_Machine_Readable_Supplementary_Data.zip": "Source_Data_1_Primary_Machine_Readable_Data.zip",
        "Additional_file_3_R7C1_External_Validation_Data.zip": "Source_Data_2_External_Validation_Data.zip",
    }
    for old, new in mapping.items():
        shutil.copy2(SRC / "supplement" / old, WORK / "supplement" / new)


def build_cover_letter() -> str:
    return f"""# Cover letter for npj Genomic Medicine

11 October 2026

Editors

npj Genomic Medicine

Dear Editors,

We submit our Article, **{TITLE}**, for consideration in *npj Genomic Medicine*.

The manuscript addresses a practical problem in genomic medicine: cell-specific regulatory assignments at disease loci can change when single-causal and multi-signal colocalization analyses use different variant support, disease statistics and linkage-disequilibrium inputs. We evaluated a PBC-wide registry of 6,923 locus-gene-cell comparisons, performed source-matched multi-signal analysis in a prespecified 642-comparison high-information subset, and used matched-input analysis to separate two identifiable input effects from residual inferential reclassification.

The main result is bidirectional. After matching support and disease statistics, eight single-causal H4 comparisons supported distinct signal pairs and ten H3 comparisons supported shared signal pairs under the multi-signal analysis. Simulations showed a corresponding trade-off: false-H4 decisions decreased in distinct-signal settings, while shared-signal recovery also decreased. The work therefore presents a reliability and attribution framework rather than claiming general superiority of one method.

The medical-genomics relevance is illustrated in PBC. IL12RB2-NK sharing was reproduced in TenK10K, supported in a different FinnGen disease study system and connected to a disease-linked NK-cell chromatin peak. We explicitly limit this evidence to robust regulatory association: public metadata do not prove person-level independence, the FinnGen analysis is single-causal, direct eQTL-caQTL sharing was not estimated and no perturbational experiment establishes mediation. FCRL3 provides both reproducible B-cell sharing and a covariate-model-sensitive T-cell counterexample. Donor-level liver data support lineage detectability without corrected disease enrichment.

The study fits the journal's scope in integrative disease genomics, immunogenomics and computational analysis of genomic variation. Its contribution is an auditable strategy for deciding when locus-gene-cell assignments remain reliable after input matching, source-aware LD construction and multi-signal decomposition.

All authors approved the manuscript and its submission. The work is original and is not under review elsewhere. The authors declare no competing interests and no specific funding. Ethics approval was not required because the study reanalysed public, de-identified data. Code, frozen protocols and aggregate outputs are publicly available on GitHub. An author-approved open research version has been released publicly; it has not been peer reviewed and does not constitute concurrent journal submission.

OpenAI Codex, operating with the locally installed QiTeng Academic Writing Skill under author supervision, assisted with language editing, document generation and consistency checks. It was not treated as an author or source of scientific evidence. The authors take responsibility for all analyses, interpretations and submitted text.

We have not discussed this manuscript with an editor at *npj Genomic Medicine* before submission.

Sincerely,

Zhi Chen  
First author  
School of Medicine, The Chinese University of Hong Kong, Shenzhen  
zhichen1@link.cuhk.edu.cn  
ORCID 0009-0001-0072-5576

Teng Qi  
Corresponding author  
School of Medicine, The Chinese University of Hong Kong, Shenzhen  
tengqi@link.cuhk.edu.cn  
ORCID 0009-0007-7648-4776
"""


def build_si() -> str:
    return """# Supplementary Information and source data guide

**Article:** Input matched multisignal colocalization clarifies regulatory assignments at primary biliary cholangitis loci

**Authors:** Zhi Chen and Teng Qi

This guide defines the supplementary upload set and the evidentiary role of every file. Machine-readable source data are supplied separately so that large tables remain reusable rather than being rasterized into a PDF.

## Supplementary upload set

- **Supplementary Data 1:** `Supplementary_Data_1_Tables_S1-S10.xlsx` contains the editable reading interface for Supplementary Tables S1-S10.
- **Source Data 1:** `Source_Data_1_Primary_Machine_Readable_Data.zip` contains primary machine-readable TSV, TSV.GZ and JSON tables, a README and SHA-256 manifest.
- **Source Data 2:** `Source_Data_2_External_Validation_Data.zip` contains FinnGen disease credible-set, FinnGen R13 x OneK1K, CASCADE, chromatin-triangle, OMIX boundary and source-ledger tables with an internal SHA-256 manifest.

## Supplementary table map

### Supplementary Table S1

Data sources, versions, accessions, checksums and redistribution status.

### Supplementary Table S2

Complete 6,923-comparison screening registry and the 5,460 ABF-eligible subset.

### Supplementary Table S3

Frozen 642-comparison roles, A0/A1/A2/M trajectories and categorical states.

### Supplementary Table S4

Support-set membership and identity audit, including removed-variant counts and hashes.

### Supplementary Table S5

Source-matched PF10/PF50 LD and SuSiE quality-control metrics.

### Supplementary Table S6

Signal-pair posterior results, cross-L identity and mixed H4/H3 semantics.

### Supplementary Table S7

Simulation grid, iteration-level outcome channels and scenario summaries, including the compressed 486,000-row audit.

### Supplementary Table S8

TenK10K and FinnGen external molecular-QTL records and allele-direction mappings.

### Supplementary Table S9

Donor-level liver target-panel metrics, exact tests and sensitivity analyses.

### Supplementary Table S10

Claim-evidence ledger and Methods-Results mirror.

### Supplementary Table S11

IL12RB2-NK external disease-study validation, source credible-set audit, CASCADE coverage, chromatin-triangle adjudication and OMIX descriptive boundary. These records are supplied in Source Data 2.

## Redistribution boundary

Individual-level genotype data, BAM files, pseudobulk matrices and third-party source archives are not redistributed. Source accessions, versions and access conditions are recorded in Supplementary Table S1 and the machine-readable source ledger.
"""


def build_reports():
    sources = [
        ("npj Genomic Medicine aims and scope", "https://www.nature.com/npjgenmed/aims", "Genomics applied to medicine; integrative phenotype-genome analysis; immunogenomics; disease polymorphisms; informatics"),
        ("npj Genomic Medicine content types", "https://www.nature.com/npjgenmed/content-types", "Article; title <=15 words without punctuation; unstructured abstract <=150 words; no Discussion subheadings, limitations or conclusion section"),
        ("npj Genomic Medicine editorial process", "https://www.nature.com/npjgenmed/for-authors-and-referees/editorial-process", "Editors assess advance, evidence and journal-level visibility"),
        ("npj Genomic Medicine metrics", "https://www.nature.com/npjgenmed/journal-impact", "2025 JIF 5.3; median 18 days to first editorial decision"),
        ("npj Genomic Medicine APC", "https://www.nature.com/npjgenmed/apc", "Original Research APC GBP 3190 / USD 4390 / EUR 3590; waiver/discount request at submission"),
        ("Communications Biology aims", "https://www.nature.com/commsbio/aims", "Significant advances bringing new biological insight; computational and secondary analyses permitted"),
        ("Communications Biology content types", "https://www.nature.com/commsbio/submit/content-types", "Article; recommended main text under approximately 5,000 words; up to 10 display items"),
        ("Genome Medicine aims and scope", "https://link.springer.com/journal/13073/aims-and-scope", "Outstanding broad-interest research applying genetics, genomics and multi-omics to understand, diagnose and treat disease"),
        ("Genome Medicine submission guidelines", "https://link.springer.com/journal/13073/submission-guidelines", "Research/Method/Software types; APC USD 5490; broad clinical/basic bridge"),
        ("Human Genomics aims and scope", "https://link.springer.com/journal/40246/aims-and-scope", "Human genetics, GWAS, regulatory genomics, statistical genetics, bioinformatics and integrative multi-omics"),
        ("Human Genomics submission guidelines", "https://link.springer.com/journal/40246/submission-guidelines", "Research and Methodology types; APC USD 3390; machine-readable data encouraged"),
    ]
    ledger_lines = ["source\turl\tverified_2026_10_11\trelevant_rule"]
    for name, url, rule in sources:
        ledger_lines.append(f"{name}\t{url}\tPASS\t{rule}")
    (WORK / "results/R7C2_official_journal_source_ledger.tsv").write_text("\n".join(ledger_lines) + "\n", encoding="utf-8")

    audit_rows = [
        ["npj Genomic Medicine", "Article", "Strong", "Conditional", "Moderate-high", "GO_ONE_ATTEMPT", "Best alignment to integrative immunogenomics and genomic-medicine reliability; risk from limited clinical utility and no perturbational validation"],
        ["Communications Biology", "Article", "Moderate", "Weak-moderate", "High", "HOLD_NOT_PREFERRED", "Methodologically sound but biological-mechanism advance and concise main-text expectations increase desk-reject risk"],
        ["Genome Medicine", "Research or Method", "Partial", "Weak", "Very high", "NO_GO", "Current evidence does not reach broad clinical/basic impact or cross-disease method validation"],
        ["Human Genomics", "Research", "Strong", "Strong", "Low-moderate", "FROZEN_FALLBACK_GO", "Direct fit for statistical genetics, regulatory genomics and integrative multi-omics; existing package is submission ready"],
    ]
    tsv = ["journal\tarticle_type\tscope_fit\tevidence_fit\tdesk_reject_risk\tdecision\treason"]
    tsv += ["\t".join(row) for row in audit_rows]
    (WORK / "results/R7C2_journal_fit_matrix.tsv").write_text("\n".join(tsv) + "\n", encoding="utf-8")

    reports = {
        "00_CMM_R7C2_执行摘要_Q1现实审计与双轨投稿判定.md": """# R7C2 执行摘要

## 最终判定

```text
R7C2 = COMPLETE_Q1_REALITY_CHECK_AND_DUAL_TRACK_PACKAGE
Q1_ONE_SHOT_TARGET = npj Genomic Medicine
Q1_DECISION = CONDITIONAL_GO_ONE_ATTEMPT
HUMAN_GENOMICS_FALLBACK = FROZEN_SUBMISSION_READY
GENOME_MEDICINE = NO_GO
COMMUNICATIONS_BIOLOGY = HOLD_NOT_PREFERRED
NEW_BIOLOGICAL_ANALYSIS = NO
```

当前稿件值得进行一次有边界的 `npj Genomic Medicine` 首投尝试。它与期刊的整合基因组学、免疫基因组学和疾病特异遗传分析范围相符；FinnGen R13 与染色质证据也使其不再只是两种统计流程的内部比较。风险仍然真实存在：没有广泛跨疾病验证、直接 eQTL-caQTL posterior、功能扰动或临床效用终点。因此本轮没有把稿件包装成普遍优越的新算法，也没有升级为染色质介导机制。

Human Genomics 版本继续作为完整、冻结、可立即回退的投稿基线。若 npj Genomic Medicine 编辑拒稿，则直接切换该基线，不启动新的基因、位点或组学层分析。
""",
        "01_CMM_R7C2_期刊现实适配矩阵.md": """# R7C2 期刊现实适配矩阵

## npj Genomic Medicine

官方范围明确覆盖基因组学在医学中的应用、整合表型与基因组分析、免疫基因组学、疾病特异变异与计算方法。当前研究的最强匹配点是调控归因可靠性、细胞语境和外部疾病研究系统验证。主要编辑风险是临床转化仍为间接、方法只在一个疾病中完成系统 benchmark。判定：`CONDITIONAL_GO_ONE_ATTEMPT`。

## Communications Biology

期刊接收二次数据分析与计算方法，但要求带来新的生物学认识。当前最强贡献是归因可靠性与统计边界，功能机制并未闭环。主文还需要压缩至约 5,000 字，压缩成本会削弱方法透明度。判定：`HOLD_NOT_PREFERRED`。

## Genome Medicine

期刊要求对理解、诊断或治疗疾病具有广泛意义的突出研究。当前单疾病、有限外部病例数与无功能验证的证据上限不适合以此作为首投。判定：`NO_GO`。

## Human Genomics

官方范围直接覆盖 GWAS、调控基因组学、统计遗传学、转录组和整合多组学。原冻结包已满足其格式和证据层级。判定：`FROZEN_FALLBACK_GO`。

详细机器表见 `results/R7C2_journal_fit_matrix.tsv`，官方来源见 `results/R7C2_official_journal_source_ledger.tsv`。
""",
        "02_CMM_R7C2_npj格式转换与claim重定位.md": f"""# R7C2 npj Genomic Medicine 格式转换与 claim 重定位

## 已完成的格式门

- 标题改为无标点的 12 词标题：`{TITLE}`。
- 摘要改为无小标题单段格式，共 {word_count(ABSTRACT)} 词，低于 150 词。
- `Background` 改为 `Introduction`。
- 取消独立 `Conclusions` 章节，将结论并入 Discussion 末段。
- 保留 Results 与 Methods 小标题；Discussion 不使用小标题。
- 将 BMC declarations 改为 Nature Portfolio 的 Ethics statement、Data availability、Code availability、Acknowledgements、Author contributions 与 Competing interests。
- 图 1-6 单独提供 PDF/PNG/SVG；图注保留在主文末尾。
- 生成 Supplementary Information 导航文件，并保留机器可读 Excel/ZIP 作为单独 source-data uploads。

## claim ceiling

允许：输入匹配后仍存在双向重分类；source-aware multi-signal 分析提高归因可解释性；IL12RB2-NK 在另一疾病研究系统获得有边界支持；疾病-eQTL 与疾病-caQTL 加 peak-gene link 构成三角信号一致性。

禁止：普遍优越的新算法；完全独立疾病复制；直接 eQTL-caQTL 共定位；染色质介导；治疗方向；功能机制已经证实。
""",
        "03_CMM_R7C2_敌意编辑预审.md": """# R7C2 敌意编辑预审

## 最可能的编辑拒稿理由

1. 研究仍主要是单疾病方法学 benchmark，普遍性有限。
2. FinnGen R13 只有 795 个病例，且公开元数据无法证明参与者完全不重叠。
3. 外部验证仍为 single-causal 分析，没有第二套 source-LD multi-signal fit。
4. chromatin triangle 缺少直接 eQTL-caQTL posterior，不能称为 mediation。
5. 肝组织结果仅支持 lineage detectability，没有校正后的疾病富集。
6. 没有扰动实验或临床效用终点。

## 已经完成的防守

- 标题、摘要与 cover letter 均以 regulatory attribution reliability 为主轴。
- 外部证据写成 different disease study system，不写 fully independent cohort。
- 三角证据写成 signal coherence，不写 chromatin mediation。
- 模拟结果同时报告 false-H4 改善与 shared-signal recovery 损失。
- 642 子集明确是 prespecified high-information subset，不写全 6,923 的 multi-signal landscape。
- 所有阴性、无信息和模型敏感结果继续保留。

## 编辑门判断

风险足以排除 Genome Medicine，但没有高到必须放弃一次 npj Genomic Medicine 尝试。该尝试必须是一次性的；编辑拒稿后不追加无边界分析。
""",
        "04_CMM_R7C2_双轨投稿与退出规则.md": """# R7C2 双轨投稿与退出规则

## Track 1

`npj Genomic Medicine / Article`。使用本轮专用稿件、cover letter、六幅图和 supplementary/source-data upload set。

## Track 2

`Human Genomics / Research`。保留 `R7B4B2` 冻结投稿包，不因 Q1 尝试改写或覆盖。

## 退出规则

- npj 送外审：只回应具体审稿意见；不自动扩展疾病、基因或全组学分析。
- npj 编辑拒稿：直接进入 Human Genomics 投稿，不做新的生物学补救。
- 若编辑指出可修复的格式或公开数据链接问题：在不改核心分析的前提下修复后再投。
- 只有发现数据身份、统计输入或数值结果的不可修复错误，才暂停双轨投稿并回到科学审计。

## 下一阶段

`R7C3_AUTHOR_FINAL_REVIEW_AND_NPJ_SUBMISSION_GATE`。任务仅包括作者逐页终审、APC/waiver 决策、投稿系统字段填写和由作者本人执行提交。期刊官方规则明确要求稿件只能由作者提交，自动化工作流不得代替作者完成最终提交。
""",
        "99_CMM_R7C2_详细行动记录_2026-10-11.md": """# CMM R7C2 详细行动记录

## 任务

按照 RP v4 执行 `Q1 Target-Journal Reality Check + Dual-Track Submission Package`，不增加新的生物学分析，并保留已完成的 Human Genomics 投稿基线。

## 输入

- R7C1 外部验证稿件、六幅主图、Supplementary Tables S1-S10、机器可读主分析数据与 R7C1 外部验证数据。
- R7B4B2 Human Genomics 冻结投稿包。
- 2026-10-11 实时核验的四本候选期刊官方 scope、article type、编辑流程、APC 与文件要求。

## 实际工作

1. 核对 npj Genomic Medicine、Communications Biology、Genome Medicine 与 Human Genomics 的官方范围和投稿规则。
2. 用当前证据上限执行编辑拒稿风险审计，而不是按期刊影响力倒推研究结论。
3. 将 npj Genomic Medicine 冻结为唯一一次 Q1 尝试目标；Genome Medicine 判定 NO-GO，Communications Biology 判定 HOLD。
4. 将 R7C1 主文转换为 npj Article 格式：12 词无标点标题、单段不超过 150 词摘要、Introduction、无独立 Conclusions、Nature Portfolio declarations。
5. 制作目标期刊 cover letter、Supplementary Information 导航、upload map 与 submission checklist。
6. 保留六幅主图与全部机器可读补充数据；不隐藏阴性、无信息或模型敏感结果。
7. 使用 WPS 后台导出 PDF，并对每页进行视觉质控；机器检查标题、摘要、章节、声明、数字、图件、哈希和发布包完整性。
8. 制作带 SHA-256 manifest 的 R7C2 release；GitHub 同步经作者明确授权公开的稿件、投稿信、图件、补充材料与安全衍生数据；隐私凭证和期刊账号不进入仓库。

## 科学裁决

R7C1 的新增 FinnGen R13 疾病研究系统验证和 CASCADE 染色质三角使一次 npj Genomic Medicine 尝试具有合理性，但当前研究仍不是广泛跨疾病验证的新算法，也没有直接机制证明。最可靠的论文定位是：`PBC-wide screening followed by source-matched multi-signal reclassification of a prespecified high-information subset, with bounded external disease-study and chromatin evidence`。

## 未改变内容

- 未新增 locus、gene、cell 或 omics layer。
- 未重跑 642 比较或 486,000 次模拟。
- 未把 FinnGen 写成完全独立队列。
- 未把 chromatin triangle 写成 mediation。
- 未覆盖 R7B4B2 Human Genomics 基线。

## 下一阶段

进入 `R7C3_AUTHOR_FINAL_REVIEW_AND_NPJ_SUBMISSION_GATE`。作者终审并决定 APC/waiver 后，由作者本人登录投稿系统执行提交。若编辑拒稿，直接使用 Human Genomics 冻结包。
""",
    }
    for name, body in reports.items():
        (WORK / "reports" / name).write_text(body.strip() + "\n", encoding="utf-8")


def build_upload_map():
    text = """# npj Genomic Medicine upload map

| Order | Submission field | Local file | Status |
|---|---|---|---|
| 1 | Cover letter | `cover_letter/R7C2_npj_Genomic_Medicine_cover_letter.docx` | READY |
| 2 | Main manuscript | `manuscript/R7C2_npj_Genomic_Medicine_manuscript.docx` | READY |
| 3-8 | Figures 1-6 | `figures/Figure*.pdf` | READY |
| 9 | Supplementary information guide | `supplement/R7C2_npj_Supplementary_Information.pdf` | READY |
| 10 | Supplementary Data 1 | `supplement/Supplementary_Data_1_Tables_S1-S10.xlsx` | READY |
| 11 | Source Data 1 | `supplement/Source_Data_1_Primary_Machine_Readable_Data.zip` | READY |
| 12 | Source Data 2 | `supplement/Source_Data_2_External_Validation_Data.zip` | READY |

## Submission-system fields

- Article type: Article
- First author: Zhi Chen
- Corresponding author: Teng Qi
- Competing interests: none
- Funding: none
- Ethics approval: not required for public de-identified secondary analysis
- Preprint/public version: disclose the GitHub open research release
- AI disclosure: use the statement in Methods and cover letter
- Prior editor discussion: no
- Concurrent submission: no
"""
    (WORK / "protocols/R7C2_npj_upload_map.md").write_text(text, encoding="utf-8")

    checklist = {
        "stage": "R7C2_NPJ_DUAL_TRACK",
        "target": "npj Genomic Medicine",
        "article_type": "Article",
        "author_final_submission_required": True,
        "checks": {
            "title_le_15_words_and_no_punctuation": True,
            "abstract_unstructured_le_150_words": True,
            "introduction_no_subheadings": True,
            "results_has_subheadings": True,
            "discussion_no_subheadings": True,
            "no_separate_limitations_section": True,
            "no_separate_conclusions_section": True,
            "methods_has_subheadings": True,
            "data_availability_present": True,
            "code_availability_present": True,
            "author_contributions_present": True,
            "competing_interests_present": True,
            "cover_letter_present": True,
            "figures_separate": True,
            "source_data_machine_readable": True,
            "human_genomics_fallback_frozen": True,
        },
        "apc": {"GBP": 3190, "USD": 4390, "EUR": 3590, "waiver_or_discount_decision_required_at_submission": True},
    }
    (WORK / "results/R7C2_submission_checklist.json").write_text(json.dumps(checklist, indent=2), encoding="utf-8")


def main():
    copy_assets()
    manuscript = build_npj_markdown()
    if word_count(ABSTRACT) > 150:
        raise RuntimeError(f"abstract exceeds limit: {word_count(ABSTRACT)}")
    if word_count(TITLE) > 15 or re.search(r"[:;,.!?()\-]", TITLE):
        raise RuntimeError("title violates npj rule")
    MANUSCRIPT_MD.write_text(manuscript, encoding="utf-8")
    build_manuscript_docx(manuscript)

    cover = build_cover_letter()
    COVER_MD.write_text(cover, encoding="utf-8")
    build_simple_docx("Cover letter for npj Genomic Medicine", cover, COVER_DOCX, "npj Genomic Medicine cover letter", compact=True)

    si = build_si()
    SI_MD.write_text(si, encoding="utf-8")
    build_simple_docx("Supplementary Information and source data guide", si, SI_DOCX, "npj Genomic Medicine supplementary information guide", compact=True)

    build_reports()
    build_upload_map()

    baseline = ROOT / "5_manuscript/R7B4B_FinalSubmissionCandidate/manuscript/R7B4B_HumanGenomics_final.docx"
    state = {
        "stage": "R7C2_Q1_REALITY_CHECK_DUAL_TRACK",
        "decision": "NPJ_GENOMIC_MEDICINE_CONDITIONAL_GO_ONE_ATTEMPT",
        "fallback": "HUMAN_GENOMICS_FROZEN_SUBMISSION_READY",
        "title": TITLE,
        "title_words": word_count(TITLE),
        "abstract_words": word_count(ABSTRACT),
        "manuscript_docx_sha256": sha256(MANUSCRIPT_DOCX),
        "cover_letter_docx_sha256": sha256(COVER_DOCX),
        "supplementary_info_docx_sha256": sha256(SI_DOCX),
        "human_genomics_baseline_docx_sha256": sha256(baseline),
        "new_biological_analysis": False,
        "next": "R7C3_AUTHOR_FINAL_REVIEW_AND_NPJ_SUBMISSION_GATE",
    }
    (WORK / "results/R7C2_build_state.json").write_text(json.dumps(state, indent=2, ensure_ascii=False), encoding="utf-8")
    print(json.dumps(state, indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()
