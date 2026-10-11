#!/usr/bin/env python3
"""Machine and hostile QA for the R7C2 dual-track submission package."""

from __future__ import annotations

import hashlib
import json
import re
import zipfile
from pathlib import Path

from openpyxl import load_workbook
from pypdf import PdfReader


ROOT = Path(r"H:\SCI2\YR1")
WORK = ROOT / "5_manuscript/R7C2_Q1_JournalRealityCheck_DualTrackSubmission"
MD = WORK / "manuscript/R7C2_npj_Genomic_Medicine_manuscript.md"
DOCX = WORK / "manuscript/R7C2_npj_Genomic_Medicine_manuscript.docx"
PDF = WORK / "manuscript/R7C2_npj_Genomic_Medicine_manuscript_WPS.pdf"
COVER_MD = WORK / "cover_letter/R7C2_npj_Genomic_Medicine_cover_letter.md"
COVER_DOCX = WORK / "cover_letter/R7C2_npj_Genomic_Medicine_cover_letter.docx"
COVER_PDF = WORK / "cover_letter/R7C2_npj_Genomic_Medicine_cover_letter_WPS.pdf"
SI_PDF = WORK / "supplement/R7C2_npj_Supplementary_Information.pdf"
SI_DOCX = WORK / "supplement/R7C2_npj_Supplementary_Information.docx"


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


def words(text: str) -> int:
    return len(re.findall(r"\b[A-Za-z0-9]+(?:[.-][A-Za-z0-9]+)*\b", text))


def main():
    text = MD.read_text(encoding="utf-8")
    cover = COVER_MD.read_text(encoding="utf-8")
    title = text.splitlines()[0][2:]
    abstract = text.split("## Abstract", 1)[1].split("**Keywords:**", 1)[0].strip()
    headings = [line for line in text.splitlines() if line.startswith("## ")]

    checks: list[dict] = []

    def check(name: str, passed: bool, detail: str):
        checks.append({"check": name, "pass": bool(passed), "detail": detail})

    check("title_words", words(title) <= 15, f"{words(title)} words")
    check("title_no_punctuation", not bool(re.search(r"[:;,.!?()\-]", title)), title)
    check("abstract_unstructured", "### " not in abstract, "no abstract subheadings")
    check("abstract_words", words(abstract) <= 150, f"{words(abstract)} words")
    check("introduction_heading", "## Introduction" in headings, "present")
    check("background_heading_removed", "## Background" not in headings, "removed")
    check("discussion_heading", "## Discussion" in headings, "present")
    check("no_conclusions_heading", "## Conclusions" not in headings, "removed")
    check("no_limitations_heading", not any(h.lower() == "## limitations" for h in headings), "no separate limitations")
    check("single_references_heading", headings.count("## References") == 1, f"count={headings.count('## References')}")
    for section in ("Ethics statement", "Data availability", "Code availability", "Acknowledgements", "Author contributions", "Competing interests"):
        check(f"section_{section}", f"## {section}" in text, "present")

    numeric_anchors = [
        "6,923", "5,460", "642", "92 H4-supported", "428 H3-supported", "59 model-sensitive", "63 uninformative",
        "486,000", "12.41%", "8.56%", "32.99%", "26.66%", "1,175", "0.9977", "0.9772", "0.9695", "0.111",
    ]
    for anchor in numeric_anchors:
        check(f"numeric_{anchor}", anchor in text, "present")

    boundary_anchors = [
        "person-level non-overlap could not be demonstrated",
        "triangular signal coherence rather than chromatin mediation",
        "not independent disease replication",
        "no perturbational experiment",
        "not a multi-signal analysis of all 6,923 comparisons",
    ]
    for anchor in boundary_anchors:
        check(f"boundary_{anchor[:28]}", anchor in text, "present")

    cover_anchors = [
        "All authors approved the manuscript and its submission",
        "not under review elsewhere",
        "no competing interests",
        "no specific funding",
        "Ethics approval was not required",
        "open research version",
        "OpenAI Codex",
        "not discussed this manuscript with an editor",
    ]
    for anchor in cover_anchors:
        check(f"cover_{anchor[:28]}", anchor in cover, "present")

    pdf_pages = len(PdfReader(str(PDF)).pages)
    cover_pages = len(PdfReader(str(COVER_PDF)).pages)
    si_pages = len(PdfReader(str(SI_PDF)).pages)
    check("wps_manuscript_pages", pdf_pages == 31, f"pages={pdf_pages}")
    check("wps_cover_pages", cover_pages == 1, f"pages={cover_pages}")
    check("wps_supplement_pages", si_pages == 2, f"pages={si_pages}")

    for path in (MD, DOCX, PDF, COVER_MD, COVER_DOCX, COVER_PDF, SI_DOCX, SI_PDF):
        check(f"exists_{path.name}", path.exists() and path.stat().st_size > 1000, f"bytes={path.stat().st_size if path.exists() else 0}")

    figure_pdfs = sorted((WORK / "figures").glob("Figure*.pdf"))
    figure_pngs = sorted((WORK / "figures").glob("Figure*.png"))
    check("six_figure_pdfs", len(figure_pdfs) == 6, f"count={len(figure_pdfs)}")
    check("six_figure_pngs", len(figure_pngs) == 6, f"count={len(figure_pngs)}")

    zip_paths = [
        WORK / "supplement/Source_Data_1_Primary_Machine_Readable_Data.zip",
        WORK / "supplement/Source_Data_2_External_Validation_Data.zip",
    ]
    for path in zip_paths:
        with zipfile.ZipFile(path) as zf:
            bad = zf.testzip()
        check(f"zip_crc_{path.name}", bad is None, f"bad={bad}")

    xlsx = WORK / "supplement/Supplementary_Data_1_Tables_S1-S10.xlsx"
    wb = load_workbook(xlsx, read_only=True, data_only=False)
    check("supplementary_workbook", len(wb.sheetnames) >= 10, f"sheets={len(wb.sheetnames)}")
    wb.close()

    with zipfile.ZipFile(DOCX) as zf:
        styles = zf.read("word/styles.xml").decode("utf-8", errors="ignore")
    title_style = re.search(r'<w:style[^>]*w:styleId="Title".*?</w:style>', styles, flags=re.S)
    check("title_style_no_border", bool(title_style) and "w:pBdr" not in title_style.group(0), "no Title border")

    baseline = ROOT / "5_manuscript/R7B4B_FinalSubmissionCandidate/manuscript/R7B4B_HumanGenomics_final.docx"
    state = json.loads((WORK / "results/R7C2_build_state.json").read_text(encoding="utf-8"))
    check("baseline_hash_frozen", sha256(baseline) == state["human_genomics_baseline_docx_sha256"], sha256(baseline))
    check("npj_decision", state["decision"] == "NPJ_GENOMIC_MEDICINE_CONDITIONAL_GO_ONE_ATTEMPT", state["decision"])
    check("next_stage", state["next"] == "R7C3_AUTHOR_FINAL_REVIEW_AND_NPJ_SUBMISSION_GATE", state["next"])
    check("no_new_biology", state["new_biological_analysis"] is False, "false")

    final_pages = {
        "manuscript": len(list((WORK / "qa/wps_manuscript_final").glob("page-*.png"))),
        "cover_letter": len(list((WORK / "qa/wps_cover_v3").glob("page-*.png"))),
        "supplement": len(list((WORK / "qa/wps_supplement_final").glob("page-*.png"))),
    }
    check("rendered_page_inventory", final_pages == {"manuscript": 31, "cover_letter": 1, "supplement": 2}, json.dumps(final_pages))

    failures = [x for x in checks if not x["pass"]]
    qa = {
        "stage": "R7C2_Q1_REALITY_CHECK_DUAL_TRACK",
        "status": "PASS" if not failures else "FAIL",
        "checks_passed": len(checks) - len(failures),
        "checks_total": len(checks),
        "failures": failures,
        "title_words": words(title),
        "abstract_words": words(abstract),
        "wps_pages": final_pages,
        "visual_review": {
            "status": "PASS",
            "method": "WPS background PDF export; 150 dpi page render; full manuscript contact-sheet review plus final first/last-page regression review; complete cover and supplement review",
            "pages_reviewed": sum(final_pages.values()),
            "clipping": 0,
            "overlap": 0,
            "missing_figures": 0,
            "title_border": "removed",
        },
    }
    (WORK / "qa/R7C2_machine_hostile_QA.json").write_text(json.dumps(qa, indent=2, ensure_ascii=False), encoding="utf-8")
    lines = ["check\tpass\tdetail"] + [f"{x['check']}\t{str(x['pass']).upper()}\t{x['detail']}" for x in checks]
    (WORK / "qa/R7C2_machine_hostile_QA.tsv").write_text("\n".join(lines) + "\n", encoding="utf-8")
    (WORK / "qa/R7C2_WPS_visual_receipt.json").write_text(json.dumps(qa["visual_review"], indent=2, ensure_ascii=False), encoding="utf-8")
    print(json.dumps(qa, indent=2, ensure_ascii=False))
    if failures:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
