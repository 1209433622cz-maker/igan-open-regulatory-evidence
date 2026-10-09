#!/usr/bin/env python3
"""Independent R7B4B pre-author gate for the frozen R7B4A release."""

from __future__ import annotations

import hashlib
import json
import re
import subprocess
import zipfile
from pathlib import Path, PurePosixPath

import pandas as pd
from docx import Document
from pypdf import PdfReader


ROOT = Path(r"H:\SCI2\YR1")
UPSTREAM = ROOT / "5_manuscript" / "R7B4A_HumanGenomics_SubmissionInterface"
OUT = ROOT / "5_manuscript" / "R7B4B_PreauthorFinalSubmissionGate"
REPO = ROOT / "github" / "igan-open-regulatory-evidence"
ARCHIVE = ROOT / "6_release" / "CMM_R7B4A_HumanGenomics_SubmissionInterface_2026-10-10.zip"
EXPECTED_ARCHIVE_SHA = "0b12b5f390b0492d31a1e93d588b07cd3ba8f70670240a1a46fd5898e754a91e"
EXPECTED_ARCHIVE_BYTES = 18_289_096
EXPECTED_ENTRIES = 56
EXPECTED_INTERNAL = 55
TAG = "r7b4a-human-genomics-interface-2026-10-10"
EXPECTED_TAG_COMMIT = "c768ff4c02481331c6563cf8e040062a8e31b6db"
EXPECTED_MAIN = "375dd6d6e1cb3949fed50089a14e446eae9c9b4e"


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def run_git(*args: str) -> str:
    return subprocess.run(
        ["git", *args], cwd=REPO, check=True, capture_output=True, text=True
    ).stdout.strip()


def citation_first_appearance(text: str) -> list[int]:
    body = text.split("\n## References\n", 1)[0]
    order: list[int] = []
    for match in re.finditer(r"\[([0-9][0-9,;\-– ]*)\]", body):
        payload = match.group(1).replace("–", "-")
        for part in re.split(r"[,;]\s*", payload):
            part = part.strip()
            if "-" in part:
                start, end = map(int, part.split("-", 1))
                values = range(start, end + 1)
            else:
                values = [int(part)]
            for value in values:
                if value not in order:
                    order.append(value)
    return order


def verify_release() -> dict:
    audit = {
        "path": str(ARCHIVE),
        "bytes": ARCHIVE.stat().st_size,
        "sha256": sha256_file(ARCHIVE),
    }
    audit["exact_release_identity"] = (
        audit["bytes"] == EXPECTED_ARCHIVE_BYTES
        and audit["sha256"] == EXPECTED_ARCHIVE_SHA
    )
    with zipfile.ZipFile(ARCHIVE) as archive:
        names = archive.namelist()
        audit["zip_entries"] = len(names)
        audit["unique_members"] = len(names) == len(set(names))
        bad = archive.testzip()
        audit["zip_crc"] = "PASS" if bad is None else "FAIL"
        audit["bad_crc_member"] = bad
        manifests = [name for name in names if PurePosixPath(name).name.lower() == "checksums.sha256"]
        if len(manifests) != 1:
            raise RuntimeError(f"Expected exactly one internal manifest, observed {len(manifests)}")
        manifest = manifests[0]
        parent = PurePosixPath(manifest).parent
        exceptions = []
        checked = 0
        for line in archive.read(manifest).decode("utf-8-sig").splitlines():
            if not line.strip():
                continue
            expected, relative = line.split(maxsplit=1)
            relative = relative.lstrip("*").replace("\\", "/")
            rel = PurePosixPath(relative)
            if rel.is_absolute() or ".." in rel.parts:
                raise RuntimeError(f"Unsafe internal path: {relative}")
            choices = [str(parent / rel), str(rel)]
            member = next((name for name in choices if name in names), None)
            checked += 1
            if member is None:
                exceptions.append({"file": relative, "error": "MISSING"})
                continue
            actual = hashlib.sha256(archive.read(member)).hexdigest()
            if actual.lower() != expected.lower():
                exceptions.append({"file": relative, "error": "HASH_MISMATCH"})
        audit["internal_checksums"] = {
            "checked": checked,
            "passed": checked - len(exceptions),
            "exceptions": exceptions,
        }
    audit["status"] = (
        "PASS"
        if audit["exact_release_identity"]
        and audit["zip_entries"] == EXPECTED_ENTRIES
        and audit["unique_members"]
        and audit["zip_crc"] == "PASS"
        and audit["internal_checksums"]["checked"] == EXPECTED_INTERNAL
        and not audit["internal_checksums"]["exceptions"]
        else "FAIL"
    )
    return audit


def main() -> int:
    (OUT / "results").mkdir(parents=True, exist_ok=True)
    (OUT / "qa").mkdir(parents=True, exist_ok=True)
    release = verify_release()
    (OUT / "results" / "R7B4A_release_identity_audit.json").write_text(
        json.dumps(release, indent=2, ensure_ascii=False), encoding="utf-8"
    )

    manuscript_path = UPSTREAM / "manuscript" / "R7B4A_HumanGenomics_manuscript.md"
    docx_path = UPSTREAM / "manuscript" / "R7B4A_HumanGenomics_manuscript.docx"
    pdf_path = UPSTREAM / "manuscript" / "R7B4A_HumanGenomics_manuscript_WPS.pdf"
    cover_md_path = UPSTREAM / "cover_letter" / "R7B4A_HumanGenomics_cover_letter_DRAFT.md"
    cover_pdf_path = UPSTREAM / "cover_letter" / "R7B4A_HumanGenomics_cover_letter_DRAFT_WPS.pdf"
    text = manuscript_path.read_text(encoding="utf-8")
    cover_text = cover_md_path.read_text(encoding="utf-8")
    document = Document(docx_path)
    manuscript_pdf = PdfReader(str(pdf_path))
    cover_pdf = PdfReader(str(cover_pdf_path))
    upstream_qa = json.loads((UPSTREAM / "qa" / "R7B4A_final_QA.json").read_text(encoding="utf-8"))

    placeholder_patterns = [
        r"TO BE COMPLETED",
        r"AUTHOR CONFIRMATION REQUIRED",
        r"LOCAL INSTITUTIONAL DETERMINATION",
        r"CRediT CONTRIBUTIONS TO BE COMPLETED",
        r"ADDITIONAL ACKNOWLEDGEMENTS TO BE COMPLETED",
        r"OPTIONAL AUTHOR BIOGRAPHICAL INFORMATION TO BE COMPLETED",
        r"FULL AUTHOR NAMES TO BE COMPLETED",
        r"NUMBERED INSTITUTIONAL ADDRESSES TO BE COMPLETED",
        r"POSTAL ADDRESS, EMAIL, AND TELEPHONE TO BE COMPLETED",
    ]
    manuscript_placeholders = [
        pattern for pattern in placeholder_patterns if re.search(pattern, text, flags=re.I)
    ]
    cover_placeholders = re.findall(r"\[[A-Z][A-Z0-9 /,.'\-]+\]", cover_text)
    first_order = citation_first_appearance(text)
    reference_numbers = [int(value) for value in re.findall(r"(?m)^(\d+)\. ", text)]

    repo_audit = {
        "local_head": run_git("rev-parse", "HEAD"),
        "origin_main": run_git("rev-parse", "origin/main"),
        "tag_object": run_git("rev-parse", f"refs/tags/{TAG}"),
        "tag_commit": run_git("rev-parse", f"refs/tags/{TAG}^{{}}"),
        "worktree_clean": run_git("status", "--porcelain") == "",
    }
    repo_audit["status"] = (
        "PASS"
        if repo_audit["local_head"] == EXPECTED_MAIN
        and repo_audit["origin_main"] == EXPECTED_MAIN
        and repo_audit["tag_commit"] == EXPECTED_TAG_COMMIT
        and repo_audit["worktree_clean"]
        else "FAIL"
    )
    (OUT / "results" / "R7B4A_repository_identity_audit.json").write_text(
        json.dumps(repo_audit, indent=2, ensure_ascii=False), encoding="utf-8"
    )

    checks: list[dict[str, str]] = []

    def check(name: str, status: str, detail) -> None:
        checks.append({"check": name, "status": status, "detail": str(detail)})

    check("release_identity", "PASS" if release["status"] == "PASS" else "FAIL", release["status"])
    check("repository_identity", "PASS" if repo_audit["status"] == "PASS" else "FAIL", repo_audit)
    check("historical_r7b4a_qa", "PASS" if upstream_qa["status"] == "PASS" and upstream_qa["summary"] == {"pass": 42, "total": 42} else "FAIL", upstream_qa["summary"])
    check("reference_list_continuity", "PASS" if reference_numbers == list(range(1, 23)) else "FAIL", reference_numbers)
    check("reference_first_appearance", "PASS" if first_order == list(range(1, 23)) else "FAIL", first_order)
    check("author_review_placeholders_explicit", "PASS" if len(manuscript_placeholders) >= 6 else "FAIL", manuscript_placeholders)
    check("author_review_internal_gate_present", "PASS" if any(paragraph.text == "Author metadata completion gate" for paragraph in document.paragraphs) else "FAIL", "historical author-review condition")
    check("author_review_cover_confirmations", "PASS" if len(cover_placeholders) >= 6 else "FAIL", len(cover_placeholders))
    check("historical_wps_manuscript", "PASS" if len(manuscript_pdf.pages) == 31 and "wps" in str(manuscript_pdf.metadata.get("/Creator", "")).lower() else "FAIL", len(manuscript_pdf.pages))
    check("historical_wps_cover", "PASS" if len(cover_pdf.pages) == 1 and "wps" in str(cover_pdf.metadata.get("/Creator", "")).lower() else "FAIL", len(cover_pdf.pages))
    check("final_author_metadata", "HOLD", "verified author record not supplied")
    check("final_placeholder_zero", "HOLD", f"manuscript={len(manuscript_placeholders)} cover={len(cover_placeholders)}")
    check("final_internal_gate_removed", "HOLD", "historical review DOCX correctly retains internal gate")
    check("final_declarations_confirmed", "HOLD", "ethics/funding/COI/CRediT/acknowledgements pending")
    check("institutional_q2_confirmation", "HOLD", "institution-specific evidence not supplied")
    check("final_wps_parity", "HOLD", "final DOCX/PDF not generated")
    check("all_author_final_approval", "HOLD", "version-bound approvals not supplied")
    check("submission_authorization", "HOLD", "author-operated submission remains unauthorized")

    failures = [item for item in checks if item["status"] == "FAIL"]
    holds = [item for item in checks if item["status"] == "HOLD"]
    status = "FAIL" if failures else ("HOLD_HUMAN_INPUT_REQUIRED" if holds else "PASS")
    result = {
        "stage": "R7B4B_PREAUTHOR_INDEPENDENT_GATE",
        "status": status,
        "summary": {
            "pass": sum(item["status"] == "PASS" for item in checks),
            "hold": len(holds),
            "fail": len(failures),
            "total": len(checks),
        },
        "checks": checks,
        "release_identity": release["status"],
        "repository_identity": repo_audit["status"],
        "new_biological_analysis": "NO",
        "submission_authorized": "NO",
        "next": "R7B4B_VERIFIED_AUTHOR_INPUT_AND_FINAL_CANDIDATE_BUILD",
    }
    (OUT / "qa" / "R7B4B_preauthor_gate_QA.json").write_text(
        json.dumps(result, indent=2, ensure_ascii=False), encoding="utf-8"
    )
    pd.DataFrame(checks).to_csv(
        OUT / "qa" / "R7B4B_preauthor_gate_QA.tsv", sep="\t", index=False
    )
    (OUT / "results" / "R7B4B_preauthor_gate_state.json").write_text(
        json.dumps(
            {
                "stage": "R7B4B_PREAUTHOR_INDEPENDENT_GATE",
                "status": status,
                "primary_journal_route": "HUMAN_GENOMICS_RESEARCH_LOCKED",
                "institutional_q2_confirmation": "PENDING",
                "final_author_metadata_injection": "NOT_STARTED",
                "final_submission_qa": "NOT_EXECUTED",
                "final_wps_parity": "NOT_EXECUTED",
                "new_biological_analysis": "NO",
                "submission_authorized": "NO",
                "next": "R7B4B_VERIFIED_AUTHOR_INPUT_AND_FINAL_CANDIDATE_BUILD",
            },
            indent=2,
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )
    print(json.dumps(result["summary"] | {"status": status}, ensure_ascii=False))
    return 1 if failures else 0


if __name__ == "__main__":
    raise SystemExit(main())
