#!/usr/bin/env python3
"""Fail-closed FINAL_SUBMISSION QA for an R7B4B Human Genomics candidate."""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import subprocess
import zipfile
from pathlib import Path

from PIL import Image
from docx import Document
from pypdf import PdfReader


ROOT = Path(r"H:\SCI2\YR1")
UPSTREAM = ROOT / "5_manuscript" / "R7B4A_HumanGenomics_SubmissionInterface"
DEFAULT_CANDIDATE = ROOT / "5_manuscript" / "R7B4B_FinalSubmissionCandidate"
DEFAULT_RECORD = ROOT / "5_manuscript" / "R7B4B_PreauthorFinalSubmissionGate" / "submission" / "R7B4B_author_record.EMPTY.json"
DEFAULT_OUT = ROOT / "5_manuscript" / "R7B4B_PreauthorFinalSubmissionGate" / "qa" / "R7B4B_final_submission_QA.json"
REPO = ROOT / "github" / "igan-open-regulatory-evidence"
TAG = "r7b4a-human-genomics-interface-2026-10-10"
EXPECTED_TAG_COMMIT = "c768ff4c02481331c6563cf8e040062a8e31b6db"
TITLE = "Input-matched multi-signal colocalization clarifies immune-cell regulatory assignments at primary biliary cholangitis risk loci"

PLACEHOLDER_PATTERNS = [
    r"TO BE COMPLETED",
    r"AUTHOR CONFIRMATION REQUIRED",
    r"LOCAL INSTITUTIONAL DETERMINATION",
    r"CRediT CONTRIBUTIONS TO BE COMPLETED",
    r"ADDITIONAL ACKNOWLEDGEMENTS TO BE COMPLETED",
    r"OPTIONAL AUTHOR BIOGRAPHICAL INFORMATION TO BE COMPLETED",
    r"FULL AUTHOR NAMES TO BE COMPLETED",
    r"NUMBERED INSTITUTIONAL ADDRESSES TO BE COMPLETED",
    r"CORRESPONDING AUTHOR NAME",
    r"SIGNATURE BLOCK",
]


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def citation_first_appearance(text: str) -> list[int]:
    body = text.split("\n## References\n", 1)[0]
    order: list[int] = []
    for match in re.finditer(r"\[([0-9][0-9,;\-– ]*)\]", body):
        for part in re.split(r"[,;]\s*", match.group(1).replace("–", "-")):
            if "-" in part:
                start, end = map(int, part.strip().split("-", 1)); values = range(start, end + 1)
            else:
                values = [int(part.strip())]
            for value in values:
                if value not in order:
                    order.append(value)
    return order


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--candidate", type=Path, default=DEFAULT_CANDIDATE)
    parser.add_argument("--author-record", type=Path, default=DEFAULT_RECORD)
    parser.add_argument("--out", type=Path, default=DEFAULT_OUT)
    args = parser.parse_args()

    paths = {
        "md": args.candidate / "manuscript" / "R7B4B_HumanGenomics_final.md",
        "docx": args.candidate / "manuscript" / "R7B4B_HumanGenomics_final.docx",
        "pdf": args.candidate / "manuscript" / "R7B4B_HumanGenomics_final_WPS.pdf",
        "cover_md": args.candidate / "cover_letter" / "R7B4B_HumanGenomics_cover_letter_final.md",
        "cover_docx": args.candidate / "cover_letter" / "R7B4B_HumanGenomics_cover_letter_final.docx",
        "cover_pdf": args.candidate / "cover_letter" / "R7B4B_HumanGenomics_cover_letter_final_WPS.pdf",
        "render_receipt": args.candidate / "qa" / "WPS_render_receipt.json",
        "graphical_abstract": args.candidate / "graphical_abstract" / "Graphical_Abstract_HumanGenomics_920x300.png",
        "supp1": args.candidate / "supplement" / "Additional_file_1_Supplementary_Tables_S1-S10.xlsx",
        "supp2": args.candidate / "supplement" / "Additional_file_2_Machine_Readable_Supplementary_Data.zip",
    }
    checks = []

    def check(name: str, status: str, detail) -> None:
        checks.append({"check": name, "status": status, "detail": str(detail)})

    missing = [name for name, path in paths.items() if not path.exists()]
    check("candidate_file_inventory", "PASS" if not missing else "HOLD", missing)
    if not args.author_record.exists():
        check("author_record", "HOLD", "private author record not supplied")
        record = {}
    else:
        record = json.loads(args.author_record.read_text(encoding="utf-8"))
        check("author_record", "PASS" if record.get("authors") else "HOLD", f"authors={len(record.get('authors', []))}")

    if not missing:
        text = paths["md"].read_text(encoding="utf-8")
        cover_text = paths["cover_md"].read_text(encoding="utf-8")
        document = Document(paths["docx"])
        docx_text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        placeholders = [pattern for pattern in PLACEHOLDER_PATTERNS if any(re.search(pattern, value, flags=re.I) for value in [text, cover_text, docx_text])]
        check("placeholder_zero", "PASS" if not placeholders else "FAIL", placeholders)
        check("internal_author_gate_removed", "PASS" if "Author metadata completion gate" not in docx_text else "FAIL", "must be absent")
        check("title_identity", "PASS" if text.startswith(f"# {TITLE}\n") and document.core_properties.title == TITLE else "FAIL", document.core_properties.title)
        abstract_start = text.index("## Abstract") + len("## Abstract")
        abstract_end = text.index("\n## Background\n", abstract_start)
        abstract = text[abstract_start:abstract_end]
        abstract_heads = re.findall(r"(?m)^### (.+)$", abstract)
        abstract_words = len(re.findall(r"\b[A-Za-z0-9][A-Za-z0-9'/-]*\b", abstract))
        check("abstract", "PASS" if abstract_heads == ["Background", "Results", "Conclusions"] and abstract_words <= 350 and not re.search(r"\[\d+\]", abstract) else "FAIL", f"heads={abstract_heads} words={abstract_words}")
        references = [int(value) for value in re.findall(r"(?m)^(\d+)\. ", text)]
        first = citation_first_appearance(text)
        check("reference_list_continuity", "PASS" if references == list(range(1, 23)) else "FAIL", references)
        check("reference_first_appearance", "PASS" if first == list(range(1, 23)) else "FAIL", first)
        declarations = ["Ethics approval and consent to participate", "Consent for publication", "Availability of data and materials", "Competing interests", "Funding", "Authors' contributions", "Acknowledgements"]
        check("declaration_headings", "PASS" if all(f"### {heading}" in text for heading in declarations) else "FAIL", declarations)
        check("llm_disclosure", "PASS" if "OpenAI Codex" in text and "not treated as an author" in text else "FAIL", "supervised disclosure")
        check("tag_text", "PASS" if TAG in text else "FAIL", TAG)
        tag_commit = subprocess.run(["git", "rev-parse", f"refs/tags/{TAG}^{{}}"], cwd=REPO, check=True, capture_output=True, text=True).stdout.strip()
        check("tag_resolution", "PASS" if tag_commit == EXPECTED_TAG_COMMIT else "FAIL", tag_commit)
        pdf = PdfReader(str(paths["pdf"])); cover_pdf = PdfReader(str(paths["cover_pdf"]))
        check("wps_manuscript", "PASS" if "wps" in str(pdf.metadata.get("/Creator", "")).lower() and len(pdf.pages) > 20 else "FAIL", f"pages={len(pdf.pages)} creator={pdf.metadata.get('/Creator')}")
        check("wps_cover", "PASS" if "wps" in str(cover_pdf.metadata.get("/Creator", "")).lower() and len(cover_pdf.pages) == 1 else "FAIL", f"pages={len(cover_pdf.pages)} creator={cover_pdf.metadata.get('/Creator')}")
        receipt = json.loads(paths["render_receipt"].read_text(encoding="utf-8"))
        check("wps_version_binding", "PASS" if receipt.get("manuscript_docx_sha256") == sha256(paths["docx"]) and receipt.get("manuscript_pdf_sha256") == sha256(paths["pdf"]) and receipt.get("cover_docx_sha256") == sha256(paths["cover_docx"]) and receipt.get("cover_pdf_sha256") == sha256(paths["cover_pdf"]) and receipt.get("visual_inspection") == "PASS_ALL_PAGES" else "FAIL", receipt.get("visual_inspection"))
        check("graphical_abstract", "PASS" if Image.open(paths["graphical_abstract"]).size == (920, 300) else "FAIL", Image.open(paths["graphical_abstract"]).size)
        check("supplement_identity", "PASS" if sha256(paths["supp1"]) == sha256(UPSTREAM / "supplement" / paths["supp1"].name) and sha256(paths["supp2"]) == sha256(UPSTREAM / "supplement" / paths["supp2"].name) else "FAIL", "frozen Additional files")
        with zipfile.ZipFile(paths["supp2"]) as archive:
            bad = archive.testzip()
        check("supplement_zip_crc", "PASS" if bad is None else "FAIL", bad)

        approval = record.get("final_version_approval", {})
        check("approved_docx_hashes", "PASS" if approval.get("manuscript_docx_sha256") == sha256(paths["docx"]) and approval.get("cover_letter_docx_sha256") == sha256(paths["cover_docx"]) else "HOLD", "must bind exact final DOCX files")
        authors = record.get("authors", [])
        initials = {author.get("initials") for author in authors}
        check("all_author_approval", "PASS" if approval.get("all_authors_approved") is True and set(approval.get("approved_author_initials", [])) == initials and initials else "HOLD", approval.get("approved_author_initials"))
        qualification = record.get("institutional_journal_qualification", {})
        check("institutional_q2_confirmation", "PASS" if qualification.get("status") == "PASS" and qualification.get("evidence_reference") else "HOLD", qualification.get("status"))
        authorization = record.get("submission_authorization", {})
        corresponding = record.get("corresponding_author", {})
        check("author_operated_submission_authorization", "PASS" if authorization.get("authorized") is True and authorization.get("authorized_by_author_initials") == corresponding.get("author_initials") and authorization.get("author_operated_submission_confirmed") is True else "HOLD", authorization.get("authorized"))

    failures = [item for item in checks if item["status"] == "FAIL"]
    holds = [item for item in checks if item["status"] == "HOLD"]
    status = "FAIL" if failures else ("HOLD_HUMAN_INPUT_REQUIRED" if holds else "PASS_FINAL_SUBMISSION")
    result = {
        "mode": "FINAL_SUBMISSION",
        "candidate": str(args.candidate),
        "status": status,
        "summary": {"pass": sum(item["status"] == "PASS" for item in checks), "hold": len(holds), "fail": len(failures), "total": len(checks)},
        "checks": checks,
        "page_count_rule": "not fixed; must be WPS-rendered, version-bound and visually inspected",
        "placeholder_rule": "specific author/declaration markers only; numeric citation brackets are allowed",
    }
    args.out.parent.mkdir(parents=True, exist_ok=True)
    args.out.write_text(json.dumps(result, indent=2, ensure_ascii=False), encoding="utf-8")
    print(json.dumps(result["summary"] | {"status": status}, ensure_ascii=False))
    return 1 if failures else 0


if __name__ == "__main__":
    raise SystemExit(main())
