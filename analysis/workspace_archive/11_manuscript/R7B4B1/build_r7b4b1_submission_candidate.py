#!/usr/bin/env python3
"""Build the private R7B4B1 Human Genomics author-metadata submission candidate."""

from __future__ import annotations

import json
import re
import shutil
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(r"H:\SCI2\YR1")
UPSTREAM = ROOT / "5_manuscript" / "R7B4A_HumanGenomics_SubmissionInterface"
BASE = ROOT / "5_manuscript" / "R7B4B_FinalSubmissionCandidate"
CODE = ROOT / "2_code" / "11_manuscript" / "R7B4B1"
TITLE = "Input-matched multi-signal colocalization clarifies immune-cell regulatory assignments at primary biliary cholangitis risk loci"
AFFILIATION = (
    "School of Medicine, The Chinese University of Hong Kong, Shenzhen, "
    "2001 Longxiang Boulevard, Longgang District, Shenzhen, Guangdong 518172, China"
)


def set_cell_margins(cell, top=90, start=90, bottom=90, end=90):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for key, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{key}"))
        if node is None:
            node = OxmlElement(f"w:{key}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    paragraph._p.append(field)


def add_line_numbering(section):
    sect_pr = section._sectPr
    node = sect_pr.find(qn("w:lnNumType"))
    if node is None:
        node = OxmlElement("w:lnNumType")
        sect_pr.append(node)
    node.set(qn("w:countBy"), "1")
    node.set(qn("w:restart"), "continuous")
    node.set(qn("w:distance"), "360")


def add_rich_text(paragraph, text: str):
    parts = re.split(r"(\*\*.*?\*\*|`.*?`|\*.*?\*)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Courier New"
            run.font.size = Pt(9)
        elif part.startswith("*") and part.endswith("*"):
            run = paragraph.add_run(part[1:-1])
            run.italic = True
        else:
            paragraph.add_run(part)


def configure_styles(doc: Document):
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.first_line_indent = Cm(0.74)
    normal.paragraph_format.widow_control = True
    for name, size in (("Title", 17), ("Heading 1", 14), ("Heading 2", 12), ("Heading 3", 11)):
        style = doc.styles[name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(12 if name != "Title" else 0)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.line_spacing = 1.05
    caption = doc.styles["Figure Caption"] if "Figure Caption" in doc.styles else doc.styles.add_style("Figure Caption", WD_STYLE_TYPE.PARAGRAPH)
    caption.font.name = "Times New Roman"
    caption.font.size = Pt(9)
    caption.paragraph_format.line_spacing = 1.0
    caption.paragraph_format.space_after = Pt(6)
    caption.paragraph_format.first_line_indent = Cm(0)


def make_author_record() -> dict:
    return {
        "record_version": "R7B4B-1.0",
        "privacy": "PRIVATE_DO_NOT_COMMIT_WHEN_COMPLETED",
        "metadata_status": "AUTHOR_SUPPLIED_IDENTITY_AND_DECLARATIONS_CONFIRMED_2026-10-10",
        "authors": [
            {
                "publication_name": "Zhi Chen",
                "initials": "ZC",
                "author_order": 1,
                "corresponding": False,
                "academic_title": "MSc student in Bioinformatics",
                "affiliation_ids": ["1"],
                "orcid": "https://orcid.org/0009-0001-0072-5576",
                "orcid_verification": "AUTHOR_SUPPLIED",
                "credit_roles": [
                    "Conceptualization", "Methodology", "Software", "Formal analysis",
                    "Investigation", "Data curation", "Visualization", "Writing - original draft"
                ],
            },
            {
                "publication_name": "Teng Qi",
                "initials": "TQ",
                "author_order": 2,
                "corresponding": True,
                "academic_title": "MSc student in Bioinformatics",
                "affiliation_ids": ["1"],
                "orcid": "https://orcid.org/0009-0007-7648-4776",
                "orcid_verification": "PUBLIC_ORCID_NAME_AND_EMAIL_MATCH",
                "credit_roles": [
                    "Conceptualization", "Methodology", "Validation", "Supervision",
                    "Project administration", "Writing - review and editing"
                ],
            },
        ],
        "affiliations": [
            {
                "id": "1",
                "full_institutional_address": AFFILIATION,
                "address_source": "CUHK-Shenzhen School of Medicine and university official contact pages",
            }
        ],
        "corresponding_author": {
            "author_initials": "TQ",
            "postal_address": AFFILIATION,
            "institutional_email": "tengqi@link.cuhk.edu.cn",
            "telephone": None,
        },
        "private_contact": {
            "ZC_email": "zhichen1@link.cuhk.edu.cn",
            "TQ_email": "tengqi@link.cuhk.edu.cn",
        },
        "authors_information": {
            "ZC": (
                "ZC is an MSc student in Bioinformatics at The Chinese University of Hong Kong, Shenzhen. "
                "His research focuses on multi-omics analysis and clinical cancer research, including "
                "integrative multi-omics approaches and the tumor microenvironment. He holds a BSc in "
                "Biomedical Sciences from Queen Mary University of London and an MB in Clinical Medicine "
                "from Nanchang University."
            )
        },
        "declarations": {
            "ethics": {
                "determination": "Ethics approval not required for secondary analysis of public de-identified data",
                "institution_or_committee": "Not applicable; no new participant recruitment or identifiable data collection",
                "reference_number_or_not_applicable": "Not applicable",
                "approved_wording": "Ethics approval was not required because this study exclusively reanalysed publicly available, de-identified data and involved no new participant recruitment or identifiable data collection. The original studies reported their applicable ethics approvals and participant consent.",
                "evidence_reference": "Author confirmation dated 2026-10-10",
                "confirmed": True,
            },
            "consent_for_publication": {"statement": "Not applicable.", "confirmed": True},
            "competing_interests": {
                "statement": "The authors declare that they have no competing interests.",
                "all_authors_confirmed": True,
            },
            "research_funding": {
                "statement": "This research received no specific grant from any funding agency in the public, commercial or not-for-profit sectors.",
                "funder_role_statement": "No funder had any role in study design, data analysis, interpretation, manuscript preparation or the decision to submit.",
                "confirmed": True,
            },
            "authors_contributions": {
                "statement_using_initials": "ZC: Conceptualization, Methodology, Software, Formal analysis, Investigation, Data curation, Visualization and Writing - original draft. TQ: Conceptualization, Methodology, Validation, Supervision, Project administration and Writing - review and editing. Both authors interpreted the results and approved the final manuscript.",
                "confirmed": True,
            },
            "acknowledgements": {
                "statement_or_not_applicable": "Not applicable.",
                "permissions_confirmed": True,
            },
            "ai_disclosure": {"statement_approved": True},
            "originality": {"not_published_or_under_review_elsewhere": True},
        },
        "publication_arrangements": {
            "apc_route": "Request a discretionary APC waiver or discount at initial submission because no research funding is available; otherwise resolve through the authors or institution before acceptance",
            "waiver_or_discount_requested_at_initial_submission": True,
            "research_funding_separate_from_apc": True,
            "licence_choice": "CC BY",
            "confirmed": True,
        },
        "institutional_journal_qualification": {
            "evaluation_system": None,
            "year": None,
            "category": None,
            "status": "PENDING",
            "evidence_reference": None,
        },
        "final_version_approval": {
            "manuscript_docx_sha256": None,
            "cover_letter_docx_sha256": None,
            "approval_date": None,
            "approved_author_initials": [],
            "all_authors_approved": False,
        },
        "submission_authorization": {
            "authorized": False,
            "authorized_by_author_initials": None,
            "authorization_date": None,
            "author_operated_submission_confirmed": False,
        },
    }


def build_manuscript_markdown() -> Path:
    source = (UPSTREAM / "manuscript" / "R7B4A_HumanGenomics_manuscript.md").read_text(encoding="utf-8")
    replacements = {
        "**Authors:** [FULL AUTHOR NAMES TO BE COMPLETED]  ": (
            "**Authors:** Zhi Chen¹ (ORCID: 0009-0001-0072-5576); "
            "Teng Qi¹* (ORCID: 0009-0007-7648-4776)  "
        ),
        "**Affiliations:** [NUMBERED INSTITUTIONAL ADDRESSES TO BE COMPLETED]  ": f"**Affiliations:** ¹ {AFFILIATION}  ",
        "**Corresponding author:** [NAME, POSTAL ADDRESS, EMAIL, AND TELEPHONE TO BE COMPLETED]": (
            f"**Corresponding author:** Teng Qi, {AFFILIATION}; tengqi@link.cuhk.edu.cn"
        ),
        "This study used previously published, de-identified datasets. The original studies reported their applicable ethics approvals and participant consent. [LOCAL INSTITUTIONAL DETERMINATION, EXEMPTION OR WAIVER WORDING TO BE CONFIRMED BY THE AUTHORS.]": (
            "Ethics approval was not required because this study exclusively reanalysed publicly available, de-identified data and involved no new participant recruitment or identifiable data collection. "
            "The original studies reported their applicable ethics approvals and participant consent."
        ),
        "Not applicable to the de-identified secondary analyses. [AUTHOR CONFIRMATION REQUIRED.]": "Not applicable.",
        "[TO BE COMPLETED FROM VERIFIED AUTHOR INFORMATION.]": "__DECLARATION_PLACEHOLDER__",
        "[CRediT CONTRIBUTIONS TO BE COMPLETED AFTER AUTHOR LIST AND ROLES ARE VERIFIED.]": (
            "ZC: Conceptualization, Methodology, Software, Formal analysis, Investigation, Data curation, Visualization and Writing - original draft. "
            "TQ: Conceptualization, Methodology, Validation, Supervision, Project administration and Writing - review and editing. "
            "Both authors interpreted the results and approved the final manuscript."
        ),
        "We thank the participants and investigators of the PBC GWAS, OneK1K, TenK10K, FinnGen and HRA008003 resources. [ADDITIONAL ACKNOWLEDGEMENTS TO BE COMPLETED.]": (
            "Not applicable."
        ),
        "[OPTIONAL AUTHOR BIOGRAPHICAL INFORMATION TO BE COMPLETED IF REQUIRED BY THE TARGET JOURNAL.]": (
            "ZC is an MSc student in Bioinformatics at The Chinese University of Hong Kong, Shenzhen. "
            "His research focuses on multi-omics analysis and clinical cancer research, including integrative multi-omics approaches and the tumor microenvironment. "
            "He holds a BSc in Biomedical Sciences from Queen Mary University of London and an MB in Clinical Medicine from Nanchang University."
        ),
    }
    for old, new in replacements.items():
        if old not in source:
            raise RuntimeError(f"Expected source text not found: {old[:80]}")
        source = source.replace(old, new)
    declaration_values = [
        "The authors declare that they have no competing interests.",
        "This research received no specific grant from any funding agency in the public, commercial or not-for-profit sectors. No funder had any role in study design, data analysis, interpretation, manuscript preparation or the decision to submit.",
    ]
    for value in declaration_values:
        if "__DECLARATION_PLACEHOLDER__" not in source:
            raise RuntimeError("Declaration placeholder count changed")
        source = source.replace("__DECLARATION_PLACEHOLDER__", value, 1)
    out = BASE / "manuscript" / "R7B4B_HumanGenomics_final.md"
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(source, encoding="utf-8")
    return out


def build_cover_markdown() -> Path:
    source = (UPSTREAM / "cover_letter" / "R7B4A_HumanGenomics_cover_letter_DRAFT.md").read_text(encoding="utf-8")
    source = source.replace(
        "**From:** [CORRESPONDING AUTHOR NAME, DEGREE, AFFILIATION, POSTAL ADDRESS, EMAIL, TELEPHONE]",
        f"**From:** Teng Qi, MSc student in Bioinformatics, {AFFILIATION}; tengqi@link.cuhk.edu.cn",
    )
    source = source.replace(
        "[CORRESPONDING AUTHOR NAME AND SIGNATURE BLOCK]",
        f"Teng Qi\n\nCorresponding author\n\n{AFFILIATION}\n\ntengqi@link.cuhk.edu.cn",
    )
    confirmation_block = """Before submission, the corresponding author must replace and confirm the following statements:

- [ALL AUTHORS APPROVE THIS VERSION AND ITS SUBMISSION.]
- [THE MANUSCRIPT IS ORIGINAL AND IS NOT UNDER CONSIDERATION ELSEWHERE.]
- [ALL FINANCIAL AND NON-FINANCIAL COMPETING INTERESTS ARE DISCLOSED.]
- [ETHICS/WAIVER WORDING HAS BEEN CONFIRMED BY THE RESPONSIBLE INSTITUTION.]
- [FUNDING, AUTHOR CONTRIBUTIONS, ACKNOWLEDGEMENTS, APC ROUTE AND LICENCE ARE CONFIRMED.]
"""
    confirmed_block = """Both authors have approved the manuscript, figures, supplementary materials and this cover letter. The work is original, has not been published, and is not under consideration by another journal. The authors declare no competing interests and no research funding. Ethics approval was not required because the study exclusively reanalysed publicly available, de-identified data without new participant recruitment or identifiable data collection. The authors approve the disclosed, supervised use of OpenAI Codex with the locally installed QiTeng Academic Writing Skill for argument structuring, language editing, file generation and consistency checking; generative AI was not treated as an author or a source of scientific evidence.
"""
    if confirmation_block not in source:
        raise RuntimeError("Cover-letter confirmation block changed")
    source = source.replace(confirmation_block, confirmed_block)
    out = BASE / "cover_letter" / "R7B4B_HumanGenomics_cover_letter_final.md"
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(source, encoding="utf-8")
    return out


def figure_path(number: int) -> Path:
    names = {
        1: "Figure1_study_scope_and_model_integrity.png",
        2: "Figure2_input_matched_reclassification.png",
        3: "Figure3_simulation_discovery_and_inference.png",
        4: "Figure4_IL12RB2_external_evidence.png",
        5: "Figure5_FCRL3_support_and_counterexample.png",
        6: "Figure6_liver_tissue_boundary.png",
    }
    return BASE / "figures" / names[number]


def build_manuscript_docx(md_path: Path) -> Path:
    text = md_path.read_text(encoding="utf-8")
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.2)
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.header_distance = Cm(1)
    section.footer_distance = Cm(1)
    add_line_numbering(section)
    add_page_number(section.footer.paragraphs[0])
    configure_styles(doc)
    props = doc.core_properties
    props.title = TITLE
    props.subject = "Human Genomics Research submission candidate pending author declarations"
    props.author = "Zhi Chen; Teng Qi"
    props.keywords = "PBC; colocalization; eQTL; SuSiE; source-matched LD; multi-signal"

    in_figure_legends = False
    for line in text.splitlines():
        raw = line.rstrip()
        if not raw:
            continue
        if raw.startswith("# "):
            paragraph = doc.add_paragraph(style="Title")
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.first_line_indent = Cm(0)
            add_rich_text(paragraph, raw[2:])
            continue
        if raw == "## Figure legends":
            doc.add_page_break()
            doc.add_heading("Figure legends", level=1)
            in_figure_legends = True
            continue
        if raw.startswith("## "):
            doc.add_heading(raw[3:], level=1)
            continue
        if raw.startswith("### "):
            doc.add_heading(raw[4:], level=2)
            continue
        if in_figure_legends and raw.startswith("**Figure "):
            match = re.match(r"\*\*Figure (\d+)\.", raw)
            if match:
                number = int(match.group(1))
                if number > 1:
                    doc.add_page_break()
                image_paragraph = doc.add_paragraph()
                image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                image_paragraph.paragraph_format.first_line_indent = Cm(0)
                image_paragraph.add_run().add_picture(str(figure_path(number)), width=Inches(6.45))
                caption = doc.add_paragraph(style="Figure Caption")
                add_rich_text(caption, raw)
                continue
        if in_figure_legends and raw.startswith("**Abbreviations:**"):
            paragraph = doc.add_paragraph(style="Figure Caption")
            add_rich_text(paragraph, raw)
            continue
        paragraph = doc.add_paragraph()
        if raw.startswith(("**Authors:", "**Affiliations:", "**Corresponding author:", "**Article type:")):
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.first_line_indent = Cm(0)
            paragraph.paragraph_format.line_spacing = 1.15
        if raw.startswith("G* =") or raw.startswith("and R_QTL"):
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.first_line_indent = Cm(0)
        add_rich_text(paragraph, raw.rstrip("  "))

    out = BASE / "manuscript" / "R7B4B_HumanGenomics_final.docx"
    doc.save(out)
    return out


def build_cover_docx(md_path: Path) -> Path:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(1.65)
    section.bottom_margin = Cm(1.65)
    section.left_margin = Cm(2.1)
    section.right_margin = Cm(2.1)
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(3)
    normal.paragraph_format.line_spacing = 1.0
    for raw in md_path.read_text(encoding="utf-8").splitlines():
        line = raw.strip()
        if not line:
            continue
        if line.startswith("# "):
            paragraph = doc.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_after = Pt(4)
            run = paragraph.add_run(line[2:].replace("—", "–"))
            run.bold = True
            run.font.name = "Arial"
            run.font.size = Pt(13)
            run.font.color.rgb = RGBColor(0, 0, 0)
        elif line.startswith("- "):
            paragraph = doc.add_paragraph(style="List Bullet")
            paragraph.add_run(line[2:])
            paragraph.paragraph_format.left_indent = Cm(0.45)
            paragraph.paragraph_format.first_line_indent = Cm(-0.25)
            paragraph.paragraph_format.space_after = Pt(0)
            for run in paragraph.runs:
                run.font.size = Pt(9.5)
        else:
            clean = line.replace("**", "")
            clean = re.sub(r"\*([^*]+)\*", r"\1", clean).replace("`", "")
            paragraph = doc.add_paragraph(clean)
            if line.startswith("**"):
                paragraph.runs[0].bold = True
    add_page_number(section.footer.paragraphs[0])
    doc.core_properties.title = "Cover letter for Human Genomics"
    doc.core_properties.author = "Teng Qi"
    doc.core_properties.comments = "Author-metadata candidate; declarations and submission authorization pending."
    out = BASE / "cover_letter" / "R7B4B_HumanGenomics_cover_letter_final.docx"
    doc.save(out)
    return out


def main() -> None:
    for folder in ("figures", "graphical_abstract", "supplement", "qa", "private", "submission"):
        (BASE / folder).mkdir(parents=True, exist_ok=True)
    for source in (UPSTREAM / "figures").glob("*"):
        if source.is_file():
            shutil.copy2(source, BASE / "figures" / source.name)
    shutil.copy2(
        UPSTREAM / "graphical_abstract" / "Graphical_Abstract_HumanGenomics_920x300.png",
        BASE / "graphical_abstract" / "Graphical_Abstract_HumanGenomics_920x300.png",
    )
    for source in (UPSTREAM / "supplement").glob("*"):
        if source.is_file():
            shutil.copy2(source, BASE / "supplement" / source.name)

    record = make_author_record()
    record_path = BASE / "private" / "R7B4B_author_record.PRIVATE.json"
    record_path.write_text(json.dumps(record, indent=2, ensure_ascii=False), encoding="utf-8")
    md = build_manuscript_markdown()
    cover_md = build_cover_markdown()
    docx = build_manuscript_docx(md)
    cover_docx = build_cover_docx(cover_md)
    import hashlib
    def digest(path: Path) -> str:
        h = hashlib.sha256()
        with path.open("rb") as handle:
            for block in iter(lambda: handle.read(1024 * 1024), b""):
                h.update(block)
        return h.hexdigest()
    record["final_version_approval"] = {
        "manuscript_docx_sha256": digest(docx),
        "cover_letter_docx_sha256": digest(cover_docx),
        "approval_date": "2026-10-10",
        "approved_author_initials": ["ZC", "TQ"],
        "all_authors_approved": True,
        "approval_basis": "User confirmed authorship information and approved the manuscript, supplements, figures, data, cover letter, originality, single submission and AI-use disclosure on 2026-10-10.",
    }
    record_path.write_text(json.dumps(record, indent=2, ensure_ascii=False), encoding="utf-8")
    state = {
        "stage": "R7B4B1_AUTHOR_METADATA_INJECTED_SUBMISSION_CANDIDATE",
        "status": "READY_FOR_WPS_AND_FINAL_QA; JOURNAL_SUBMISSION_MUST_BE_AUTHOR_OPERATED",
        "authors": ["Zhi Chen", "Teng Qi"],
        "corresponding_author": "Teng Qi",
        "private_record": str(record_path),
        "manuscript_docx": str(docx),
        "cover_letter_docx": str(cover_docx),
        "public_git_policy": "Do not commit completed author record, private emails, declarations evidence or final approval receipts.",
    }
    (BASE / "qa" / "R7B4B1_build_state.json").write_text(json.dumps(state, indent=2, ensure_ascii=False), encoding="utf-8")
    print(json.dumps(state, indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()
