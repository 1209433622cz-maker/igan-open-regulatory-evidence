"""Independent table recomputation, document invariants and WPS PDF rasterization."""
from pathlib import Path
import collections,csv,gzip,hashlib,importlib.util,io,json,re,statistics,zipfile
import pymupdf as fitz
from PIL import Image,ImageDraw
from docx import Document
ROOT=Path(r'H:\SCI2\YR1')
OUT=ROOT/'5_manuscript/R7C3B_QiTeng_WPS_SubmissionRefinement_20261011'
SRC=ROOT/'0_admin/intake/R7C3A_20261011'
checks=[]
def check(name,test,evidence):
    checks.append({'check':name,'status':'PASS' if test else 'FAIL','evidence':evidence})
def table(z,n):
    b=z.read(n)
    if n.endswith('.gz'):b=gzip.decompress(b)
    return list(csv.DictReader(io.StringIO(b.decode('utf-8-sig')),delimiter='\t'))
def counter(rows,key):return dict(collections.Counter(r[key] for r in rows))
def numeq(x,v,tol=1e-8):return abs(float(x)-v)<tol
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
with zipfile.ZipFile(OUT/'05_Source_Data/Source_Data_1_Primary_Machine_Readable_Data.zip') as z:
    registry=table(z,'S2/S2_PBCwide_6923_screening_registry.tsv.gz')
    check('screen_registry_rows',len(registry)==6923,len(registry))
    counts=counter(registry,'classification')
    check('screen_eligible_coverage',counts['INSUFFICIENT_VARIANT_OVERLAP']==1463 and len(registry)-1463==5460,counts)
    pf=table(z,'S5/S5_642_PF10_PF50_classifications.tsv')
    check('multisignal_642_state_counts',counter(pf,'PF10_multisignal_state')=={'H4_SUPPORTED_STABLE':92,'H3_SUPPORTED_STABLE':428,'MODEL_SENSITIVE':59,'UNINFORMATIVE':63},counter(pf,'PF10_multisignal_state'))
    traj=table(z,'S3/S3_642_A0_A1_A2_M_trajectories.tsv')
    counts={k:sum(r[k]=='True' for r in traj) for k in ['A0_to_A1_changed','A1_to_A2_changed','A2_to_M_changed']}
    check('matched_arm_changes_15_19_139',list(counts.values())==[15,19,139],counts)
    mat=collections.Counter((r['A2_simple_state'],r['M_simple_state']) for r in traj)
    check('matched_strict_reversals_8_10',mat['H4','H3']==8 and mat['H3','H4']==10,{'H4_to_H3':mat['H4','H3'],'H3_to_H4':mat['H3','H4']})
    a2=counter(traj,'A2_simple_state')
    check('matched_denominators_113_452',a2['H4']==113 and a2['H3']==452,a2)
    rm=[int(r['removed_n']) for r in traj]
    check('variant_removal_19_114_median65',(min(rm),max(rm),statistics.median(rm))==(19,114,65),(min(rm),max(rm),statistics.median(rm)))
    check('subset_47_loci_200_genes_14_cells',tuple(len({r[k] for r in traj}) for k in ['locus_index','gene','cell_type'])==(47,200,14),tuple(len({r[k] for r in traj}) for k in ['locus_index','gene','cell_type']))
    check('fixed_184_455_3_roles',sorted(counter(traj,'R7B1B_v2_role').values())==[3,184,455],counter(traj,'R7B1B_v2_role'))
    definite=[r for r in pf if r['PF10_multisignal_state'] in ['H4_SUPPORTED_STABLE','H3_SUPPORTED_STABLE']]
    sensitivity=counter(definite,'PF10_PF50_sensitivity')
    check('PF50_conditional_denominator',len(definite)==520 and sensitivity.get('MATCH')==491 and sorted(sensitivity.values())==[4,25,491],sensitivity)
    sem=table(z,'S6/S6_signal_semantics_642.tsv')
    stable=[r for r in sem if r['historical_PF10_multisignal_state'] in ['H4_SUPPORTED_STABLE','H3_SUPPORTED_STABLE']]
    mixed=[r for r in sem if r['historical_PF10_multisignal_state']=='H4_SUPPORTED_STABLE' and r['mixed_h4_h3_within_any_config']=='True']
    check('pair_identity_520_and_mixed32',len(stable)==520 and all(r['same_signal_pair_confirmed']=='True' for r in stable) and len(mixed)==32,{'stable':len(stable),'mixed_H4_H3':len(mixed)})
    channels=table(z,'S7/S7_channel_summary.tsv')
    primary={r['matched_channel']:int(r['n']) for r in channels if r['scope']=='PRIMARY_MATCHED_OVERALL'}
    evaluable=sum(primary.get(k,0) for k in ['PAIR_H3','PAIR_H4','PAIR_UNINFORMATIVE'])
    check('simulation_486000_and160782',sum(primary.values())==486000 and evaluable==160782,primary)
    scenarios=table(z,'S7/S7_scenario_summary.tsv')
    distinct=[r for r in scenarios if r['true_shared']=='False'];shared=[r for r in scenarios if r['true_shared']=='True']
    rates=[statistics.mean(float(r[k])*100 for r in rows) for rows in [distinct,shared] for k in ['abf_h4_rate','matched_h4_rate']]
    check('simulation_tradeoff_rates',list(map(lambda v:round(v,2),rates))==[12.41,8.56,32.99,26.66],rates)
    tissue=table(z,'S9/S9_exact_5vs5_tissue.tsv')
    check('liver_5_vs5_detectability',len(tissue)==2 and all(all(r[k]=='5' for k in ['PBC_n','control_n','PBC_positive_donors','control_positive_donors']) for r in tissue),tissue)
    check('liver_exact_p_q',numeq(tissue[0]['exact_label_permutation_p'],1/9) and numeq(tissue[1]['exact_label_permutation_p'],4/63) and all(numeq(r['BH_q_two_prespecified_targets'],1/9) for r in tissue),{'p':[r['exact_label_permutation_p'] for r in tissue],'q':[r['BH_q_two_prespecified_targets'] for r in tissue]})
    target=[r for r in pf if r['gene']=='FCRL3' and r['cell_type']=='CD8_ET']
    check('FCRL3_CD8_model_counterexample',len(target)==1 and numeq(target[0]['pf10_l10_best_h4'],.003636,1e-5) and numeq(target[0]['pf50_l10_best_h4'],.991894,1e-5),target)
with zipfile.ZipFile(OUT/'05_Source_Data/Source_Data_2_External_Regulatory_Evidence.zip') as z:
    ext=table(z,'data/R7C0_FinnGenR13_x_OneK_NK_IL12RB2_coloc.tsv')
    check('external_1175_prior_posteriors',all(r['n_variants']=='1175' for r in ext) and [round(float(r['PP_H4']),4) for r in ext]==[.9772,.9977,.9998],ext)
science=list(csv.DictReader((OUT/'06_QA/Scientific_Source_Data_48_Identity.tsv').open(encoding='utf-8-sig'),delimiter='\t'))
check('scientific_member_48_byte_identity',len(science)==48 and all(r['independent_status']=='PASS' for r in science),len(science))
check('workbook_byte_identity',sha(OUT/'03_Supplement/Supplementary_Data_1_Tables_S1-S10.xlsx')==sha(SRC/'03_Supplement/Supplementary_Data_1_Tables_S1-S10.xlsx'),sha(OUT/'03_Supplement/Supplementary_Data_1_Tables_S1-S10.xlsx'))
figs=json.loads((OUT/'06_QA/Figure_Production_Derivative_Identity.json').read_text('utf-8'))
check('all_six_figure_delta_restricted_to_panel_letters',len(figs)==6 and all(r['outside_label_raster_delta']=='ZERO' for r in figs),figs)

# Execute actual skill QA; report semantic review separately, not as 570 executable tests.
spec=importlib.util.spec_from_file_location('qiteng_qa',ROOT/'0_admin/skills/qiteng_runtime_v03242/04_EXECUTABLE_QA/qiteng_qa.py')
mod=importlib.util.module_from_spec(spec);spec.loader.exec_module(mod)
docreports=[]
for rel in ['01_Manuscript/R7C3B_npj_Manuscript.docx','01_Manuscript/R7C3B_AuthorReview_InlineFigures.docx','03_Supplement/R7C3B_Supplementary_Information.docx','04_Cover/R7C3B_Cover_Letter_Candidate.docx']:
    path=OUT/rel;structure=mod.docx_structure_qa(path)
    check('clean_docx_'+path.stem,not structure['comments_present'] and structure['tracked_insertions']==0 and structure['tracked_deletions']==0,structure)
    if '/01_Manuscript' in str(path).replace('\\','/'):
        citation=mod.docx_citation_qa(path);docreports.append({'file':rel,'structure':structure,'citation':citation})
        check('six_scientific_objects_'+path.stem,structure['drawing_elements']==6,structure['drawing_elements'])
        check('references_and_first_use_'+path.stem,citation['reference_count']==19 and citation['unique_cited_count']==19 and citation['reference_continuity'] and citation['first_appearance_order_pass'] and not citation['missing_reference_numbers'] and not citation['orphan_reference_numbers'],citation)
        d=Document(path);text='\n'.join(p.text for p in d.paragraphs)
        check('abstract_title_limit_'+path.stem,len(d.paragraphs[0].text.split())==12 and len(d.paragraphs[6].text.split())<=150,{'title':len(d.paragraphs[0].text.split()),'abstract':len(d.paragraphs[6].text.split())})
        h1=[p.text for p in d.paragraphs if p.style.name=='Heading 1']
        check('npj_section_order_'+path.stem,[x for x in h1 if x in ['Abstract','Introduction','Results','Discussion','Methods']]==['Abstract','Introduction','Results','Discussion','Methods'],h1)
        check('no_inherited_fresh_approval_'+path.stem,'approved the final manuscript' not in text,text.count('approved'))
        check('essential_liver_reproduction_visible_'+path.stem,all(s in text for s in ['6,000','99th-percentile','CD79A','XCL2','T-panel count','252','Benjamini-Hochberg']),True)
        captions=[p.text for p in d.paragraphs if p.text.startswith('Figure ') and re.match(r'^Figure \d\.',p.text)]
        check('six_legends_under_350_'+path.stem,len(captions)==6 and max(len(s.split()) for s in captions)<=350,[len(s.split()) for s in captions])
        check('Figure6_boxplot_semantics_'+path.stem,'1.5 times the interquartile range' in text and 'whiskers' in text,True)
        check('citation_source_conservation_'+path.stem,len(list(csv.DictReader((OUT/'06_QA/Citation_Source_Conservation.tsv').open(encoding='utf-8-sig'),delimiter='\t')))==25,{'numbered':19,'resource_URLs':6})
        check('volume_years_2020_2025_'+path.stem,'17, 93-110 (2020)' in text and '49, e22592 (2025)' in text, 'Print/volume year; earlier online dates retained in metadata receipts')
        cp=next(p.text for p in d.paragraphs if p.text.startswith('Analysis code, frozen'))
        check('no_duplicate_code_URL_'+path.stem,cp.count('github.com')==1,cp.count('github.com'))
        for token in ['0.9977','32.99%','26.66%','12.41%','8.56%','0.111','0.0635','92','428','59','63','PF50','642']:
            check('claim_token_'+path.stem+'_'+token,token in text,token)

# PDF rasterization and page-level technical parity, based on final WPS exports.
pdfrows=[];pagechecks=[]
for path in sorted(OUT.rglob('*.pdf')):
    if '09_Archive' in str(path) or '02_Figures' in str(path):continue
    d=fitz.open(path);relative=str(path.relative_to(OUT)).replace('\\','/')
    check('wps_renderer_'+path.stem,'WPS' in d.metadata.get('creator',''),d.metadata)
    target=OUT/'06_QA/pages'/path.stem;target.mkdir(parents=True,exist_ok=True)
    # Only current page files are counted; remove stale same-stage render derivatives safely.
    for stale in target.glob('page_*.png'):stale.unlink()
    pdftext='\n'.join(p.get_text() for p in d)
    for i,page in enumerate(d):
        page.get_pixmap(matrix=fitz.Matrix(1.5,1.5),alpha=False).save(target/f'page_{i+1:03d}.png')
        outliers=[]
        for block in page.get_text('blocks'):
            if block[6]==0 and (block[0]<-1 or block[1]<-1 or block[2]>page.rect.width+1 or block[3]>page.rect.height+1):outliers.append(block[:4])
        pagechecks.append({'file':relative,'page':i+1,'text_blocks_outside_page':len(outliers),'technical_status':'PASS' if not outliers else 'FAIL','visual_status':'PENDING_HUMAN_MODEL_INSPECTION'})
    # Four-page sheets for a complete layout sweep, individual pages are also retained at 150%.
    for stale in (OUT/'06_QA').glob(path.stem+'_pages_*.jpg'):stale.unlink()
    for start in range(0,len(d),4):
        sheet=Image.new('RGB',(1228,1760),'white')
        for i in range(start,min(len(d),start+4)):
            im=Image.open(target/f'page_{i+1:03d}.png');im.thumbnail((594,842))
            tile=Image.new('RGB',(614,880),'#dde0e4');tile.paste(im,((614-im.width)//2,22));ImageDraw.Draw(tile).text((8,860),f'Page {i+1}',fill='black')
            j=i-start;sheet.paste(tile,((j%2)*614,(j//2)*880))
        sheet.save(OUT/'06_QA'/f'{path.stem}_pages_{start+1:02d}-{min(start+4,len(d)):02d}.jpg',quality=94)
    mapping={}
    if '/01_Manuscript' in '/'+relative:
        for n in range(1,7):
            imagepages=[i+1 for i,p in enumerate(d) if len(p.get_images())>0 and f'Figure {n}.' in p.get_text()]
            mapping[n]=imagepages
        check('six_figures_same_page_legends_'+path.stem,all(len(v)==1 for v in mapping.values()),mapping)
        complete=[]
        for n,pages in mapping.items():
            if len(pages)!=1:complete.append(False);continue
            page=d[pages[0]-1];pt=page.get_text()
            complete.append(len(page.get_images())==1 and 'Abbreviations:' in pt and ('interquartile range' in pt if n==6 else True))
        check('complete_legend_abbreviations_single_image_'+path.stem,all(complete),complete)
        normalized=lambda s:re.sub(r'\s+','',s).replace('−','-').replace('–','-')
        docx=Document(path.with_suffix('.docx'));body=normalized('\n'.join(p.text for p in docx.paragraphs))
        # WPS discretionary line-break hyphenation and typography make full byte-text identity inappropriate.
        check('pdf_critical_claim_parity_'+path.stem,all(token in normalized(pdftext) for token in ['6,923','5,460','642','486,000','0.9977','PF50','FCRL3','IL12RB2']),True)
    pdfrows.append({'file':relative,'pages':len(d),'metadata':d.metadata,'figure_page_mapping':mapping,'sha256':sha(path)})
check('all_pdf_blocks_within_pages',all(r['technical_status']=='PASS' for r in pagechecks),len(pagechecks))
with (OUT/'06_QA/Page_Level_WPS_QA.tsv').open('w',encoding='utf-8-sig',newline='') as f:
    w=csv.DictWriter(f,fieldnames=list(pagechecks[0]),delimiter='\t');w.writeheader();w.writerows(pagechecks)
for name,obj in [('Numerical_Document_QA.json',{'checks':checks,'total':len(checks),'passes':sum(r['status']=='PASS' for r in checks),'failures':[r for r in checks if r['status']=='FAIL'],'semantic_review':'separate documented review; not executable unit tests'}),('QiTeng_Executed_DOCX_QA.json',docreports),('PDF_Metadata_Pages.json',pdfrows)]:
    (OUT/'06_QA'/name).write_text(json.dumps(obj,ensure_ascii=False,indent=2),encoding='utf-8')
print('QA',sum(r['status']=='PASS' for r in checks),'/',len(checks))
for row in checks:
    if row['status']=='FAIL':print('FAIL',row)
for row in pdfrows:print('PDF',row['file'],row['pages'],row['figure_page_mapping'])
if any(r['status']=='FAIL' for r in checks):raise SystemExit(1)
