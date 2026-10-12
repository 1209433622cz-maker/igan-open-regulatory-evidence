from pathlib import Path
import csv, hashlib, json, re, zipfile
from docx import Document
import pymupdf
from lxml import etree

ROOT=Path(r'H:\SCI2\YR1')
SRC=ROOT/'5_manuscript/R7C3H_StatisticalReproducibilityClosure_20261013'
OUT=ROOT/'5_manuscript/R7C3H1B_ExactVersionHandoff_20261013'
WORK=ROOT/'5_analysis/R7C3H1B_20261013'
checks=[]
def sha(p): return hashlib.sha256(Path(p).read_bytes()).hexdigest()
def check(name, value, expected):
    checks.append(dict(check=name, actual=value, expected=expected, passed=value==expected))
baseline=json.loads((OUT/'06_QA/BASELINE_INPUT_IDENTITY.json').read_text())
check('Original_R7C3H_files_unchanged',sum(sha(SRC/r['relative'])==r['sha256'] for r in baseline),115)
doc=Document(OUT/'01_Manuscript/Manuscript_SubmissionCandidate.docx')
old=Document(SRC/'01_Manuscript/Manuscript_ExpertRebuilt.docx')
before=[p.text for p in old.paragraphs];after=[p.text for p in doc.paragraphs]
check('Main_paragraph_count_unchanged',len(after),len(before))
ledger=json.loads((OUT/'06_QA/EDITORIAL_CHANGE_LEDGER.json').read_text(encoding='utf-8'))
expected=before.copy()
for change in ledger[:3]:
    expected=[p.replace(change['before'],change['after']) for p in expected]
def normalize_notation(s):
    return re.sub(r'(?:[x×])?10\^[-+]?\d+|S_[AM]',lambda m:m.group().replace('^','').replace('_','').replace('x','×'),s)
eq=json.loads((OUT/'06_QA/NATIVE_EQUATION_LEDGER.json').read_text(encoding='utf-8'))
expected_text=[normalize_notation(p) if p not in [r['before'] for r in eq['equations']] else '' for p in expected]
check('Only_documented_main_text_notation_equation_edits',after,expected_text)
check('Editorially_changed_paragraphs_before_notation',sum(a!=b for a,b in zip(expected,before)),4)
typo=json.loads((OUT/'06_QA/TYPOGRAPHY_CHANGE_LEDGER.json').read_text(encoding='utf-8'))
check('Raw_power_and_support_notation_removed',bool(re.search(r'10\^[-+]?\d+|S_[AM]','\n'.join(after))),False)
check('Native_exponent_and_subscript_runs_added',len(doc.element.xpath('.//w:vertAlign')),len(old.element.xpath('.//w:vertAlign'))+len(typo['changes']))
text='\n'.join(after)
check('Old_release_links_removed','r7c3f-full-audit-rebuilt-submission-2026-10-12' in text,False)
check('Precise_scientific_tree_links',text.count('/tree/r7c3h0-statistical-reproducibility-closure-2026-10-13'),2)
check('Precise_scientific_release_link',text.count('/releases/tag/r7c3h0-statistical-reproducibility-closure-2026-10-13'),1)
paragraph=next(p for p in after if p.startswith('FCRL3 showed stable OneK1K'))
check('FCRL3_TenK_ref_is_Figure3d','Figure 3d' in paragraph and 'Figure 5b' not in paragraph,True)
check('FinnGen_CS_Figure5b_retained',any('both members of the OneK QTL credible set' in p and 'Figure 5b' in p for p in after),True)
check('All_source_archives_declared','Source Data 1-3 provide' in text and 'Source Data 3 contains' in text,True)
title=after[0];abstract=after[after.index('Abstract')+1]
token=lambda s: len(re.findall(r"[A-Za-z0-9]+(?:[.'’\-][A-Za-z0-9]+)*",s))
check('Title_word_count',token(title),12)
check('Abstract_within_150',token(abstract)<=150,True)
check('Abstract_unchanged',abstract,before[before.index('Abstract')+1])
legends=[p.text for p in doc.paragraphs if p.text.startswith(tuple(f'Figure {i}. ' for i in range(1,7)))]
check('Six_full_legends',len(legends),6)
check('All_legend_counts_within_350',all(token(p)<=350 for p in legends),True)
check('Legends_scientific_text_unchanged',legends,[normalize_notation(p.text) for p in old.paragraphs if p.text.startswith(tuple(f'Figure {i}. ' for i in range(1,7)))])
def protected_zipped(p):
    with zipfile.ZipFile(p) as z:
        return {n:hashlib.sha256(z.read(n)).hexdigest() for n in z.namelist() if n not in ['word/document.xml','word/_rels/document.xml.rels']}
check('OOXML_math_images_styles_other_parts_unchanged',protected_zipped(OUT/'01_Manuscript/Manuscript_SubmissionCandidate.docx'),protected_zipped(SRC/'01_Manuscript/Manuscript_ExpertRebuilt.docx'))
check('Math_object_count',len(doc.element.xpath('.//m:oMath')),len(old.element.xpath('.//m:oMath'))+2)
check('Native_equation_OMML_identity',[hashlib.sha256(etree.tostring(m)).hexdigest() for m in doc.element.xpath('.//m:oMath')],[r['omml_sha256'] for r in eq['equations']])
check('Track_changes_absent',len(doc.element.xpath('.//w:ins | .//w:del')),0)
check('Authors_correspondence_ORCIDs',all(s in text for s in ['Zhi Chen','Teng Qi','tengqi@link.cuhk.edu.cn','0009-0001-0072-5576','0009-0007-7648-4776']),True)
check('Funding_conflict_AI_ethics_declarations',all(s in text for s in ['no competing interests','no specific grant','OpenAI Codex','Ethics approval was not required']),True)
for folder,number in [('02_Figures',27),('05_Source_Data',3)]:
    candidates=[p for p in (SRC/folder).glob('*') if p.is_file()]
    check(folder+'_unchanged_files',sum(sha(p)==sha(OUT/folder/p.name) for p in candidates),number)
suppold=Document(SRC/'03_Supplement/Supplementary_Information_ExpertRebuilt.docx')
suppnew=Document(OUT/'03_Supplement/Supplementary_Information_ExpertRebuilt.docx')
check('Supplement_only_native_notation_edits',[p.text for p in suppnew.paragraphs],[normalize_notation(p.text) for p in suppold.paragraphs])
check('Supplement_math_figures_styles_other_parts_unchanged',protected_zipped(OUT/'03_Supplement/Supplementary_Information_ExpertRebuilt.docx'),protected_zipped(SRC/'03_Supplement/Supplementary_Information_ExpertRebuilt.docx'))
check('Supplement_workbook_and_crosswalk_unchanged',sum(sha(SRC/'03_Supplement'/name)==sha(OUT/'03_Supplement'/name) for name in ['Supplementary_Data_1_Tables_S1-S10.xlsx','Supplementary_Editorial_Panel_Crosswalk.tsv']),2)
suppwps=json.loads((OUT/'06_QA/WPS_SUPPLEMENT_EXPORT_RECEIPT.json').read_text(encoding='utf-8-sig'))
check('Supplement_WPS_DOCX_and_PDF_identity',all(sha(OUT/'03_Supplement'/('Supplementary_Information_ExpertRebuilt.'+ext))==suppwps[key] for ext,key in [('docx','docx_sha256'),('pdf','pdf_sha256')]),True)
check('Cover_letter_unchanged',sha(OUT/'04_Cover/Cover_Letter_SubmissionCandidate.docx'),sha(SRC/'04_Cover/Cover_Letter_ExpertRebuilt_DRAFT.docx'))
with open(OUT/'06_QA/PANEL_EVIDENCE_MAP.tsv',encoding='utf-8-sig',newline='') as f: panels=list(csv.DictReader(f,delimiter='\t'))
check('Panel_map_count',len(panels),34)
row=next(r for r in panels if r['figure']=='3' and r['panel']=='d')
check('Figure3d_public_locator',row['source_member'].startswith('https://github.com/') and '4fb58180c23af48e337e9a14e4073d9071e44d16' in row['source_member'],True)
check('Figure3d_public_source_identity',sha(OUT/'05_Source_Data/R7A1C_FCRL3_TenK_Bintermediate_coloc.tsv'),ledger[3]['source_sha256'])
archives=[]
for p in sorted((OUT/'05_Source_Data').glob('*.zip')):
    with zipfile.ZipFile(p) as z:
        check(p.name+'_CRC',z.testzip(),None)
        lines=z.read('checksums.sha256').decode('utf-8-sig').splitlines()
        n=0
        for line in lines:
            if not line.strip():continue
            h,rel=line.split(None,1);rel=rel.strip().lstrip('*')
            assert hashlib.sha256(z.read(rel)).hexdigest()==h,rel
            n+=1
        archives.append(dict(archive=p.name,passed=n))
check('Three_source_archives_internal_checksums',sum(a['passed'] for a in archives),60)
wps=json.loads((OUT/'06_QA/WPS_EXPORT_RECEIPT.json').read_text(encoding='utf-8-sig'))
check('WPS_readonly_DOCX_identity',sha(OUT/'01_Manuscript/Manuscript_SubmissionCandidate.docx'),wps['docx_sha256'])
check('WPS_PDF_identity',sha(OUT/'01_Manuscript/Manuscript_SubmissionCandidate.pdf'),wps['pdf_sha256'])
upload=[('Main Manuscript','01_Manuscript/Manuscript_SubmissionCandidate.docx')]
upload += [('Figure '+str(i),f'02_Figures/Figure_{i}.pdf') for i in range(1,7)]
upload += [('Supplementary Information','03_Supplement/Supplementary_Information_ExpertRebuilt.pdf'),('Supplementary Data 1','03_Supplement/Supplementary_Data_1_Tables_S1-S10.xlsx')]
upload += [(f'Source Data {i}',p.relative_to(OUT).as_posix()) for i,p in enumerate(sorted((OUT/'05_Source_Data').glob('*.zip')),1)]
upload += [('Cover Letter','04_Cover/Cover_Letter_SubmissionCandidate.docx')]
manifest=[dict(order=i+1,portal_role=role,relative=rel,bytes=(OUT/rel).stat().st_size,sha256=sha(OUT/rel),scope='Candidate role mapping; actual portal fields remain to be checked') for i,(role,rel) in enumerate(upload)]
with open(OUT/'06_QA/UPLOAD_CANDIDATE_MANIFEST.tsv','w',encoding='utf-8-sig',newline='') as f:
    w=csv.DictWriter(f,manifest[0].keys(),delimiter='\t');w.writeheader();w.writerows(manifest)
check('Unique_main_manuscript',sum(m['portal_role']=='Main Manuscript' for m in manifest),1)
check('Candidate_upload_files',len(manifest),13)
page_records=[];issues=[]
pdfs=[('Main','01_Manuscript/Manuscript_SubmissionCandidate.pdf'),('Supplement','03_Supplement/Supplementary_Information_ExpertRebuilt.pdf'),('Cover','04_Cover/Cover_Letter_SubmissionCandidate.pdf')]
for role,rel in pdfs:
    d=WORK/'pages'/role;d.mkdir(parents=True,exist_ok=True)
    with pymupdf.open(OUT/rel) as pdf:
        for i,page in enumerate(pdf):
            ptxt=page.get_text();outside=[]
            for b in page.get_text('blocks'):
                if b[0]<-1 or b[1]<-1 or b[2]>page.rect.width+1 or b[3]>page.rect.height+1:outside.append(list(b[:4]))
            if outside or len(ptxt.strip())<20:issues.append(dict(role=role,page=i+1,bounds=outside,chars=len(ptxt.strip())))
            page.get_pixmap(matrix=pymupdf.Matrix(1.7,1.7),alpha=False).save(d/f'page-{i+1:03d}.png')
            page_records.append(dict(role=role,page=i+1,chars=len(ptxt.strip()),image_count=len(page.get_images()),outside_bounds=outside))
        (d/'fulltext.txt').write_text('\n'.join(p.get_text() for p in pdf),encoding='utf-8')
check('PDF_text_bounds_and_nonblank_pages',issues,[])
state=dict(status='PASS' if all(c['passed'] for c in checks) else 'FAIL',checks_passed=sum(c['passed'] for c in checks),checks_total=len(checks),checks=checks,word_counts=dict(title=token(title),abstract=token(abstract),abstract_whitespace=len(abstract.split()),abstract_alphanumeric_pieces=len(re.findall('[A-Za-z0-9]+',abstract)),legends=[token(p) for p in legends]),source_archive_checks=archives,page_records=page_records,upload_total_bytes=sum(r['bytes'] for r in manifest),visual_review='PENDING_MANUAL_PAGE_INSPECTION',publisher_portal_proof='NOT_EXECUTED')
(OUT/'06_QA/MACHINE_QA.json').write_text(json.dumps(state,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({k:state[k] for k in ['status','checks_passed','checks_total','word_counts','upload_total_bytes']}))
print(json.dumps({'pages':{role:sum(r['role']==role for r in page_records) for role,_ in pdfs}}))
assert state['status']=='PASS',[(r['check'],r['actual']) for r in checks if not r['passed']]
