from pathlib import Path
import hashlib,json,csv,zipfile
from docx import Document
import pymupdf
from PIL import Image,ImageDraw,ImageFont
R=Path(r'H:\SCI2\YR1'); O=R/'5_analysis/R7C3G_ExpertNarrative_20261012'; O.mkdir(parents=True,exist_ok=True)
M=R/'0_admin/repositories/igan-open-regulatory-evidence_REBUILD_20261012'
inputs={'manuscript':R/'5_manuscript/R7C3F_SubmissionCandidate_20261012/01_Main/Manuscript.docx','plan':R/'9.Version/PBC_研究计划书_RP_v4_20261011.docx','supplement':R/'5_manuscript/R7C3F_SubmissionCandidate_20261012/03_Supplement/Supplementary_Information_Editable.docx'}
receipt=[]
for k,p in inputs.items():
 d=Document(p); txt='\n'.join(f'[{i:03d}] ({q.style.name}) {q.text}' for i,q in enumerate(d.paragraphs))
 txt+='\n\nTABLES\n'+'\n\n'.join('\n'.join(' | '.join(c.text for c in row.cells) for row in t.rows) for t in d.tables)
 (O/f'{k}_extracted.txt').write_text(txt,encoding='utf-8');receipt.append({'role':k,'path':str(p),'sha256':hashlib.sha256(p.read_bytes()).hexdigest(),'paragraphs':len(d.paragraphs),'tables':len(d.tables)})
z=Path(r'C:\Users\Administrator\Desktop\QiTeng_Academic_Writing_Skill_v0.3.24.2_CompressionGovernance_Release_2026-09-20.zip')
with zipfile.ZipFile(z) as f:
 assert f.testzip() is None
 names=[n for n in f.namelist() if n.endswith('00_CONTROL/VERSION_REGISTRY.json')]
 assert len(names)==1
 v=json.loads(f.read(names[0]));assert v['release_version']=='v0.3.24.2'
 receipt.append({'role':'QiTeng_user_zip','sha256':hashlib.sha256(z.read_bytes()).hexdigest(),'bytes':z.stat().st_size,'crc':'PASS','active_registry':v['release_version']})
(O/'INPUT_RECEIPT.json').write_text(json.dumps(receipt,ensure_ascii=False,indent=2),encoding='utf-8')
images=[];font=ImageFont.truetype('C:/Windows/Fonts/arial.ttf',18)
for p in sorted((R/'4_figures').rglob('*')):
 if not p.is_file() or p.suffix.lower() not in ['.png','.pdf','.svg','.jpg','.jpeg']:continue
 rel=p.relative_to(R/'4_figures'); info={'id':f'I{len(images)+1:02d}','path':str(p),'relative':str(rel),'sha256':hashlib.sha256(p.read_bytes()).hexdigest(),'bytes':p.stat().st_size}
 try:
  if p.suffix.lower()=='.pdf':
   d=pymupdf.open(p);info['pages']=len(d);info['text']='\n'.join(q.get_text() for q in d);pix=d[0].get_pixmap(matrix=pymupdf.Matrix(1.2,1.2),alpha=False);im=Image.frombytes('RGB',[pix.width,pix.height],pix.samples);d.close()
  elif p.suffix.lower()=='.svg':
   info['text']=p.read_text(encoding='utf-8')[:2000]; im=None
  else:
   im=Image.open(p).convert('RGB');info['pixel_size']=list(im.size)
  if im:
   im.thumbnail((860,610));canvas=Image.new('RGB',(900,680),'white');canvas.paste(im,((900-im.width)//2,32));dr=ImageDraw.Draw(canvas);dr.text((12,8),info['id']+'  '+str(rel.parts[0])+' / '+p.stem[:60],font=font,fill='black');canvas.save(O/f"{info['id']}_preview.jpg")
 except Exception as e:info['error']=str(e)
 images.append(info)
(O/'ALL_FIGURE_CONTENT.json').write_text(json.dumps(images,ensure_ascii=False,indent=2),encoding='utf-8')
for start in range(0,len(images),6):
 page=Image.new('RGB',(1800,2040),'#eeeeee')
 for i,v in enumerate(images[start:start+6]):
  p=O/f"{v['id']}_preview.jpg"
  if p.exists():page.paste(Image.open(p),((i%2)*900,(i//2)*680))
 page.save(O/f'CONTACT_{start//6+1:02d}.jpg')
files=[]
for root in [M/'results',M/'protocols']:
 for p in sorted(root.rglob('*')):
  if p.is_file() and p.suffix in ['.tsv','.json','.csv']:
   files.append({'path':str(p.relative_to(M)),'bytes':p.stat().st_size,'head':p.read_text(encoding='utf-8-sig',errors='replace')[:800]})
(O/'RESULT_SCHEMA_INDEX.json').write_text(json.dumps(files,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({'inputs':len(receipt),'images':len(images),'result_objects':len(files),'output':str(O)},ensure_ascii=False))
