#!/usr/bin/env python3
"""Build the R7C3H exact-version manuscript package from frozen R7C3G."""

from __future__ import annotations

import hashlib
import json
import shutil
import zipfile
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph


ROOT = Path(r"H:\SCI2\YR1")
SOURCE = ROOT / "5_manuscript/R7C3G_ExpertNarrative_20261012"
DESTINATION = ROOT / "5_manuscript/R7C3H_StatisticalReproducibilityClosure_20261013"
AUDIT = ROOT / "3_results/00_audit/R7C3H0"
ADJUDICATION = ROOT / "3_results/04_integration/R7C3H0/neff_probe/adjudication"
CODE = ROOT / "2_code/06_intake/r7c3h0"


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def insert_after(paragraph: Paragraph, text: str, style: str = "Normal") -> Paragraph:
    node = OxmlElement("w:p")
    paragraph._p.addnext(node)
    created = Paragraph(node, paragraph._parent)
    created.style = style
    created.add_run(text)
    return created


def find_paragraph(document: Document, needle: str) -> Paragraph:
    matches = [paragraph for paragraph in document.paragraphs if needle in paragraph.text]
    if len(matches) != 1:
        raise RuntimeError(f"Expected one paragraph containing {needle!r}; found {len(matches)}")
    return matches[0]


RESULT_TEXT = (
    "The disease sample-size convention was then tested without changing the primary source-matched analysis. "
    "Gjoka and Cordell fitted SuSiE-RSS to logistic-regression summary statistics using the total sample size of "
    "24,510. A protocol-frozen secondary rerun instead used the case-control effective sample size "
    "4/(1/8,021+1/16,489)=21,584.38 for 33 high-risk comparisons across 14 loci. All 33 final states and all four "
    "manuscript key axes were unchanged. Disease credible-set membership was identical in 164 of 164 "
    "comparison-configuration-set records, and the largest absolute change in the PF10 L=10 best PP.H4 was "
    "1.63x10^-4 (Source Data 3)."
)

METHOD_TEXT = (
    "The disease value N=24,510 followed the exact source analysis of the GJOKA logistic-regression statistics, "
    "rather than a generic case-control transformation [13]. As a bounded robustness test, we replaced it with "
    "the case-control effective sample size Neff=4/(1/Ncase+1/Ncontrol)=21,584.38. The 33 comparisons were frozen "
    "before this rerun and comprised all 18 strict matched-input H4/H3 reversals, four manuscript key axes, six "
    "comparisons nearest an H4 decision boundary and six high-information model-sensitive or uninformative "
    "comparisons, with overlap among these groups. The original N=24,510 results remained primary."
)

LIMIT_TEXT = (
    "The effective-N analysis was a bounded stress test of 33 deliberately high-risk comparisons, not a rerun of "
    "all 642 comparisons and not evidence that total and effective sample sizes are generally interchangeable. "
    "The study also did not compare every available multi-signal colocalization method or ancestry-aware extension; "
    "its method claims are restricted to the frozen ABF and SuSiE-based procedures and the tested LD settings."
)


def edit_manuscript(path: Path) -> None:
    document = Document(path)
    insert_after(find_paragraph(document, "Ten of 12 historical directional transitions"), RESULT_TEXT)
    insert_after(find_paragraph(document, "Disease and QTL association summaries were fine-mapped"), METHOD_TEXT)
    insert_after(find_paragraph(document, "Several limitations define the statistical claim ceiling"), LIMIT_TEXT)
    availability = find_paragraph(document, "Source Data 1 supplies the complete primary machine-readable tables")
    availability.add_run(
        " Source Data 3 contains the frozen 33-comparison disease-N robustness workload, state comparison and "
        "disease credible-set identity audit."
    )
    document.save(path)


def edit_supplement(path: Path) -> None:
    document = Document(path)
    introduction = find_paragraph(document, "This document defines the units, provenance and interpretation limits")
    introduction.text = introduction.text.replace("S1-S11", "S1-S12") + (
        " The S12 disease-N robustness family is supplied in Source Data 3."
    )
    upload_anchor = find_paragraph(document, "Source Data 2 —")
    inserted = insert_after(
        upload_anchor,
        "Source Data 3 — Source_Data_3_Disease_N_Robustness.zip. Frozen 33-comparison effective-N stress test, "
        "disease credible-set identity audit and machine-readable decision state.",
        upload_anchor.style.name,
    )
    s11 = find_paragraph(document, "Evidence family S11")
    s11_index = next(i for i, paragraph in enumerate(document.paragraphs) if paragraph.text == s11.text)
    s11_body = document.paragraphs[s11_index + 1]
    heading = insert_after(s11_body, "Evidence family S12. Disease sample-size source identity and robustness.", "Heading 2")
    insert_after(
        heading,
        "The primary disease N=24,510 reproduces the source GJOKA SuSiE-RSS convention. A prespecified secondary "
        "stress test used Neff=21,584.38 in 33 high-risk comparisons across 14 loci. All 33 final classifications, "
        "all four manuscript key axes and all 164 disease credible-set membership records were unchanged; the "
        "largest absolute PF10 L=10 best-PP.H4 change was 1.63x10^-4. This bounded analysis is neither a correction "
        "to the primary model nor a new validation cohort. File map: Source Data 3.",
    )
    sharing = find_paragraph(document, "The workbook and both Source Data archives")
    sharing.text = sharing.text.replace("both Source Data archives", "the three Source Data archives")
    document.save(path)


def write_markdown_from_docx(docx_path: Path, markdown_path: Path) -> None:
    document = Document(docx_path)
    lines: list[str] = []
    for paragraph in document.paragraphs:
        if not paragraph.text:
            continue
        prefix = "# " if paragraph.style.name == "Title" else "## " if paragraph.style.name == "Heading 1" else "### " if paragraph.style.name == "Heading 2" else ""
        lines.append(prefix + paragraph.text)
    markdown_path.write_text("\n\n".join(lines) + "\n", encoding="utf-8")


def build_source_data_3(destination: Path) -> None:
    staging = destination.parent / "_Source_Data_3_staging"
    staging.mkdir(parents=True, exist_ok=False)
    sources = [
        AUDIT / "R7C3H0_neff_probe_selection_state.json",
        AUDIT / "R7C3H0_neff_probe_comparisons.tsv",
        ADJUDICATION / "R7C3H0_33_neff_state_comparison.tsv",
        ADJUDICATION / "R7C3H0_disease_CS_identity.tsv",
        ADJUDICATION / "R7C3H0_neff_probe_state.json",
    ]
    for source in sources:
        shutil.copy2(source, staging / source.name)
    (staging / "README.md").write_text(
        "# Source Data 3 Disease N robustness\n\n"
        "The primary disease sample size of 24,510 matches the source GJOKA SuSiE-RSS analysis. This archive records "
        "a frozen secondary stress test using case-control Neff=21,584.38. It does not replace the primary model. "
        "The 33-comparison workload was selected before the rerun.\n",
        encoding="utf-8",
    )
    members = sorted(path for path in staging.iterdir() if path.is_file())
    (staging / "checksums.sha256").write_text(
        "\n".join(f"{sha256(path)}  {path.name}" for path in members) + "\n", encoding="utf-8"
    )
    with zipfile.ZipFile(destination, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9) as archive:
        for path in sorted(staging.iterdir()):
            archive.write(path, path.name)
    shutil.rmtree(staging)


def main() -> None:
    if DESTINATION.exists():
        raise FileExistsError(f"Refusing to overwrite existing exact version: {DESTINATION}")
    shutil.copytree(SOURCE, DESTINATION)
    (DESTINATION / "checksums.sha256").unlink(missing_ok=True)
    edit_manuscript(DESTINATION / "01_Manuscript/Manuscript_ExpertRebuilt.docx")
    edit_manuscript(DESTINATION / "01_Manuscript/Manuscript_AuthorReview_InlineFigures.docx")
    edit_supplement(DESTINATION / "03_Supplement/Supplementary_Information_ExpertRebuilt.docx")
    write_markdown_from_docx(
        DESTINATION / "01_Manuscript/Manuscript_ExpertRebuilt.docx",
        DESTINATION / "01_Manuscript/Manuscript_ExpertRebuilt.md",
    )
    build_source_data_3(DESTINATION / "05_Source_Data/Source_Data_3_Disease_N_Robustness.zip")
    for path in sorted(CODE.glob("*.py")):
        shutil.copy2(path, DESTINATION / "09_Code" / path.name)
    start = DESTINATION / "00_START_HERE.md"
    text = start.read_text(encoding="utf-8")
    text = text.replace("R7C3G", "R7C3H").replace(
        "Source Data 1–2", "Source Data 1–3"
    )
    text += (
        "\n\n## R7C3H statistical closure\n\nThe source-matched disease sample size remains N=24,510. "
        "A frozen 33-comparison Neff stress test produced no state or disease credible-set membership changes. "
        "Figures 1–6 therefore remain byte-identical; Source Data 3 records the new robustness analysis.\n"
    )
    start.write_text(text, encoding="utf-8")
    state = {
        "schema": "R7C3H_EXACT_VERSION_1.0",
        "source_package": str(SOURCE),
        "destination": str(DESTINATION),
        "source_matched_disease_N": 24510,
        "secondary_effective_N": 21584.376825785395,
        "state_changes": 0,
        "figure_scientific_bytes_changed": False,
        "author_reconfirmation_required": True,
    }
    (DESTINATION / "08_Protocol/R7C3H_EDITORIAL_STATE.json").write_text(
        json.dumps(state, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print(json.dumps(state, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
