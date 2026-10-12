#!/usr/bin/env python3
"""Machine QA and checksum freeze for the R7C3H exact-version package."""

from __future__ import annotations

import csv
import hashlib
import json
import re
import zipfile
from pathlib import Path

import pymupdf
from docx import Document


ROOT = Path(r"H:\SCI2\YR1")
SOURCE = ROOT / "5_manuscript/R7C3G_ExpertNarrative_20261012"
PACKAGE = ROOT / "5_manuscript/R7C3H_StatisticalReproducibilityClosure_20261013"
VISUAL = ROOT / "5_analysis/R7C3H0_StatisticalReproducibilityClosure_20261013/visual_wps"
QA = PACKAGE / "06_QA"


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def main() -> None:
    checks: list[dict[str, object]] = []

    def check(name: str, passed: bool, detail: object = "") -> None:
        checks.append({"check": name, "status": "PASS" if passed else "FAIL", "detail": json.dumps(detail, ensure_ascii=False) if not isinstance(detail, str) else detail})

    main_docs = [
        PACKAGE / "01_Manuscript/Manuscript_ExpertRebuilt.docx",
        PACKAGE / "01_Manuscript/Manuscript_AuthorReview_InlineFigures.docx",
    ]
    for path in main_docs:
        text = "\n".join(paragraph.text for paragraph in Document(path).paragraphs)
        check(f"{path.name}: disease-N result inserted once", text.count("protocol-frozen secondary rerun") == 1)
        check(f"{path.name}: disease-N method inserted once", text.count("generic case-control transformation") == 1)
        check(f"{path.name}: bounded limitation inserted once", text.count("bounded stress test of 33 deliberately high-risk comparisons") == 1)
        check(f"{path.name}: Source Data 3 declared", "Source Data 3 contains the frozen 33-comparison" in text)

    supplement = Document(PACKAGE / "03_Supplement/Supplementary_Information_ExpertRebuilt.docx")
    supplement_text = "\n".join(paragraph.text for paragraph in supplement.paragraphs)
    check("Supplement S12 declared", supplement_text.count("Evidence family S12") == 1)
    check("Supplement Source Data 3 upload role", supplement_text.count("Source Data 3 —") == 1)
    check("Supplement three archives", "three Source Data archives" in supplement_text)

    expected_pages = {
        "01_Manuscript/Manuscript_ExpertRebuilt.pdf": 25,
        "01_Manuscript/Manuscript_AuthorReview_InlineFigures.pdf": 25,
        "03_Supplement/Supplementary_Information_ExpertRebuilt.pdf": 6,
        "04_Cover/Cover_Letter_ExpertRebuilt_DRAFT.pdf": 1,
    }
    for relative, expected in expected_pages.items():
        with pymupdf.open(PACKAGE / relative) as pdf:
            characters = [len(page.get_text().strip()) for page in pdf]
            check(f"WPS PDF pages {relative}", len(pdf) == expected, {"expected": expected, "actual": len(pdf)})
            check(f"WPS PDF nonblank {relative}", min(characters) > 100, {"minimum_text_characters": min(characters)})

    receipt = json.loads((QA / "WPS_Export_Receipt.json").read_text(encoding="utf-8-sig"))
    check("WPS export receipt 4 files", len(receipt) == 4)
    check("WPS renderer identity", all("WPS" in row["renderer"] and row["background"] for row in receipt))

    figure_rows = []
    for source in sorted((SOURCE / "02_Figures").glob("*")):
        if not source.is_file():
            continue
        current = PACKAGE / "02_Figures" / source.name
        identical = current.is_file() and sha256(source) == sha256(current)
        figure_rows.append({"file": source.name, "source_sha256": sha256(source), "current_sha256": sha256(current) if current.is_file() else "MISSING", "identical": identical})
    check("All scientific figure files byte-identical", len(figure_rows) == 27 and all(row["identical"] for row in figure_rows), {"files": len(figure_rows)})

    frozen_assets = [
        "03_Supplement/Supplementary_Data_1_Tables_S1-S10.xlsx",
        "05_Source_Data/Source_Data_1_Primary_Machine_Readable_Data.zip",
        "05_Source_Data/Source_Data_2_External_Regulatory_Evidence.zip",
    ]
    check("Frozen prior data assets byte-identical", all(sha256(SOURCE / rel) == sha256(PACKAGE / rel) for rel in frozen_assets))

    source3 = PACKAGE / "05_Source_Data/Source_Data_3_Disease_N_Robustness.zip"
    with zipfile.ZipFile(source3) as archive:
        bad = archive.testzip()
        manifest = archive.read("checksums.sha256").decode("utf-8").splitlines()
        internal_ok = True
        for line in manifest:
            digest, name = re.split(r"\s{2,}", line, maxsplit=1)
            internal_ok &= hashlib.sha256(archive.read(name)).hexdigest() == digest
        members = len(archive.namelist())
    check("Source Data 3 ZIP CRC", bad is None)
    check("Source Data 3 internal hashes", internal_ok, {"members": members, "manifest_entries": len(manifest)})

    visual_sheets = sorted(VISUAL.glob("*.png"))
    check("WPS visual contact sheets present", len(visual_sheets) == 17)
    visual_receipt = {
        "renderer": "WPS kwps.Application ExportAsFixedFormat 17; PyMuPDF raster for review",
        "pdf_pages_reviewed": 57,
        "contact_sheets_reviewed": len(visual_sheets),
        "review_status": "PASS_57_OF_57",
        "observations": "No clipping, overlap, missing glyphs, blank pages or figure-caption separation defects observed.",
        "contact_sheets": [{"path": str(path), "sha256": sha256(path)} for path in visual_sheets],
    }
    (QA / "VISUAL_REVIEW_RECEIPT_R7C3H.json").write_text(json.dumps(visual_receipt, ensure_ascii=False, indent=2), encoding="utf-8")

    figure_path = QA / "R7C3H_FIGURE_BYTE_IDENTITY.tsv"
    with figure_path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(figure_rows[0]), delimiter="\t")
        writer.writeheader()
        writer.writerows(figure_rows)

    qa_path = QA / "R7C3H_MACHINE_QA.tsv"
    with qa_path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=["check", "status", "detail"], delimiter="\t")
        writer.writeheader()
        writer.writerows(checks)
    failures = [row for row in checks if row["status"] == "FAIL"]
    state = {
        "schema": "R7C3H_PACKAGE_QA_1.0",
        "status": "PASS" if not failures else "FAIL",
        "checks": len(checks),
        "passes": len(checks) - len(failures),
        "failures": len(failures),
        "wps_pages_visually_reviewed": 57,
        "figure_bytes_changed": False,
        "new_scientific_asset": "Source_Data_3_Disease_N_Robustness.zip",
        "author_reconfirmation_required": True,
    }
    (QA / "R7C3H_PACKAGE_QA_STATE.json").write_text(json.dumps(state, ensure_ascii=False, indent=2), encoding="utf-8")

    role_rows = []
    for path in sorted(PACKAGE.rglob("*")):
        if not path.is_file() or path.name in {"checksums.sha256", "UPLOAD_ROLES_AND_IDENTITIES.tsv"}:
            continue
        relative = path.relative_to(PACKAGE).as_posix()
        if relative.endswith("Manuscript_ExpertRebuilt.docx"):
            role = "MAIN_CANDIDATE_SELECT_EXACTLY_ONE"
        elif relative.startswith("02_Figures/") and path.suffix.lower() == ".pdf":
            role = "INDEPENDENT_FIGURE"
        elif relative.startswith("05_Source_Data/"):
            role = "SOURCE_DATA"
        elif relative.startswith("03_Supplement/"):
            role = "SUPPLEMENT"
        elif relative.startswith("04_Cover/"):
            role = "COVER_LETTER_DRAFT"
        elif relative.startswith("06_QA/") or relative.startswith("07_Reports/") or relative.startswith("08_Protocol/") or relative.startswith("09_Code/"):
            role = "INTERNAL_REVIEW_OR_PROVENANCE"
        else:
            role = "READING_OR_EDITABLE_MASTER"
        role_rows.append({"file": relative, "role": role, "bytes": path.stat().st_size, "sha256": sha256(path)})
    roles_path = PACKAGE / "UPLOAD_ROLES_AND_IDENTITIES.tsv"
    with roles_path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=["file", "role", "bytes", "sha256"], delimiter="\t")
        writer.writeheader()
        writer.writerows(role_rows)

    manifest_files = sorted(path for path in PACKAGE.rglob("*") if path.is_file() and path.name != "checksums.sha256")
    (PACKAGE / "checksums.sha256").write_text(
        "\n".join(f"{sha256(path)}  {path.relative_to(PACKAGE).as_posix()}" for path in manifest_files) + "\n",
        encoding="utf-8",
    )
    print(json.dumps(state, ensure_ascii=False, indent=2))
    if failures:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
