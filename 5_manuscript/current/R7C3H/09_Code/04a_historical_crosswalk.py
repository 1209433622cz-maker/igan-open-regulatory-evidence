from pathlib import Path
import zipfile,csv,io
from docx import Document
R=Path(r'H:\SCI2\YR1');O=R/'5_manuscript/R7C3G_ExpertNarrative_20261012';C=R/'2_code/11_manuscript/R7C3G_ExpertNarrative_20261012'
# Original S10 data/figures refer to a historical layout, which is preserved byte-for-byte.
targets={
'C01':'Figure1a;E01','C02':'Figure1a,d;E02','C03':'Figure2c,d;E04','C04':'Figure2c,d;E04','C05':'Figure2c,d;E04','C06':'Figure3b;E06','C07':'Figure3a;E05','C08':'Figure3a;E05','C09':'Methods mult-signal diagnostic adjudication;FigureS1d','C10':'Methods mult-signal diagnostic adjudication;FigureS1d','C11':'Figure4c;E10','C12':'Figure4c;E10','C13':'Figure4c;FigureS2;E10','C14':'Discussion simulation limits;E10','C15':'Results IL12RB2 external module;Figure5a;E11','C16':'Results aligned-expression direction;Methods external-QTL;E11','C17':'Figure5d;S11 regional-edge family;E13','C18':'Figure3d;E08','C19':'Results FCRL3 direction;E08','C20':'Figure3c;E07','C21':'Results FCRL3 coverage;E12','C22':'Figure6a,b;E14','C23':'Figure2b;FigureS1a;E03','C24':'Figure2b;FigureS1b;E03','C25':'Figure2c,d;E04','C26':'Results/Methods matched-input sensitivity;no separate panel','C27':'Figure4a,b,d;E09',
'MR01':'Figure1a,d;E01/E02','MR02':'Figure2a;E03','MR03':'Results input matching and Methods replay;no separate panel','MR04':'Figure2b;FigureS1a;E03','MR05':'Figure2b;FigureS1b;E03','MR06':'Figure2c,d;E04','MR07':'Results/Methods sensitivity;no separate panel','MR08':'Figure4a,d;E09','MR09':'Figure4b;E09','MR10':'Figures3d/5/6;E08/E11/E14'}
out=[]
with zipfile.ZipFile(O/'05_Source_Data/Source_Data_1_Primary_Machine_Readable_Data.zip') as z:
 for member in ['S10/S10_claim_evidence_ledger.tsv','S10/S10_methods_results_mirror.tsv']:
  rows=list(csv.DictReader(io.StringIO(z.read(member).decode('utf-8-sig')),delimiter='\t'))
  for row in rows:
   key=row.get('claim_id') or row['module_id'];assert key in targets
   out.append({'historical_record_id':key,'source_member':member,'historical_figure_metadata':str(row.get('figure',''))+str(row.get('panel','')) if 'claim_id' in row else row['figure_panel'],'current_destination_and_editorial_id':targets[key],'scientific_content_status':'CHROMATIN_CEILING_SUPERSEDED_BY_S11_REGIONAL_CA_QTL_RECORD;MEDIATION_STILL_UNPROVEN' if key=='C17' else 'SCIENTIFIC_ANCHOR_RETAINED;PANEL_METADATA_RELOCATED','interpretation_note':'C17 position-only audit is historical. S11 now supports NK disease-caQTL regional sharing; no direct eQTL-caQTL posterior. Link best cell is PBMC.' if key=='C17' else 'Current editorial E-series IDs and historical S10 C/MR IDs use separate namespaces. Source workbook and ZIP remain unchanged.'})
assert len(out)==37 and len({x['historical_record_id'] for x in out})==37
p=O/'03_Supplement/Supplementary_Editorial_Panel_Crosswalk.tsv'
with p.open('w',encoding='utf-8-sig',newline='') as f:
 w=csv.DictWriter(f,fieldnames=list(out[0]),delimiter='\t');w.writeheader();w.writerows(out)
# Namespace the new editorial 15-claim map; do not reuse original C01-C27 identifiers.
p=O/'06_QA/CENTRAL_CLAIM_EVIDENCE_CEILING_MAP.tsv'
s=p.read_text(encoding='utf-8-sig')
import re
s=re.sub(r'(?m)^C(\d{2})\t',r'E\1\t',s);p.write_text('\ufeff'+s,encoding='utf-8')
note='The figure and panel columns in the preserved S10 workbook and source ledgers refer to the historical manuscript layout. The companion Supplementary_Editorial_Panel_Crosswalk.tsv maps all 27 historical C-series claims and ten MR-series Methods–Results records to the current figures or text. Current editorial claim IDs use an E-series namespace. Historical C17 retains its earlier positional chromatin ceiling; the S11 regional disease-caQTL evidence updates that ceiling without establishing mediation. Scientific source values and the original workbook/ZIP bytes remain unchanged.'
p=O/'03_Supplement/Supplementary_Information_ExpertRebuilt.docx';d=Document(p)
if not any(x.text==note for x in d.paragraphs):
 # Place this before the source-version explanation and supplementary figure pages.
 for x in d.paragraphs:
  if x.text=='Source-version interpretation':x.insert_paragraph_before(note);break
 d.save(p)
print('S10_CROSSWALK_PASS 27 claims +10 MethodsResults rows; distinct E namespace; source unchanged')
